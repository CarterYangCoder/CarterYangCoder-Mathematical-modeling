"""问题二科研图：温湿场、径向分布与变物性影响。"""
from __future__ import annotations


import argparse
import math
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field
from itertools import zip_longest
from pathlib import Path
from typing import Final

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import load_workbook
from PIL import Image


# 路径、模型常量与绘图参数。

ROOT: Final = Path(__file__).resolve().parent
DEFAULT_RESULTS: Final = ROOT / "result2.xlsx"
DEFAULT_OUTPUT: Final = ROOT / "generated_figures_q2"

TEMPERATURE_SHEET: Final = "温度"
MOISTURE_SHEET: Final = "水分浓度"
INITIAL_T_C: Final = 28.0
INITIAL_C_KG_KG: Final = 2.55
RADIUS_CM: Final = 2.0
RADIUS_M: Final = RADIUS_CM / 100.0
Q2_END_TIME_S: Final = 10_800.0

DEFAULT_MAX_TIME_BINS: Final = 1_440
MAX_TIME_BINS_HARD: Final = 3_000
MAX_HEATMAP_PIXELS: Final = 2_000_000
MAX_PLOT_POINTS_PER_SERIES: Final = 3_000
MIN_PNG_BYTES: Final = 5_000
MAX_PNG_BYTES: Final = 50 * 1024 * 1024
DPI: Final = 600

FULL_WIDTH_IN: Final = 5.71

REPORT_TIMES_S: Final = (0.0, 1_800.0, 3_600.0, 5_400.0, 7_200.0, 9_000.0, Q2_END_TIME_S)
PROPERTY_TIMES_S: Final = (1_800.0, 5_400.0, Q2_END_TIME_S)
STATE_PATH_RADII_CM: Final = (0.0, 0.5, 1.0, 1.5, 2.0)

TEXT: Final = "#1F2933"
MUTED: Final = "#667784"
GRID: Final = "#D7E0E5"
TEMP: Final = "#C84B31"
MOISTURE: Final = "#1677A8"
ENVIRONMENT: Final = "#5C6770"
CENTRE: Final = "#3B6EA5"
SURFACE: Final = "#C85045"
PROFILE_COLOURS: Final = ("#3B4CC0", "#3387BC", "#65B96E", "#D9B534", "#E07A3F", "#C85145", "#7A3B70")
PROFILE_STYLES: Final = ("-", "--", "-.", ":", (0, (4, 1.5)), (0, (2, 1.3)), (0, (6, 1.5, 1, 1.5)))
RADIUS_COLOURS: Final = ("#3B4CC0", "#3387BC", "#65B96E", "#E07A3F", "#8A3B6D")


@dataclass(frozen=True)
class EnvironmentHistory:


    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture_kg_kg: np.ndarray


@dataclass
class FieldBuckets:


    bins: int
    radii_cm: np.ndarray
    count: np.ndarray = field(init=False)
    time_first_s: np.ndarray = field(init=False)
    time_last_s: np.ndarray = field(init=False)
    temperature_sum: np.ndarray = field(init=False)
    moisture_sum: np.ndarray = field(init=False)
    temperature_min: np.ndarray = field(init=False)
    temperature_max: np.ndarray = field(init=False)
    moisture_min: np.ndarray = field(init=False)
    moisture_max: np.ndarray = field(init=False)

    def __post_init__(self) -> None:
        radial_count = self.radii_cm.size
        self.count = np.zeros(self.bins, dtype=np.int64)
        self.time_first_s = np.full(self.bins, np.nan, dtype=float)
        self.time_last_s = np.full(self.bins, np.nan, dtype=float)
        shape = (self.bins, radial_count)
        self.temperature_sum = np.zeros(shape, dtype=float)
        self.moisture_sum = np.zeros(shape, dtype=float)
        self.temperature_min = np.full(shape, np.inf, dtype=float)
        self.temperature_max = np.full(shape, -np.inf, dtype=float)
        self.moisture_min = np.full(shape, np.inf, dtype=float)
        self.moisture_max = np.full(shape, -np.inf, dtype=float)

    def add(self, bin_index: int, time_s: float, temperature: np.ndarray, moisture: np.ndarray) -> None:
        if self.count[bin_index] == 0:
            self.time_first_s[bin_index] = time_s
        self.time_last_s[bin_index] = time_s
        self.count[bin_index] += 1
        self.temperature_sum[bin_index] += temperature
        self.moisture_sum[bin_index] += moisture
        self.temperature_min[bin_index] = np.minimum(self.temperature_min[bin_index], temperature)
        self.temperature_max[bin_index] = np.maximum(self.temperature_max[bin_index], temperature)
        self.moisture_min[bin_index] = np.minimum(self.moisture_min[bin_index], moisture)
        self.moisture_max[bin_index] = np.maximum(self.moisture_max[bin_index], moisture)

    def finalise(self) -> "AggregatedField":
        if np.any(self.count <= 0):
            missing = np.flatnonzero(self.count <= 0).tolist()
            raise ValueError(f"时间聚合桶为空，拒绝绘图：{missing[:10]}")
        divisor = self.count[:, None]
        return AggregatedField(
            time_s=(self.time_first_s + self.time_last_s) / 2.0,
            time_edges_s=time_bin_edges(self.time_first_s, self.time_last_s),
            temperature=self.temperature_sum / divisor,
            moisture=self.moisture_sum / divisor,
            temperature_min=self.temperature_min,
            temperature_max=self.temperature_max,
            moisture_min=self.moisture_min,
            moisture_max=self.moisture_max,
        )


@dataclass(frozen=True)
class AggregatedField:


    time_s: np.ndarray
    time_edges_s: np.ndarray
    temperature: np.ndarray
    moisture: np.ndarray
    temperature_min: np.ndarray
    temperature_max: np.ndarray
    moisture_min: np.ndarray
    moisture_max: np.ndarray


@dataclass(frozen=True)
class Q2Data:


    radii_cm: np.ndarray
    field: AggregatedField
    snapshots: dict[float, tuple[np.ndarray, np.ndarray]]
    raw_rows: int
    rendered_bins: int
    initial_row_added: bool


# 结果表读取与结构校验。

def as_finite_float(value: object, name: str) -> float:

    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{name}不是可用数值：{value!r}") from exc
    if not math.isfinite(number):
        raise ValueError(f"{name}不是有限数值：{number!r}")
    return number


def numeric_header(row: Sequence[object], sheet_name: str) -> np.ndarray:

    if len(row) < 3:
        raise ValueError(f"工作表“{sheet_name}”列数不足，无法组成时间—径向场。")
    if row[0] is None:
        raise ValueError(f"工作表“{sheet_name}”缺少时间列标题。")
    radii = np.asarray([as_finite_float(value, f"{sheet_name}表头") for value in row[1:]], dtype=float)
    if not np.all(np.diff(radii) > 0):
        raise ValueError(f"工作表“{sheet_name}”的径向坐标必须严格递增。")
    if not np.isclose(radii[0], 0.0, atol=1e-10) or not np.isclose(radii[-1], RADIUS_CM, atol=1e-10):
        raise ValueError(
            f"工作表“{sheet_name}”应覆盖 0–{RADIUS_CM:g} cm；当前为 {radii[0]:g}–{radii[-1]:g} cm。"
        )
    return radii


def validate_input_schema(workbook_path: Path) -> tuple[np.ndarray, int]:

    if not workbook_path.is_file():
        raise FileNotFoundError(f"未找到问题 2 结果文件：{workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    try:
        missing = [name for name in (TEMPERATURE_SHEET, MOISTURE_SHEET) if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"结果文件缺少必需工作表：{missing}；实际工作表为 {workbook.sheetnames}")
        temperature_sheet = workbook[TEMPERATURE_SHEET]
        moisture_sheet = workbook[MOISTURE_SHEET]
        if temperature_sheet.max_row < 2 or moisture_sheet.max_row < 2:
            raise ValueError("结果工作表至少应包含表头和一条数值记录。")
        if temperature_sheet.max_row != moisture_sheet.max_row:
            raise ValueError("温度和水分浓度工作表的记录行数不一致。")
        if temperature_sheet.max_column != moisture_sheet.max_column:
            raise ValueError("温度和水分浓度工作表的列数不一致。")
        temperature_header = next(temperature_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        moisture_header = next(moisture_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        radii_t = numeric_header(temperature_header, TEMPERATURE_SHEET)
        radii_c = numeric_header(moisture_header, MOISTURE_SHEET)
        if not np.array_equal(radii_t, radii_c):
            raise ValueError("温度和水分浓度工作表的径向坐标不一致。")
        raw_rows = temperature_sheet.max_row - 1
        return radii_t, raw_rows
    finally:
        workbook.close()


def load_environment(path: Path) -> EnvironmentHistory:

    if not path.is_file():
        raise FileNotFoundError(f"未找到环境输入附件：{path}")
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook.active
        if sheet.max_row < 3 or sheet.max_column < 3:
            raise ValueError("环境附件至少需要时间、温度和水分浓度三列。")
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        expected_keywords = (("时间",), ("温度",), ("水分", "浓度"))
        for index, words in enumerate(expected_keywords):
            label = str(header[index]) if index < len(header) else ""
            if not any(word in label for word in words):
                raise ValueError(f"环境附件第 {index + 1} 列标题应包含 {words}，当前为 {label!r}。")

        time_values: list[float] = []
        temperature_values: list[float] = []
        moisture_values: list[float] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=3, values_only=True), start=2):
            if all(value is None for value in row):
                continue
            time_values.append(as_finite_float(row[0], f"环境附件第 {row_number} 行时间"))
            temperature_values.append(as_finite_float(row[1], f"环境附件第 {row_number} 行温度"))
            moisture_values.append(as_finite_float(row[2], f"环境附件第 {row_number} 行水分浓度"))
        time_s = np.asarray(time_values, dtype=float)
        temperature_c = np.asarray(temperature_values, dtype=float)
        moisture_kg_kg = np.asarray(moisture_values, dtype=float)
        if time_s.size < 2 or not np.all(np.diff(time_s) > 0):
            raise ValueError("环境附件时间必须至少有两个严格递增节点。")
        if not (np.all(np.isfinite(temperature_c)) and np.all(np.isfinite(moisture_kg_kg))):
            raise ValueError("环境附件存在非有限温度或水分浓度。")
        return EnvironmentHistory(time_s, temperature_c, moisture_kg_kg)
    finally:
        workbook.close()


def _coerce_result_row(row: Sequence[object], expected_width: int, row_number: int, field_name: str) -> tuple[float, np.ndarray]:
    if len(row) != expected_width:
        raise ValueError(f"{field_name}第 {row_number} 行列数为 {len(row)}，应为 {expected_width}。")
    time_s = as_finite_float(row[0], f"{field_name}第 {row_number} 行时间")
    values = np.asarray(
        [as_finite_float(value, f"{field_name}第 {row_number} 行径向值") for value in row[1:]], dtype=float
    )
    return time_s, values


def time_bin_edges(first_s: np.ndarray, last_s: np.ndarray) -> np.ndarray:

    if first_s.size == 1:
        width = max(1.0, last_s[0] - first_s[0] + 1.0)
        return np.asarray([first_s[0] - width / 2.0, last_s[0] + width / 2.0])
    centres = (first_s + last_s) / 2.0
    edges = np.empty(centres.size + 1, dtype=float)
    edges[1:-1] = (centres[:-1] + centres[1:]) / 2.0
    edges[0] = first_s[0]
    edges[-1] = last_s[-1]
    if not np.all(np.diff(edges) > 0):
        raise ValueError("聚合后的时间边界不是严格递增，拒绝绘图。")
    return edges


def stream_read(
    workbook_path: Path,
    radii_cm: np.ndarray,
    raw_rows: int,
    max_time_bins: int,
    requested_times_s: Iterable[float],
) -> Q2Data:

    if not 1 <= max_time_bins <= MAX_TIME_BINS_HARD:
        raise ValueError(f"--max-time-bins 必须在 1–{MAX_TIME_BINS_HARD} 之间。")
    requested = {float(value) for value in requested_times_s}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    try:
        temperature_sheet = workbook[TEMPERATURE_SHEET]
        moisture_sheet = workbook[MOISTURE_SHEET]
        temperature_rows = temperature_sheet.iter_rows(min_row=2, values_only=True)
        moisture_rows = moisture_sheet.iter_rows(min_row=2, values_only=True)
        try:
            first_temperature_raw = next(temperature_rows)
            first_moisture_raw = next(moisture_rows)
        except StopIteration as exc:
            raise ValueError("结果工作表没有数值记录。") from exc

        expected_width = radii_cm.size + 1
        first_time_t, first_temperature = _coerce_result_row(
            first_temperature_raw, expected_width, 2, TEMPERATURE_SHEET
        )
        first_time_c, first_moisture = _coerce_result_row(first_moisture_raw, expected_width, 2, MOISTURE_SHEET)
        if not np.isclose(first_time_t, first_time_c, atol=1e-10):
            raise ValueError("温度和水分浓度的首个时间点不一致。")
        if np.isclose(first_time_t, 0.0, atol=1e-10):
            initial_row_added = False
        elif np.isclose(first_time_t, 1.0, atol=1e-10):
            initial_row_added = True
        else:
            raise ValueError(
                f"结果时间应从模型初始时刻 0 s 或第一秒 1 s 开始，当前首点为 {first_time_t:g} s。"
            )

        total_records = raw_rows + int(initial_row_added)
        bins = min(max_time_bins, total_records)
        if bins * radii_cm.size > MAX_HEATMAP_PIXELS:
            raise ValueError(
                f"热图将有 {bins * radii_cm.size:,} 个像素，超过 {MAX_HEATMAP_PIXELS:,} 上限；"
                "请减小 --max-time-bins。"
            )
        buckets = FieldBuckets(bins=bins, radii_cm=radii_cm)
        snapshots: dict[float, tuple[np.ndarray, np.ndarray]] = {}

        def add_record(serial: int, time_s: float, temperature: np.ndarray, moisture: np.ndarray) -> None:
            if not (np.all(np.isfinite(temperature)) and np.all(np.isfinite(moisture))):
                raise ValueError(f"t={time_s:g} s 存在非有限结果，拒绝绘图。")
            bucket = min((serial * bins) // total_records, bins - 1)
            buckets.add(bucket, time_s, temperature, moisture)
            for requested_time in requested:
                if np.isclose(time_s, requested_time, atol=1e-10):
                    snapshots[requested_time] = (temperature.copy(), moisture.copy())

        serial = 0
        if initial_row_added:
            add_record(
                serial,
                0.0,
                np.full(radii_cm.size, INITIAL_T_C, dtype=float),
                np.full(radii_cm.size, INITIAL_C_KG_KG, dtype=float),
            )
            serial += 1

        previous_time = -np.inf
        for source_index, pair in enumerate(
            _paired_rows(first_temperature_raw, first_moisture_raw, temperature_rows, moisture_rows), start=2
        ):
            temperature_raw, moisture_raw = pair
            time_t, temperature = _coerce_result_row(temperature_raw, expected_width, source_index, TEMPERATURE_SHEET)
            time_c, moisture = _coerce_result_row(moisture_raw, expected_width, source_index, MOISTURE_SHEET)
            if not np.isclose(time_t, time_c, atol=1e-10):
                raise ValueError(f"第 {source_index} 行温度与水分浓度的时间不一致：{time_t} vs {time_c}。")
            if not time_t > previous_time:
                raise ValueError(f"第 {source_index} 行时间未严格递增：{time_t} <= {previous_time}。")
            previous_time = time_t
            add_record(serial, time_t, temperature, moisture)
            serial += 1

        if serial != total_records:
            raise ValueError(f"流式读取到 {serial} 条记录，但工作表元数据声称有 {total_records} 条。")
        if not np.isclose(previous_time, Q2_END_TIME_S, atol=1e-10):
            raise ValueError(
                f"问题 2 的模型输出应覆盖到 {Q2_END_TIME_S:g} s（3 h），当前最后时刻为 {previous_time:g} s。"
            )
        missing_snapshots = sorted(time for time in requested if time <= previous_time and time not in snapshots)
        if missing_snapshots:
            raise ValueError(
                "结果中没有精确的报告时刻，脚本不会插值伪造径向剖面："
                + ", ".join(f"{time:g} s" for time in missing_snapshots)
            )
        field_data = buckets.finalise()
        return Q2Data(
            radii_cm=radii_cm,
            field=field_data,
            snapshots=snapshots,
            raw_rows=raw_rows,
            rendered_bins=bins,
            initial_row_added=initial_row_added,
        )
    finally:
        workbook.close()


def _paired_rows(
    first_temperature: Sequence[object],
    first_moisture: Sequence[object],
    remaining_temperature: Iterable[Sequence[object]],
    remaining_moisture: Iterable[Sequence[object]],
) -> Iterable[tuple[Sequence[object], Sequence[object]]]:

    yield first_temperature, first_moisture
    sentinel = object()
    for temperature_row, moisture_row in zip_longest(remaining_temperature, remaining_moisture, fillvalue=sentinel):
        if temperature_row is sentinel or moisture_row is sentinel:
            raise ValueError("温度和水分浓度工作表的数据行数不一致。")
        yield temperature_row, moisture_row


# 绘图基础函数与模型派生量。

def _pick_font(candidates: tuple[str, ...], fallback: str) -> str:

    installed = {font.name for font in font_manager.fontManager.ttflist}
    return next((name for name in candidates if name in installed), fallback)


# 数学标签中的中文字符范围。
_CJK_RANGES: Final = (
    (0x2E80, 0x2EFF),
    (0x3000, 0x303F),
    (0x3400, 0x4DBF),
    (0x4E00, 0x9FFF),
    (0xF900, 0xFAFF),
    (0xFE30, 0xFE4F),
    (0xFF00, 0xFFEF),
)


def _is_cjk_symbol(symbol: str) -> bool:

    if len(symbol) != 1:
        return False
    codepoint = ord(symbol)
    return any(low <= codepoint <= high for low, high in _CJK_RANGES)


def _enable_chinese_in_mathtext(chinese_font: str) -> None:


    from matplotlib import _mathtext

    if getattr(_mathtext.UnicodeFonts._get_glyph, "_accepts_chinese", False):
        return
    font_path = font_manager.findfont(
        font_manager.FontProperties(family=[chinese_font]), fallback_to_default=False
    )
    original_get_glyph = _mathtext.UnicodeFonts._get_glyph

    def _get_glyph(self, fontname, font_class, sym):
        if fontname == "default" and _is_cjk_symbol(sym):
            font = self._get_font(font_path)
            codepoint = ord(sym)
            if font.get_char_index(codepoint):
                return font, codepoint, False
        return original_get_glyph(self, fontname, font_class, sym)

    _get_glyph._accepts_chinese = True
    _mathtext.UnicodeFonts._get_glyph = _get_glyph


def configure_matplotlib() -> None:


    latin_font = _pick_font(("Times New Roman", "Nimbus Roman", "Liberation Serif"), "DejaVu Serif")
    chinese_font = _pick_font(("SimSun", "Songti SC", "Noto Serif CJK SC", "SimHei"), "DejaVu Sans")

    mpl.rcParams.update(
        {
            # 建立中西文回退链，确保数学标签中的中文正常显示。
            "font.family": [latin_font, chinese_font, "DejaVu Serif"],
            "font.serif": [latin_font, chinese_font, "DejaVu Serif"],
            "axes.unicode_minus": False,
            "mathtext.fontset": "custom",
            "mathtext.rm": latin_font,
            "mathtext.it": f"{latin_font}:italic",
            "mathtext.bf": f"{latin_font}:bold",
            "mathtext.sf": latin_font,
            "mathtext.tt": latin_font,
            "mathtext.default": "it",
            "mathtext.fallback": "stix",
            "font.size": 8.2,
            "axes.linewidth": 0.75,
            "axes.labelcolor": TEXT,
            "xtick.color": TEXT,
            "ytick.color": TEXT,
            "xtick.major.width": 0.65,
            "ytick.major.width": 0.65,
            "legend.frameon": False,
            "figure.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )
    _enable_chinese_in_mathtext(chinese_font)
    print(
        f"字体：中文 {chinese_font}（常规字重）；西文、数字与单位 {latin_font}；"
        f"数学变量 {latin_font} 斜体（mathtext custom）。"
    )


def ordered_cmap(colours: Sequence[str], name: str) -> LinearSegmentedColormap:
    return LinearSegmentedColormap.from_list(name, list(colours), N=256)


TEMPERATURE_CMAP: Final = ordered_cmap(("#F8F5E6", "#F4C46A", "#D96B3D", "#8F1D2C"), "q2_temperature")
MOISTURE_CMAP: Final = ordered_cmap(("#F3F7F8", "#A6D5E4", "#4B9CC5", "#155B8B"), "q2_moisture")
# 变物性场统一采用低值绿、高值红的视觉梯度；各面板色标仍独立对应物理量。
PROPERTY_CMAP: Final = ordered_cmap(
    ("#1A9850", "#A6D96A", "#FEE08B", "#EF6548", "#B2182B"),
    "q2_property_green_yellow_red",
)


def radial_edges_cm(radii_cm: np.ndarray) -> np.ndarray:

    edges = np.empty(radii_cm.size + 1, dtype=float)
    edges[0] = radii_cm[0]
    edges[-1] = radii_cm[-1]
    edges[1:-1] = (radii_cm[:-1] + radii_cm[1:]) / 2.0
    if not np.all(np.diff(edges) > 0):
        raise ValueError("径向网格边界不严格递增。")
    return edges


def series_indices(radii_cm: np.ndarray, values: Iterable[float]) -> list[int]:
    indices: list[int] = []
    for radius in values:
        matched = np.flatnonzero(np.isclose(radii_cm, radius, atol=1e-10))
        if matched.size != 1:
            raise ValueError(f"结果网格中没有唯一的 r={radius:g} cm 节点。")
        indices.append(int(matched[0]))
    return indices


def hours_label(time_s: float) -> str:
    return f"{time_s / 3600.0:g} h"


def add_panel_label(axis: plt.Axes, label: str) -> None:
    axis.text(-0.13, 1.035, label, transform=axis.transAxes, fontsize=9.2, fontweight="bold", color=TEXT)


def style_axis(axis: plt.Axes, *, grid: bool = True) -> None:
    axis.spines[["top", "right"]].set_visible(False)
    if grid:
        axis.grid(axis="y", color=GRID, linewidth=0.55, alpha=0.8)
    axis.tick_params(direction="out", length=3, pad=2)


def box_axis(axis: plt.Axes) -> None:

    for spine in ("top", "right"):
        axis.spines[spine].set_visible(True)
        axis.spines[spine].set_color(TEXT)
        axis.spines[spine].set_linewidth(0.8)
    axis.tick_params(top=True, right=True, labeltop=False, labelright=False)


def nearest_environment(environment: EnvironmentHistory, time_s: np.ndarray) -> tuple[np.ndarray, np.ndarray]:

    if time_s[0] < environment.time_s[0] or time_s[-1] > environment.time_s[-1]:
        raise ValueError(
            "结果时间超出附件 1 的环境时程，脚本不会外推环境边界。"
        )
    return (
        np.interp(time_s, environment.time_s, environment.temperature_c),
        np.interp(time_s, environment.time_s, environment.moisture_kg_kg),
    )


def material_properties(temperature_c: np.ndarray, moisture_kg_kg: np.ndarray) -> dict[str, np.ndarray]:

    if np.any(moisture_kg_kg <= 0):
        raise ValueError("扩散系数关系要求 C>0；当前绘图数据存在非正水分浓度。")
    temperature_k = temperature_c + 273.15
    rho = 650.0 + 128.0 * moisture_kg_kg
    cp = 1450.0 + 2736.0 * moisture_kg_kg / (1.0 + moisture_kg_kg)
    conductivity = 0.21 + 0.38 * moisture_kg_kg / (1.0 + moisture_kg_kg)
    diffusivity = 2.4e-3 * np.exp(-0.45 / moisture_kg_kg) * np.exp(-3850.0 / temperature_k)
    heat_capacity = rho * cp
    diffusivity_initial = 2.4e-3 * np.exp(-0.45 / INITIAL_C_KG_KG) * np.exp(-3850.0 / (INITIAL_T_C + 273.15))
    log_moisture_factor = 0.45 * (1.0 / INITIAL_C_KG_KG - 1.0 / moisture_kg_kg)
    log_temperature_factor = 3850.0 * (1.0 / (INITIAL_T_C + 273.15) - 1.0 / temperature_k)
    moisture_factor = np.exp(log_moisture_factor)
    temperature_factor = np.exp(log_temperature_factor)
    return {
        "rho": rho,
        "cp": cp,
        "k": conductivity,
        "D": diffusivity,
        "W": heat_capacity,
        "D_ratio": moisture_factor * temperature_factor,
        "I_C": moisture_factor,
        "I_T": temperature_factor,
    }


def area_weighted_mean(field: np.ndarray, radii_cm: np.ndarray) -> np.ndarray:

    radius_m = radii_cm / 100.0
    integrand = field * radius_m[None, :]
    integral = np.trapezoid(integrand, x=radius_m, axis=1)
    return 2.0 * integral / (radius_m[-1] ** 2)


def finite_difference_per_minute(field: np.ndarray, time_s: np.ndarray) -> np.ndarray:
    if time_s.size < 3:
        raise ValueError("至少需要三个时间桶才能计算响应速率。")
    if not np.all(np.diff(time_s) > 0):
        raise ValueError("计算响应速率前时间必须严格递增。")
    return np.gradient(field, time_s, axis=0) * 60.0


def select_safe_contour_levels(values: np.ndarray, candidates: Sequence[float]) -> list[float]:
    lower = float(np.nanmin(values))
    upper = float(np.nanmax(values))
    levels = sorted({float(level) for level in candidates if lower < level < upper})
    if len(levels) < 2:
        raise ValueError(f"可用等值线级数不足：数据范围为 {lower:g}–{upper:g}。")
    return levels


def decimate_indices(length: int, maximum: int = MAX_PLOT_POINTS_PER_SERIES) -> np.ndarray:

    if length <= maximum:
        return np.arange(length)
    indices = np.linspace(0, length - 1, maximum, dtype=int)
    indices[0] = 0
    indices[-1] = length - 1
    return np.unique(indices)


# 各图对应论文中的独立证据链。

def figure_temperature_heatmap(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axis = plt.subplots(figsize=(FULL_WIDTH_IN, 3.33), layout="constrained")
    mesh = axis.pcolormesh(
        data.field.time_edges_s / 3600.0,
        radial_edges_cm(data.radii_cm),
        data.field.temperature.T,
        shading="auto",
        cmap=TEMPERATURE_CMAP,
    )
    contours = axis.contour(
        data.field.time_s / 3600.0,
        data.radii_cm,
        data.field.temperature.T,
        levels=select_safe_contour_levels(data.field.temperature, (30, 35, 40, 45, 49)),
        colors="#54252D",
        linewidths=0.55,
        alpha=0.82,
    )
    axis.clabel(contours, inline=True, fontsize=6.5, fmt="%g")
    colourbar = fig.colorbar(mesh, ax=axis, pad=0.018)
    colourbar.set_label("温度 T (°C；时间桶均值)")
    axis.set(xlabel="时间 t (h)", ylabel="到中心的距离 r (cm)", ylim=(0, RADIUS_CM))
    style_axis(axis, grid=False)
    return fig


def figure_moisture_heatmap(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axis = plt.subplots(figsize=(FULL_WIDTH_IN, 3.33), layout="constrained")
    mesh = axis.pcolormesh(
        data.field.time_edges_s / 3600.0,
        radial_edges_cm(data.radii_cm),
        data.field.moisture.T,
        shading="auto",
        cmap=MOISTURE_CMAP,
    )
    contours = axis.contour(
        data.field.time_s / 3600.0,
        data.radii_cm,
        data.field.moisture.T,
        levels=select_safe_contour_levels(data.field.moisture, (2.4, 2.2, 2.0, 1.8, 1.6, 1.4, 1.2)),
        colors="#17476A",
        linewidths=0.55,
        alpha=0.84,
    )
    axis.clabel(contours, inline=True, fontsize=6.5, fmt="%.1f")
    colourbar = fig.colorbar(mesh, ax=axis, pad=0.018)
    colourbar.set_label("水分浓度 C (kg/kg，干基；时间桶均值)")
    axis.set(xlabel="时间 t (h)", ylabel="到中心的距离 r (cm)", ylim=(0, RADIUS_CM))
    style_axis(axis, grid=False)
    return fig


def figure_temperature_moisture_heatmaps(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axes = plt.subplots(1, 2, figsize=(FULL_WIDTH_IN, 2.45), sharey=True, layout="constrained")
    time_edges_h = data.field.time_edges_s / 3600.0
    time_h = data.field.time_s / 3600.0
    edges_r = radial_edges_cm(data.radii_cm)
    panels = (
        (
            "(a) 温度场",
            data.field.temperature,
            TEMPERATURE_CMAP,
            (30, 35, 40, 45, 49),
            "#54252D",
            "%g",
            "温度 T (°C)",
            (INITIAL_T_C, 50.0),
            (28.0, 30.0, 35.0, 40.0, 45.0, 50.0),
            None,
        ),
        (
            "(b) 水分浓度场",
            data.field.moisture,
            MOISTURE_CMAP,
            (2.4, 2.2, 2.0, 1.8, 1.6, 1.4, 1.2),
            "#17476A",
            "%.1f",
            "水分浓度 C (kg/kg，干基)",
            None,
            None,
            (1.2, (2.55, 1.86)),
        ),
    )
    for axis, (title, values, cmap, levels, contour_colour, contour_format, colourbar_label, colour_limits, colour_ticks, manual_label) in zip(axes, panels):
        mesh_kwargs: dict[str, float] = {}
        if colour_limits is not None:
            mesh_kwargs = {"vmin": colour_limits[0], "vmax": colour_limits[1]}
        mesh = axis.pcolormesh(time_edges_h, edges_r, values.T, shading="auto", cmap=cmap, **mesh_kwargs)
        contours = axis.contour(
            time_h,
            data.radii_cm,
            values.T,
            levels=select_safe_contour_levels(values, levels),
            colors=contour_colour,
            linewidths=0.58,
            alpha=0.82,
        )
        if manual_label is None:
            axis.clabel(contours, inline=True, fontsize=6.0, fmt=contour_format)
        else:
            manual_level, manual_position = manual_label
            auto_levels = [level for level in contours.levels if not np.isclose(level, manual_level)]
            axis.clabel(contours, levels=auto_levels, inline=True, fontsize=6.0, fmt=contour_format)
            axis.clabel(
                contours,
                levels=[manual_level],
                manual=[manual_position],
                inline=True,
                fontsize=6.0,
                fmt=contour_format,
            )
        colourbar = fig.colorbar(mesh, ax=axis, pad=0.018, shrink=0.92)
        if colour_ticks is not None:
            colourbar.set_ticks(colour_ticks)
        colourbar.set_label(colourbar_label, fontsize=7.0)
        colourbar.ax.tick_params(labelsize=6.5)
        axis.set(title=title, xlabel="时间 t (h)", ylim=(0, RADIUS_CM))
        style_axis(axis, grid=False)
        box_axis(axis)
    axes[0].set_ylabel("到中心的距离 r (cm)")
    return fig


def _profile_figure(data: Q2Data, field_index: int) -> plt.Figure:
    field_name, unit, colour = (
        ("温度 T", "°C", TEMP) if field_index == 0 else ("水分浓度 C", "kg/kg（干基）", MOISTURE)
    )
    fig, axis = plt.subplots(figsize=(FULL_WIDTH_IN, 3.33), layout="constrained")
    available_times = [time for time in REPORT_TIMES_S if time in data.snapshots]
    for index, time_s in enumerate(available_times):
        values = data.snapshots[time_s][field_index]
        axis.plot(
            data.radii_cm,
            values,
            color=PROFILE_COLOURS[index],
            linestyle=PROFILE_STYLES[index],
            linewidth=1.25,
            label=hours_label(time_s),
        )
    axis.set(xlabel="到中心的距离 r (cm)", ylabel=f"{field_name} ({unit})", xlim=(0, RADIUS_CM))
    axis.legend(title="时刻", ncols=2, loc="best", fontsize=7.1, title_fontsize=7.1)
    style_axis(axis)
    return fig


def figure_temperature_profiles(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    return _profile_figure(data, 0)


def figure_moisture_profiles(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    return _profile_figure(data, 1)


def figure_temperature_moisture_profiles(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axes = plt.subplots(1, 2, figsize=(FULL_WIDTH_IN, 2.53), sharex=True)
    available_times = [time for time in REPORT_TIMES_S if time in data.snapshots]
    panels = (
        ("(a) 温度径向分布", 0, "温度 T (°C)"),
        ("(b) 水分浓度径向分布", 1, "水分浓度 C (kg/kg，干基)"),
    )
    for axis, (title, field_index, ylabel) in zip(axes, panels):
        for index, time_s in enumerate(available_times):
            axis.plot(
                data.radii_cm,
                data.snapshots[time_s][field_index],
                color=PROFILE_COLOURS[index],
                linestyle=PROFILE_STYLES[index],
                linewidth=1.25,
                label=hours_label(time_s),
            )
        axis.set(title=title, xlabel="到中心的距离 r (cm)", ylabel=ylabel, xlim=(0, RADIUS_CM))
        style_axis(axis)
        box_axis(axis)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(
        handles,
        labels,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.995),
        ncols=4,
        fontsize=6.5,
        handlelength=2.35,
        columnspacing=0.9,
    )
    fig.subplots_adjust(left=0.105, right=0.985, bottom=0.20, top=0.80, wspace=0.34)
    return fig


def _plot_binned_series(
    axis: plt.Axes,
    time_h: np.ndarray,
    mean: np.ndarray,
    lower: np.ndarray,
    upper: np.ndarray,
    *,
    label: str,
    colour: str,
    linestyle: str = "-",
    envelope_label: str | None = None,
) -> None:

    axis.plot(time_h, mean, color=colour, linewidth=1.28, linestyle=linestyle, label=label, zorder=3)
    if np.any(upper > lower):
        axis.fill_between(
            time_h,
            lower,
            upper,
            color=colour,
            alpha=0.09,
            linewidth=0,
            zorder=1,
            label=envelope_label,
        )


def figure_temperature_response(data: Q2Data, environment: EnvironmentHistory) -> plt.Figure:

    fig, axis = plt.subplots(figsize=(FULL_WIDTH_IN, 3.33), layout="constrained")
    time_h = data.field.time_s / 3600.0
    air_temperature, _ = nearest_environment(environment, data.field.time_s)
    _plot_binned_series(
        axis, time_h, data.field.temperature[:, 0], data.field.temperature_min[:, 0], data.field.temperature_max[:, 0],
        label="中心 r = 0 cm", colour=CENTRE, envelope_label="时间桶 min–max（降采样范围）",
    )
    _plot_binned_series(
        axis, time_h, data.field.temperature[:, -1], data.field.temperature_min[:, -1], data.field.temperature_max[:, -1],
        label="表面 r = 2 cm", colour=SURFACE,
    )
    axis.plot(time_h, air_temperature, color=ENVIRONMENT, linestyle="--", linewidth=1.05, label="环境温度 $T_a$")
    axis.set(xlabel="时间 t (h)", ylabel="温度 (°C)")
    axis.legend(loc="best", fontsize=7.2)
    style_axis(axis)
    return fig


def figure_moisture_response(data: Q2Data, environment: EnvironmentHistory) -> plt.Figure:

    fig, axes = plt.subplots(2, 1, figsize=(FULL_WIDTH_IN, 4.32), sharex=True, height_ratios=(2.0, 1.0), layout="constrained")
    time_h = data.field.time_s / 3600.0
    _, equilibrium_moisture = nearest_environment(environment, data.field.time_s)
    _plot_binned_series(
        axes[0], time_h, data.field.moisture[:, 0], data.field.moisture_min[:, 0], data.field.moisture_max[:, 0],
        label="中心 r = 0 cm", colour=CENTRE, envelope_label="时间桶 min–max（降采样范围）",
    )
    _plot_binned_series(
        axes[0], time_h, data.field.moisture[:, -1], data.field.moisture_min[:, -1], data.field.moisture_max[:, -1],
        label="表面 r = 2 cm", colour=SURFACE,
    )
    axes[0].set_ylabel("药材 C (kg/kg，干基)")
    axes[0].legend(loc="best", fontsize=7.2)
    axes[1].plot(time_h, equilibrium_moisture, color=ENVIRONMENT, linewidth=1.15, label="环境平衡水分 $C_e$")
    axes[1].set(xlabel="时间 t (h)", ylabel="$C_e$ (kg/kg)")
    axes[1].legend(loc="best", fontsize=7.2)
    for label, axis in zip(("(a)", "(b)"), axes):
        add_panel_label(axis, label)
        style_axis(axis)
    return fig


def figure_radial_gradients(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axes = plt.subplots(2, 1, figsize=(FULL_WIDTH_IN, 4.32), sharex=True, layout="constrained")
    time_h = data.field.time_s / 3600.0
    thermal_gradient = data.field.temperature[:, -1] - data.field.temperature[:, 0]
    moisture_gradient = data.field.moisture[:, 0] - data.field.moisture[:, -1]
    axes[0].plot(time_h, thermal_gradient, color=TEMP, linewidth=1.35)
    axes[0].set_ylabel("$T_s - T_c$ (°C)")
    axes[1].plot(time_h, moisture_gradient, color=MOISTURE, linewidth=1.35)
    axes[1].set(xlabel="时间 t (h)", ylabel="$C_c - C_s$ (kg/kg)")
    for label, axis in zip(("(a)", "(b)"), axes):
        add_panel_label(axis, label)
        style_axis(axis)
    return fig


def figure_temperature_lag(data: Q2Data, environment: EnvironmentHistory) -> plt.Figure:

    fig, axes = plt.subplots(2, 1, figsize=(FULL_WIDTH_IN, 4.32), sharex=True, layout="constrained")
    time_h = data.field.time_s / 3600.0
    air_temperature, _ = nearest_environment(environment, data.field.time_s)
    axes[0].plot(time_h, data.field.temperature[:, -1] - air_temperature, color=SURFACE, linewidth=1.25)
    axes[1].plot(time_h, data.field.temperature[:, 0] - air_temperature, color=CENTRE, linewidth=1.25)
    axes[0].axhline(0.0, color=MUTED, linewidth=0.7, linestyle="--")
    axes[1].axhline(0.0, color=MUTED, linewidth=0.7, linestyle="--")
    axes[0].set_ylabel("$T_s - T_a$ (°C)")
    axes[1].set(xlabel="时间 t (h)", ylabel="$T_c - T_a$ (°C)")
    for label, axis in zip(("(a)", "(b)"), axes):
        add_panel_label(axis, label)
        style_axis(axis)
    return fig


def figure_moisture_lag(data: Q2Data, environment: EnvironmentHistory) -> plt.Figure:

    fig, axes = plt.subplots(2, 1, figsize=(FULL_WIDTH_IN, 4.32), sharex=True, layout="constrained")
    time_h = data.field.time_s / 3600.0
    _, equilibrium_moisture = nearest_environment(environment, data.field.time_s)
    axes[0].plot(time_h, data.field.moisture[:, -1] - equilibrium_moisture, color=SURFACE, linewidth=1.25)
    axes[1].plot(time_h, data.field.moisture[:, 0] - equilibrium_moisture, color=CENTRE, linewidth=1.25)
    axes[0].axhline(0.0, color=MUTED, linewidth=0.7, linestyle="--")
    axes[1].axhline(0.0, color=MUTED, linewidth=0.7, linestyle="--")
    axes[0].set_ylabel("$C_s - C_e$ (kg/kg)")
    axes[1].set(xlabel="时间 t (h)", ylabel="$C_c - C_e$ (kg/kg)")
    for label, axis in zip(("(a)", "(b)"), axes):
        add_panel_label(axis, label)
        style_axis(axis)
    return fig


def figure_property_fields(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    properties = material_properties(data.field.temperature, data.field.moisture)
    panels = (
        ("rho", "密度 $\\rho$ (kg/m³)", 1.0),
        ("k", "导热系数 $k$ (W/(m·K))", 1.0),
        ("W", "体积热容 $W=\\rho c_p$ (10⁶ J/(m³·K))", 1.0e6),
        ("D", "扩散系数 $D$ (10⁻⁸ m²/s)", 1.0e-8),
    )
    fig, axes = plt.subplots(2, 2, figsize=(FULL_WIDTH_IN, 4.52), sharex=True, sharey=True, layout="constrained")
    edges_r = radial_edges_cm(data.radii_cm)
    for label, axis, (key, colourbar_label, display_scale) in zip(("(a)", "(b)", "(c)", "(d)"), axes.flat, panels):
        values = properties[key] / display_scale
        mesh = axis.pcolormesh(
            data.field.time_edges_s / 3600.0,
            edges_r,
            values.T,
            shading="auto",
            cmap=PROPERTY_CMAP,
            edgecolors="#303030",
            linewidth=0.025,
            antialiased=True,
        )
        colourbar = fig.colorbar(mesh, ax=axis, pad=0.015, shrink=0.9)
        colourbar.set_label(colourbar_label, fontsize=7.1)
        colourbar.ax.tick_params(labelsize=6.6)
        add_panel_label(axis, label)
        style_axis(axis, grid=False)
    for axis in axes[-1, :]:
        axis.set_xlabel("时间 t (h)")
    for axis in axes[:, 0]:
        axis.set_ylabel("r (cm)")
    return fig


def figure_diffusivity_decomposition(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axes = plt.subplots(1, 3, figsize=(FULL_WIDTH_IN, 2.70), sharex=True)
    panels = (
        ("I_C", "水分因子 $I_C$", 1.0),
        ("I_T", "温度因子 $I_T$", 1.0),
        ("D_ratio", "综合影响 $D/D_0$", 1.0),
    )
    time_colours = ("#2878B5", "#D9981E", "#C85145")
    for panel_index, (axis, (key, ylabel, reference_value)) in enumerate(zip(axes, panels)):
        for time_index, time_s in enumerate(PROPERTY_TIMES_S):
            if time_s not in data.snapshots:
                continue
            temperature, moisture = data.snapshots[time_s]
            values = material_properties(temperature, moisture)[key]
            axis.plot(
                data.radii_cm,
                values,
                color=time_colours[time_index],
                linewidth=1.45,
                label=hours_label(time_s),
                zorder=3,
            )
        axis.axhline(reference_value, color=MUTED, linewidth=0.65, linestyle="--", zorder=1)
        axis.set(xlabel="到中心的距离 r (cm)", ylabel=ylabel, xlim=(0, RADIUS_CM))
        axis.set_title(f"({chr(ord('a') + panel_index)}) {ylabel}", fontsize=8.1, pad=7)
        style_axis(axis)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(
        handles,
        labels,
        title="时刻",
        loc="upper center",
        bbox_to_anchor=(0.5, 0.995),
        ncols=3,
        fontsize=6.7,
        title_fontsize=6.7,
        handlelength=2.4,
        columnspacing=1.35,
    )
    fig.subplots_adjust(left=0.095, right=0.985, bottom=0.245, top=0.750, wspace=0.50)
    fig.text(
        0.5,
        0.030,
        "$D_0=D(C_0,T_0)$，$C_0=2.55$ kg/kg，$T_0=301.15$ K，且 $D/D_0=I_CI_T$",
        ha="center",
        va="bottom",
        fontsize=6.6,
        color=MUTED,
    )
    return fig


def figure_state_paths(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axis = plt.subplots(figsize=(FULL_WIDTH_IN, 3.51), layout="constrained")
    indices = series_indices(data.radii_cm, STATE_PATH_RADII_CM)
    reduced = decimate_indices(data.field.time_s.size)
    start_temperature, start_moisture = data.snapshots[0.0]
    end_temperature, end_moisture = data.snapshots[Q2_END_TIME_S]
    for colour, radius_cm, index in zip(RADIUS_COLOURS, STATE_PATH_RADII_CM, indices):
        temperature = data.field.temperature[reduced, index]
        moisture = data.field.moisture[reduced, index]
        axis.plot(moisture, temperature, color=colour, linewidth=1.15, label=f"r = {radius_cm:g} cm")
        axis.scatter(start_moisture[index], start_temperature[index], s=12, color=colour, marker="o", zorder=4)
        axis.scatter(end_moisture[index], end_temperature[index], s=18, color=colour, marker="s", zorder=4)
    axis.text(0.015, 0.03, "○ 起始   □ 终止", transform=axis.transAxes, color=MUTED, fontsize=7.1)
    axis.set(xlabel="水分浓度 C (kg/kg，干基)", ylabel="温度 T (°C)")
    axis.legend(loc="best", fontsize=7.1, ncols=2)
    style_axis(axis)
    return fig


def figure_temperature_fronts(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axis = plt.subplots(figsize=(FULL_WIDTH_IN, 3.33), layout="constrained")
    levels = select_safe_contour_levels(data.field.temperature, (30, 35, 40, 45, 49))
    contours = axis.contour(
        data.field.time_s / 3600.0,
        data.radii_cm,
        data.field.temperature.T,
        levels=levels,
        cmap=ordered_cmap(("#F0B35C", "#C74D3A", "#79233A"), "q2_temperature_fronts"),
        linewidths=1.15,
    )
    axis.clabel(contours, inline=True, fontsize=7, fmt=lambda value: f"{value:g} °C")
    axis.set(xlabel="时间 t (h)", ylabel="到中心的距离 r (cm)", ylim=(0, RADIUS_CM))
    style_axis(axis)
    return fig


def figure_moisture_fronts(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axis = plt.subplots(figsize=(FULL_WIDTH_IN, 3.33), layout="constrained")
    levels = select_safe_contour_levels(data.field.moisture, (2.4, 2.2, 2.0, 1.8, 1.6, 1.4, 1.2))
    contours = axis.contour(
        data.field.time_s / 3600.0,
        data.radii_cm,
        data.field.moisture.T,
        levels=levels,
        cmap=ordered_cmap(("#8CC7DC", "#2E80B6", "#17476A"), "q2_moisture_fronts"),
        linewidths=1.15,
    )
    axis.clabel(contours, inline=True, fontsize=7, fmt=lambda value: f"{value:.1f}")
    axis.set(xlabel="时间 t (h)", ylabel="到中心的距离 r (cm)", ylim=(0, RADIUS_CM))
    style_axis(axis)
    return fig


def figure_response_rates(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axes = plt.subplots(2, 1, figsize=(FULL_WIDTH_IN, 4.32), sharex=True, layout="constrained")
    time_h = data.field.time_s / 3600.0
    dtemperature_dt = finite_difference_per_minute(data.field.temperature, data.field.time_s)
    dmoisture_dt = finite_difference_per_minute(data.field.moisture, data.field.time_s)
    for axis, rates, ylabel in (
        (axes[0], dtemperature_dt, "$dT/dt$ (°C/min)"),
        (axes[1], dmoisture_dt, "$dC/dt$ (kg/(kg·min))"),
    ):
        axis.plot(time_h, rates[:, 0], color=CENTRE, linewidth=1.2, label="中心")
        axis.plot(time_h, rates[:, -1], color=SURFACE, linewidth=1.2, label="表面")
        axis.axhline(0.0, color=MUTED, linewidth=0.7, linestyle="--")
        axis.set_ylabel(ylabel)
        axis.legend(loc="best", fontsize=7.0)
        style_axis(axis)
    axes[1].set_xlabel("时间 t (h)")
    add_panel_label(axes[0], "(a)")
    add_panel_label(axes[1], "(b)")
    return fig


def figure_area_means(data: Q2Data, _: EnvironmentHistory) -> plt.Figure:

    fig, axes = plt.subplots(2, 1, figsize=(FULL_WIDTH_IN, 4.32), sharex=True, layout="constrained")
    time_h = data.field.time_s / 3600.0
    mean_temperature = area_weighted_mean(data.field.temperature, data.radii_cm)
    mean_moisture = area_weighted_mean(data.field.moisture, data.radii_cm)
    axes[0].plot(time_h, mean_temperature, color=TEMP, linewidth=1.3)
    axes[1].plot(time_h, mean_moisture, color=MOISTURE, linewidth=1.3)
    axes[0].set_ylabel("截面估计平均 T (°C)")
    axes[1].set(xlabel="时间 t (h)", ylabel="截面估计平均 C (kg/kg)")
    for label, axis in zip(("(a)", "(b)"), axes):
        add_panel_label(axis, label)
        style_axis(axis)
    fig.text(
        0.5,
        0.005,
        "由输出径向节点按 2r/R² 的梯形积分估计；不作为有限体积守恒残差。",
        ha="center",
        va="bottom",
        fontsize=6.7,
        color=MUTED,
    )
    return fig


FIGURES: Final[dict[str, tuple[str, Callable[[Q2Data, None], plt.Figure]]]] = {
    "temperature_moisture_heatmaps": (
        "fig_q2_temperature_moisture_fields.png",
        figure_temperature_moisture_heatmaps,
    ),
    "temperature_moisture_profiles": (
        "fig_q2_temperature_moisture_profiles.png",
        figure_temperature_moisture_profiles,
    ),
    "property_fields": ("fig_q2_variable_properties.png", figure_property_fields),
    "diffusivity_decomposition": ("fig_q2_diffusivity_factors.png", figure_diffusivity_decomposition),
}


# 图片导出与命令行入口。

def validate_output_directory(output_dir: Path, *, overwrite: bool) -> None:

    allowed = {".png", ".jpg", ".jpeg"}
    if output_dir.exists() and not output_dir.is_dir():
        raise ValueError(f"输出路径不是目录：{output_dir}")
    if output_dir.exists():
        unexpected = [path for path in output_dir.rglob("*") if path.is_file() and path.suffix.lower() not in allowed]
        if unexpected:
            preview = ", ".join(str(path) for path in unexpected[:5])
            raise RuntimeError(f"正式输出目录含有非法非图片产物：{preview}")
    if not overwrite and output_dir.exists() and any(output_dir.glob("*.png")):
        raise FileExistsError("输出目录已有 PNG 图片。请使用新目录，或明确传入 --overwrite。")


def save_png_or_jpg(fig: plt.Figure, destination: Path, *, overwrite: bool) -> None:

    if destination.suffix.lower() != ".png":
        raise ValueError(f"本脚本默认只输出 PNG，拒绝目标扩展名：{destination.suffix}")
    if destination.exists() and not overwrite:
        raise FileExistsError(f"拒绝静默覆盖既有图片：{destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(destination, format="png", dpi=DPI, facecolor="white", edgecolor="white")
    plt.close(fig)
    try:
        with Image.open(destination) as image:
            image.verify()
        with Image.open(destination) as image:
            if image.format != "PNG":
                raise RuntimeError(f"导出格式为 {image.format!r}，不是 PNG。")
            if image.width < 1_000 or image.height < 1_000:
                raise RuntimeError(f"图片像素不足：{image.width}×{image.height}。")
            if image.mode not in {"RGB", "RGBA"}:
                raise RuntimeError(f"图片色彩模式异常：{image.mode}。")
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    size = destination.stat().st_size
    if not MIN_PNG_BYTES <= size <= MAX_PNG_BYTES:
        destination.unlink(missing_ok=True)
        raise RuntimeError(f"PNG 文件大小 {size:,} bytes 超出 {MIN_PNG_BYTES:,}–{MAX_PNG_BYTES:,} 的 QA 范围。")


def validate_image_outputs(output_dir: Path, expected: Sequence[Path]) -> None:

    allowed = {".png", ".jpg", ".jpeg"}
    unexpected = [path for path in output_dir.rglob("*") if path.is_file() and path.suffix.lower() not in allowed]
    if unexpected:
        raise RuntimeError(f"正式输出目录存在非法产物：{unexpected}")
    missing = [path for path in expected if not path.is_file()]
    if missing:
        raise RuntimeError(f"缺少预期 PNG 输出：{missing}")
    for path in expected:
        with Image.open(path) as image:
            if image.format != "PNG" or image.width < 1_000 or image.height < 1_000:
                raise RuntimeError(f"图片终检失败：{path}")


def parse_figure_selection(value: str) -> list[str]:
    if value.strip().lower() == "all":
        return list(FIGURES)
    names = [part.strip() for part in value.split(",") if part.strip()]
    unknown = sorted(set(names) - set(FIGURES))
    if unknown:
        raise ValueError(f"未知图形名称：{unknown}。可用名称：{', '.join(FIGURES)}")
    if not names:
        raise ValueError("--figures 不能为空。")
    return names


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="为 CUMCM 问题 2 生成只含 PNG 的科研绘图。")
    parser.add_argument("--results", type=Path, default=DEFAULT_RESULTS, help="问题 2 的 result2.xlsx 路径")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="仅容纳 PNG/JPG 的正式输出目录")
    parser.add_argument(
        "--figures",
        default="all",
        help="all，或逗号分隔的图形名称。可用名称：" + ", ".join(FIGURES),
    )
    parser.add_argument(
        "--max-time-bins",
        type=int,
        default=DEFAULT_MAX_TIME_BINS,
        help=f"流式时间聚合上限（默认 {DEFAULT_MAX_TIME_BINS}，硬上限 {MAX_TIME_BINS_HARD}）",
    )
    parser.add_argument("--overwrite", action="store_true", help="明确允许替换所选图形同名的既有 PNG")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    results_path = args.results.resolve()
    output_dir = args.output.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    if output_dir == results_path.parent:
        raise ValueError("输出目录不能与结果 XLSX 输入文件所在目录相同。")
    selected = parse_figure_selection(args.figures)
    configure_matplotlib()
    validate_output_directory(output_dir, overwrite=args.overwrite)
    radii_cm, raw_rows = validate_input_schema(results_path)
    requested_times = set(REPORT_TIMES_S) | set(PROPERTY_TIMES_S)
    data = stream_read(results_path, radii_cm, raw_rows, args.max_time_bins, requested_times)

    print(f"输入结果：{results_path}")
    print(f"读取策略：openpyxl read_only + iter_rows；原始记录 {data.raw_rows:,} 行，径向节点 {data.radii_cm.size} 个。")
    print(
        f"绘图聚合：{data.rendered_bins:,} 个时间桶，热图 {data.rendered_bins * data.radii_cm.size:,} 单元；"
        "每桶保存均值及 min–max（带状区仅表示降采样范围，不是模型不确定性）。"
    )
    print(f"初始行：{'已在内存补入模型给定 t=0 状态' if data.initial_row_added else '结果文件已提供 t=0 状态'}。")

    expected_paths: list[Path] = []
    for name in selected:
        filename, builder = FIGURES[name]
        destination = output_dir / filename
        figure = builder(data, None)
        save_png_or_jpg(figure, destination, overwrite=args.overwrite)
        expected_paths.append(destination)
        print(f"已通过 PNG QA：{destination}")
    validate_image_outputs(output_dir, expected_paths)
    print(f"输出终检通过：{len(expected_paths)} 个 PNG；正式输出目录未发现非法非图片产物。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
