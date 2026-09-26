"""问题一至问题四共用的圆柱药材干燥模型核心。"""
from __future__ import annotations


# 问题一：圆柱径向有限体积离散。C 为干基含水率，环带状态为体积平均值。

from dataclasses import dataclass, field
import hashlib
import time
from pathlib import Path
from typing import Callable

import numpy as np
from numpy.polynomial.legendre import leggauss
from scipy import special
from scipy.linalg import solve_banded
from scipy.optimize import brentq


@dataclass(frozen=True)
class Parameters:
    radius_m: float = 0.02
    length_m: float = 0.25
    initial_temperature_C: float = 28.0
    initial_moisture_kg_kg: float = 2.55
    density_kg_m3: float = 820.0
    heat_capacity_J_kgK: float = 2600.0
    conductivity_W_mK: float = 0.36
    heat_transfer_W_m2K: float = 25.0
    moisture_transfer_m_s: float = 8e-7
    diffusivity_prefactor_m2_s: float = 7e-9
    diffusivity_exponent_kg_kg: float = 0.89

    @property
    def alpha(self):
        return self.conductivity_W_mK / (self.density_kg_m3*self.heat_capacity_J_kgK)

    @property
    def bi_heat(self):
        return self.heat_transfer_W_m2K*self.radius_m/self.conductivity_W_mK


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class LinearDrive:

    def __init__(self, times, values):
        self.times = np.asarray(times, dtype=float)
        self.values = np.asarray(values, dtype=float)
        if (self.times.ndim != 1 or self.times.size < 2 or
                self.times.shape != self.values.shape or
                not np.all(np.isfinite(self.times)) or
                not np.all(np.isfinite(self.values)) or
                not np.all(np.diff(self.times) > 0)):
            raise ValueError("Invalid or duplicate/missing driver nodes")

    def __call__(self, t):
        tt = np.asarray(t, dtype=float)
        if np.any(tt < self.times[0]) or np.any(tt > self.times[-1]):
            raise ValueError("Extrapolation is prohibited")
        ans = np.interp(tt, self.times, self.values)
        return float(ans) if ans.ndim == 0 else ans

    def next_knot(self, t: float):
        idx = np.searchsorted(self.times, t, side="right")
        return float(self.times[idx]) if idx < len(self.times) else np.inf


def read_drives(root: Path, config: dict):
    import openpyxl
    spec = config["inputs"]
    path = root / spec["environment"]
    if sha256(path) != spec["environment_sha256"]:
        raise ValueError("Original environmental workbook fingerprint changed")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[spec["environment_sheet"]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    vals = np.array([row[:3] for row in rows], dtype=float)
    wb.close()
    if vals.ndim != 2 or vals.shape[1] != 3 or not np.all(np.isfinite(vals)):
        raise ValueError("Missing or non-numeric environmental node")
    if not np.array_equal(vals[:, 0], np.arange(0., 14400.1, 60.)):
        raise ValueError("Environmental times differ from the audited 60 s sequence")
    q1 = vals[vals[:, 0] <= 1800].copy()
    if len(q1) != 31:
        raise ValueError("Insufficient Q1 environmental coverage")
    return LinearDrive(q1[:, 0], q1[:, 1]), LinearDrive(q1[:, 0], q1[:, 2]), q1


class RadialModes:


    def __init__(self, radius, diffusivity, bi, modes=64):
        if radius <= 0 or diffusivity <= 0 or bi < 0 or modes < 1:
            raise ValueError("Invalid radial mode parameters")
        self.radius, self.diffusivity, self.bi = radius, diffusivity, bi
        self.modes = int(modes)
        if bi == 0:
            self.roots = self.coefficients = self.beta = np.empty(0)
            return
        z0 = special.jn_zeros(0, modes)
        z1 = special.jn_zeros(1, modes-1) if modes > 1 else np.empty(0)
        left = np.r_[0., z1]
        f = lambda lam: lam*special.j1(lam)-bi*special.j0(lam)
        self.brackets = np.column_stack([left, z0])
        self.roots = np.array([brentq(f, lo, hi, xtol=5e-15,
                                     rtol=4*np.finfo(float).eps)
                               for lo, hi in self.brackets])
        if not np.all(np.diff(self.roots) > 0):
            raise ArithmeticError("Unsorted or repeated Robin root")
        lam = self.roots
        j0, j1 = special.j0(lam), special.j1(lam)
        self.coefficients = 2*j1/(lam*(j0*j0+j1*j1))
        self.beta = diffusivity*(lam/radius)**2
        self.root_residual = lam*j1-bi*j0

    def amplitudes(self, times, drive: LinearDrive, initial):
        ts = np.asarray(times, dtype=float)
        if (ts.ndim != 1 or np.any(ts < 0) or np.any(np.diff(ts) < 0)
                or drive.times[0] != 0):
            raise ValueError("Queries must be ordered and start within a drive from zero")
        drive(ts)
        if self.bi == 0:
            return np.zeros((len(ts), 0))
        b = self.coefficients*(initial-drive(0.))
        now = 0.
        out = []
        for target in ts:
            while now < target:
                endpoint = min(float(target), drive.next_knot(now))
                j = min(np.searchsorted(drive.times, now, side="right")-1,
                        len(drive.times)-2)
                slope = ((drive.values[j+1]-drive.values[j]) /
                         (drive.times[j+1]-drive.times[j]))
                dt = endpoint-now
                decay = np.exp(-self.beta*dt)
                integ = -np.expm1(-self.beta*dt)/self.beta
                b = b*decay-self.coefficients*slope*integ
                now = endpoint
            out.append(b.copy())
        return np.asarray(out)

    def values(self, times, radii, drive: LinearDrive, initial):
        ts, rr = np.asarray(times, float), np.asarray(radii, float)
        if np.any(rr < 0) or np.any(rr > self.radius):
            raise ValueError("Radius outside the cylinder")
        if self.bi == 0:
            return np.full((len(ts), len(rr)), initial, dtype=float)
        b = self.amplitudes(ts, drive, initial)
        basis = special.j0(np.outer(self.roots, rr/self.radius))
        result = drive(ts)[:, None] + b @ basis
        # 初值按题设直接给定，角点处允许与边界条件不相容。
        result[ts == 0, :] = initial
        return result

    def volume_mean(self, times, drive, initial):
        if self.bi == 0:
            return np.full(len(times), initial)
        avg_basis = 2*special.j1(self.roots)/self.roots
        ans = drive(times)+self.amplitudes(times, drive, initial)@avg_basis
        ans[np.asarray(times) == 0] = initial
        return ans


_GL_X, _GL_W = leggauss(12)


class DiffusionLaw:
    def __init__(self, prefactor=7e-9, exponent=.89):
        if prefactor <= 0 or exponent < 0:
            raise ValueError("Invalid diffusion law")
        self.prefactor, self.exponent = prefactor, exponent

    def D(self, concentration):
        c = np.asarray(concentration, dtype=float)
        if np.any(c < 0) or not np.all(np.isfinite(c)):
            raise ValueError("Nonphysical concentration passed to D; no clipping")
        ans = np.zeros_like(c)
        if self.exponent == 0:
            ans[...] = self.prefactor
        else:
            positive = c > 0
            ans[positive] = self.prefactor*np.exp(-self.exponent/c[positive])
        return float(ans) if ans.ndim == 0 else ans

    def phi(self, concentration):
        c = np.asarray(concentration, dtype=float)
        if np.any(c < 0) or not np.all(np.isfinite(c)):
            raise ValueError("Nonphysical concentration passed to Phi")
        ans = np.zeros_like(c)
        if self.exponent == 0:
            ans = self.prefactor*c
        else:
            positive = c > 0
            # 采用与积分形式等价的积分势函数闭式表达式。
            ans[positive] = (self.prefactor*c[positive] *
                             special.expn(2, self.exponent/c[positive]))
        return float(ans) if ans.ndim == 0 else ans

    def difference(self, left, right):
        a, b = np.broadcast_arrays(np.asarray(left, float), np.asarray(right, float))
        if np.any(a < 0) or np.any(b < 0) or not np.all(np.isfinite(a+b)):
            raise ValueError("Nonphysical Phi difference arguments")
        if self.exponent == 0:
            return self.prefactor*(a-b)
        ans = np.asarray(self.phi(a)-self.phi(b))
        close = (np.abs(a-b) <= .2*np.minimum(a, b)) & (a != b)
        # 仅在被积函数平缓时使用数值求积，避免低含水率下的误差放大。
        viable = close & (a > self.exponent/700) & (b > self.exponent/700)
        gentle = np.zeros_like(close)
        gentle[viable] = np.abs(self.exponent/a[viable]-self.exponent/b[viable]) <= 1.
        close = viable & gentle
        if np.any(close):
            mid = (a[close]+b[close])/2
            half = (a[close]-b[close])/2
            args = mid[:, None]+half[:, None]*_GL_X
            ans[close] = half*(self.D(args)@_GL_W)
        return float(ans) if ans.ndim == 0 else ans


class RingGrid:
    def __init__(self, cells, radius=.02, length=.25):
        if cells < 3 or radius <= 0 or length <= 0:
            raise ValueError("Invalid ring geometry")
        self.cells, self.radius, self.length = int(cells), radius, length
        self.faces = np.linspace(0., radius, cells+1)
        self.centres = (self.faces[:-1]+self.faces[1:])/2
        self.volumes = np.pi*length*np.diff(self.faces**2)
        self.areas = 2*np.pi*length*self.faces
        self.tau = self.areas[1:-1]/np.diff(self.centres)
        self.delta = radius-self.centres[-1]

    def mean_of(self, fn, order=16):
        x, w = leggauss(order)
        lo, hi = self.faces[:-1]**2, self.faces[1:]**2
        radii = np.sqrt((lo+hi)[:, None]/2+(hi-lo)[:, None]/2*x)
        return fn(radii)@w/2

    def point_weights(self, radii):


        rr = np.asarray(radii, float)
        if np.any(rr < 0) or np.any(rr > self.radius):
            raise ValueError("Query outside radius")
        out = np.zeros((len(rr), self.cells))
        edges = self.faces**2
        for j, r in enumerate(rr):
            if r == self.radius:
                continue
            if r == 0:
                out[j, :2] = [1.25, -.25]
                continue
            idx = min(np.searchsorted(self.faces, r, side="right")-1,
                      self.cells-1)
            start = max(0, min(idx-1, self.cells-3))
            ids = np.arange(start, start+3)
            origin = (edges[idx]+edges[idx+1])/2
            scale = edges[start+3]-edges[start]
            lo, hi = (edges[ids]-origin)/scale, (edges[ids+1]-origin)/scale
            moments = np.column_stack([np.ones(3), (lo+hi)/2,
                                       (lo*lo+lo*hi+hi*hi)/3])
            s = (r*r-origin)/scale
            out[j, ids] = np.linalg.solve(moments.T, [1., s, s*s])
        return out

    def reconstruct(self, averages, radii, surface):
        ans = np.asarray(averages) @ self.point_weights(radii).T
        ans[..., np.asarray(radii) == self.radius] = np.asarray(surface)[..., None]
        return ans


@dataclass
class SurfaceState:

    concentration: float
    outward_flux: float
    derivative_flux_last: float
    residual_flux: float
    root_error_C: float
    bracket_low: float
    bracket_high: float
    iterations: int


class MoistureFV:
    def __init__(self, grid: RingGrid, law: DiffusionLaw, transfer=8e-7,
                 root_xtol=5e-15, nonlinear_tol=2e-13):
        self.grid, self.law, self.transfer = grid, law, float(transfer)
        self.root_xtol, self.nonlinear_tol = root_xtol, nonlinear_tol
        if transfer < 0:
            raise ValueError("Negative transfer coefficient")

    def surface(self, last, external):
        if last < 0 or external < 0:
            raise ValueError("Negative state at surface")
        lo, hi = min(last, external), max(last, external)
        h, delta, law = self.transfer, self.grid.delta, self.law
        if h == 0 or last == external:
            return SurfaceState(last, 0., 0. if h == 0 else
                                h*law.D(last)/(law.D(last)+h*delta),
                                0., 0., lo, hi, 0)
        scale = law.prefactor
        def g(c):
            return law.difference(c, last)/scale + h*delta/scale*(c-external)
        cs, info = brentq(g, lo, hi, xtol=self.root_xtol,
                         rtol=4*np.finfo(float).eps, full_output=True)
        if not info.converged or not lo <= cs <= hi:
            raise ArithmeticError("Surface bracket solve failed")
        dcs = law.D(cs)
        flux = h*(cs-external)
        residual = law.difference(cs, last)/delta+flux
        deriv = h*law.D(last)/(dcs+h*delta)
        c_error = abs(residual)/(dcs/delta+h)
        if c_error > 10*self.root_xtol + 2e-14*abs(cs):
            raise ArithmeticError("Surface residual exceeds concentration budget")
        return SurfaceState(cs, flux, deriv, residual, c_error, lo, hi,
                            int(info.iterations))

    def fluxes_and_jacobian(self, state, external):

        grid, law = self.grid, self.law
        c = np.asarray(state)
        internal = grid.tau*law.difference(c[:-1], c[1:])
        surf = self.surface(float(c[-1]), float(external))
        flux = np.r_[0., internal, grid.areas[-1]*surf.outward_flux]
        divergence = np.diff(flux)/grid.volumes
        d = law.D(c)
        diagonal = np.zeros_like(c)
        diagonal[:-1] += grid.tau*d[:-1]
        diagonal[1:] += grid.tau*d[1:]
        diagonal[-1] += grid.areas[-1]*surf.derivative_flux_last
        diagonal /= grid.volumes
        lower = -grid.tau*d[:-1]/grid.volumes[1:]
        upper = -grid.tau*d[1:]/grid.volumes[:-1]
        return divergence, (lower, diagonal, upper), surf

    def solve_balance(self, old, dt, external, source=None, steady=False):


        old = np.asarray(old, float)
        c = old.copy()
        src = np.zeros_like(c) if source is None else np.asarray(source, float)
        weight = 0. if steady else 1.
        factor = 1. if steady else dt
        evaluations = 0
        for iteration in range(25):
            div, (lower, diagonal, upper), surf = self.fluxes_and_jacobian(c, external)
            evaluations += 1
            residual = weight*(c-old)+factor*(div-src)
            norm = float(np.max(np.abs(residual)))
            if norm <= self.nonlinear_tol:
                return c, surf, {"iterations": iteration, "residual_C": norm,
                                  "flux_evaluations": evaluations}
            ab = np.zeros((3, len(c)))
            ab[0, 1:] = factor*upper
            ab[1] = weight+factor*diagonal
            ab[2, :-1] = factor*lower
            change = solve_banded((1, 1), ab, -residual, check_finite=False)
            damping = 1.
            # 仅在全步可能越过零值时施加非负约束，避免微小修正造成数值溢出。
            negative = change < -.99*c
            if np.any(negative):
                damping = min(1., float(np.min(-.99*c[negative]/change[negative])))
            for _ in range(22):
                candidate = c+damping*change
                if np.any(candidate < 0) or not np.all(np.isfinite(candidate)):
                    damping *= .5
                    continue
                new_div, _, _ = self.fluxes_and_jacobian(candidate, external)
                evaluations += 1
                new_res = weight*(candidate-old)+factor*(new_div-src)
                if np.max(np.abs(new_res)) <= (1.-1e-4*damping)*norm:
                    c = candidate
                    break
                damping *= .5
            else:
                raise ArithmeticError(f"Newton line search failed, residual={norm:.3e}")
        raise ArithmeticError(f"Newton iteration limit, residual={norm:.3e}")


@dataclass
class MoistureResult:
    times: np.ndarray
    ring_means: np.ndarray
    surface: np.ndarray
    cumulative_exchange: np.ndarray
    balance_residual: np.ndarray
    accepted_half_steps: list = field(default_factory=list)
    rejected_steps: list = field(default_factory=list)
    diagnostics: dict = field(default_factory=dict)


def integrate_moisture(model: MoistureFV, drive: LinearDrive, initial, query_times,
                       *, atol=1e-9, rtol=1e-7, initial_dt=.1, max_dt=1.,
                       min_dt=1e-10, timeout=300., checkpoint: Callable | None=None):


    targets = np.asarray(query_times, float)
    if (targets.ndim != 1 or targets[0] != 0 or np.any(np.diff(targets) <= 0)
            or max_dt <= 0 or initial_dt <= 0 or atol <= 0 or rtol < 0):
        raise ValueError("Invalid integration settings or output times")
    drive(targets)
    c = np.broadcast_to(initial, (model.grid.cells,)).astype(float).copy()
    if np.any(c < 0) or not np.all(np.isfinite(c)):
        raise ValueError("Invalid initial concentration")
    if not np.all(c == c[0]):
        raise ValueError("This integrator requires the approved uniform initial field; "
                         "nonuniform verification needs an explicit initial surface trace")
    state0 = c.copy()
    initial_integral = float(model.grid.volumes@c)
    exchange = 0.
    now, dt = 0., initial_dt
    rings, surfaces, exchanges, balances = [c.copy()], [float(c[-1])], [0.], [0.]
    accepted, rejected = [], []
    start_clock = time.perf_counter()
    max_nl = max_root = max_root_C = 0.
    max_iters = 0
    max_step_balance = 0.
    for target in targets[1:]:
        while now < target:
            if time.perf_counter()-start_clock > timeout:
                if checkpoint:
                    checkpoint(now, c.copy(), exchange, "timeout")
                raise TimeoutError(f"Short-run time budget exhausted at t={now:.9g}")
            endpoint = min(float(target), drive.next_knot(now), now+min(dt, max_dt))
            step = endpoint-now
            if step < min_dt and target-now > 2*min_dt:
                raise ArithmeticError("Adaptive step below minimum")
            mid = now+step/2
            try:
                coarse, sc, dc = model.solve_balance(c, step, drive(endpoint))
                half, sh, dh = model.solve_balance(c, step/2, drive(mid))
                fine, sf, df = model.solve_balance(half, step/2, drive(endpoint))
                scale = atol+rtol*np.maximum(np.abs(coarse), np.abs(fine))
                error = float(np.max(np.abs(fine-coarse)/scale))
                es = abs(sf.concentration-sc.concentration)/(atol+rtol*max(
                    abs(sf.concentration), abs(sc.concentration)))
                error = max(error, es)
                low_half = min(float(c.min()), drive(mid))
                high_half = max(float(c.max()), drive(mid))
                low_fine = min(float(half.min()), drive(endpoint))
                high_fine = max(float(half.max()), drive(endpoint))
                if (fine.min() < low_fine-2e-11 or fine.max() > high_fine+2e-11 or
                        half.min() < low_half-2e-11 or half.max() > high_half+2e-11):
                    raise ArithmeticError("Discrete range principle violated")
            except (ArithmeticError, ValueError, np.linalg.LinAlgError) as exc:
                rejected.append({"t": now, "dt": step, "reason": str(exc)})
                dt = step*.25
                if dt < min_dt:
                    raise ArithmeticError("Unrecoverable implicit step") from exc
                continue
            if not np.isfinite(error) or error > 1.:
                rejected.append({"t": now, "dt": step, "error_norm": error,
                                 "reason": "local step-doubling error"})
                dt = step*max(.2, .9/np.sqrt(max(error, 1.)))
                continue
            for state_before, state_after, surface_state, diagnostic, end in (
                    (c, half, sh, dh, mid), (half, fine, sf, df, endpoint)):
                exchanged = model.grid.areas[-1]*surface_state.outward_flux*step/2
                step_balance = float(model.grid.volumes@(state_after-state_before)+exchanged)
                exchange += exchanged
                max_step_balance = max(max_step_balance, abs(step_balance))
                max_nl = max(max_nl, diagnostic["residual_C"])
                max_iters = max(max_iters, diagnostic["iterations"])
                max_root = max(max_root, abs(surface_state.residual_flux))
                max_root_C = max(max_root_C, surface_state.root_error_C)
                accepted.append([end, step/2, surface_state.concentration,
                                 surface_state.outward_flux, diagnostic["residual_C"],
                                 surface_state.residual_flux, step_balance,
                                 surface_state.bracket_low, surface_state.bracket_high,
                                 surface_state.root_error_C])
            c, now = fine, endpoint
            dt = min(max_dt, step*min(2., max(.5, .9/np.sqrt(max(error, 1e-12)))))
        rings.append(c.copy())
        surfaces.append(sf.concentration)
        exchanges.append(exchange)
        balances.append(float(model.grid.volumes@(c-state0)+exchange))
        if checkpoint:
            checkpoint(now, c.copy(), exchange, "output")
    return MoistureResult(targets, np.array(rings), np.array(surfaces),
                          np.array(exchanges), np.array(balances), accepted, rejected,
                          {"wall_seconds": time.perf_counter()-start_clock,
                           "accepted_full_steps": len(accepted)//2,
                           "accepted_half_steps": len(accepted),
                           "rejected_steps": len(rejected),
                           "maximum_nonlinear_residual_C": max_nl,
                           "maximum_surface_flux_residual_m_s": max_root,
                           "maximum_surface_residual_equivalent_C": max_root_C,
                           "maximum_newton_iterations": max_iters,
                           "maximum_step_balance_m3_C": max_step_balance,
                           "initial_integral_m3_C": initial_integral,
                           "maximum_cumulative_balance_m3_C": float(np.max(np.abs(balances))),
                           "maximum_normalized_balance": (float(np.max(np.abs(balances)))/initial_integral
                                                          if initial_integral > 0 else None),
                           "range_ring_means": [float(np.min(rings)), float(np.max(rings))]})


# 问题二：温湿耦合与变物性扩展。


from dataclasses import dataclass
from pathlib import Path
import hashlib
import json
import math
import time

import numpy as np
from scipy import special
from scipy.optimize import brentq
from scipy.sparse import coo_matrix, bmat, csr_matrix
from scipy.integrate import solve_ivp


GLX, GLW = np.polynomial.legendre.leggauss(12)


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def harmonic(x, y):
    return 2*x*y/(x+y)


def harmonic_partials(x, y):
    return 2*y*y/(x+y)**2, 2*x*x/(x+y)**2


@dataclass(frozen=True)
class Properties:

    test_q1: bool = False

    @property
    def dstar(self):
        return 7e-9 if self.test_q1 else 2.4e-3

    @property
    def gamma(self):
        return .89 if self.test_q1 else .45

    def thermal(self, c):
        c = np.asarray(c, float)
        if self.test_q1:
            return (np.full_like(c, 820*2600.), np.full_like(c, .36),
                    np.zeros_like(c), np.zeros_like(c), np.zeros_like(c))
        rho, cp = 650+128*c, 1450+2736*c/(1+c)
        cpp = 2736/(1+c)**2
        return (rho*cp, .21+.38*c/(1+c), 128*cp+rho*cpp,
                .38/(1+c)**2, 256*cpp-2*rho*2736/(1+c)**3)

    def a(self, t):
        tk = np.asarray(t, float)+273.15
        if np.any(tk <= 0) or not np.all(np.isfinite(tk)):
            raise ValueError("Nonphysical absolute temperature; no clipping")
        if self.test_q1:
            return np.ones_like(tk), np.zeros_like(tk)
        a = np.exp(-3850/tk)
        return a, a*3850/tk**2

    def b(self, c):
        c = np.asarray(c, float)
        if np.any(c < 0) or not np.all(np.isfinite(c)):
            raise ValueError("Nonphysical C; no clipping")
        ans = np.zeros_like(c)
        positive = c > 0
        ans[positive] = np.exp(-self.gamma/c[positive])
        return ans

    def psi(self, c):
        c = np.asarray(c, float)
        if np.any(c < 0) or not np.all(np.isfinite(c)):
            raise ValueError("Nonphysical potential input")
        ans = np.zeros_like(c)
        positive = c > 0
        ans[positive] = c[positive]*special.expn(2, self.gamma/c[positive])
        return ans

    def psi_difference(self, left, right):


        left, right = np.broadcast_arrays(np.asarray(left, float), np.asarray(right, float))
        if np.any(left < 0) or np.any(right < 0) or not np.all(np.isfinite(left+right)):
            raise ValueError("Nonphysical potential difference input")
        delta = left-right
        out = np.zeros_like(delta)
        unequal = delta != 0
        near = unequal & (abs(delta) <= .2*np.minimum(left, right))
        near &= (left > self.gamma/700) & (right > self.gamma/700)
        gentle = np.zeros_like(near)
        gentle[near] = abs(self.gamma/left[near]-self.gamma/right[near]) <= 1
        near &= gentle
        if np.any(near):
            points = (left[near]+right[near])[:, None]/2+delta[near, None]*GLX/2
            out[near] = delta[near]/2*(np.exp(-self.gamma/points)@GLW)
        far = unequal & ~near
        if np.any(far):
            out[far] = self.psi(left[far])-self.psi(right[far])
        return out

    def D(self, c, t):
        return self.dstar*self.a(t)[0]*self.b(c)


@dataclass
class Surface:
    T: float
    C: float
    flux_T: float
    flux_C: float
    derivatives: np.ndarray
    heat_residual_W_m2: float
    moisture_residual_m_s: float
    root_equivalent_C: float
    bracket: tuple
    root_iterations: int


class CoupledFV:
    def __init__(self, cells, *, flux="potential", hT=25., hC=8e-7,
                 properties=None, root_xtol=5e-15):
        if flux not in ("potential", "fick"):
            raise ValueError("Unknown internal moisture flux")
        self.grid = RingGrid(cells)
        self.props = Properties() if properties is None else properties
        self.flux, self.hT, self.hC = flux, float(hT), float(hC)
        self.root_xtol = root_xtol
        self.boundary_calls = 0
        self.max_root_equiv = self.max_heat_residual = self.max_moisture_residual = 0.
        self.max_root_iterations = 0

    def surface(self, ti, ci, ta, ce):


        if ci <= 0 or ce <= 0 or not all(map(math.isfinite, (ti, ci, ta, ce))):
            raise ValueError("Invalid boundary state; no clipping")
        p, delta = self.props, self.grid.delta
        _, ki, _, kip, _ = p.thermal(ci)
        ai, aip = p.a(ti)
        ds = p.dstar
        ht, hc = self.hT, self.hC
        # 在局部区间内检验唯一根条件，约束非线性迭代中的实际试探状态。
        cmin = min(ci, ce)
        _, kmin, _, kpmax, _ = p.thermal(cmin)
        kbmin = harmonic(ki, kmin)
        kbprime_bound = 2*ki**2/(ki+kmin)**2*kpmax
        tsprime_bound = delta*ht*abs(ti-ta)*kbprime_bound/(kbmin+delta*ht)**2
        tkmin = min(ti, ta)+273.15
        if tkmin <= 0:
            raise ValueError("Nonphysical boundary temperature")
        eta = (0. if p.test_q1 else 3850/tkmin**2)*tsprime_bound*abs(ci-ce)
        if eta >= 1:
            raise ArithmeticError(f"Surface branch uniqueness guard failed: eta={eta:.17g}")

        def calc(cs, derivatives=False):
            _, ks, _, ksp, _ = p.thermal(cs)
            kb = harmonic(ki, ks)
            ts = (kb*ti+delta*ht*ta)/(kb+delta*ht)
            av, ap = p.a(ts)
            ab = harmonic(ai, av)
            dp = float(p.psi_difference(cs, ci))
            g = ds*ab*dp/delta+hc*(cs-ce)
            if not derivatives:
                return float(g)
            hki, hks = harmonic_partials(ki, ks)
            hai, has = harmonic_partials(ai, av)
            kbci, kbcs = hki*kip, hks*ksp
            abi, abs_ = hai*aip, has*ap
            # 对两个罗宾边界方程进行二维隐式求导。
            j00 = kb/delta+ht
            j01 = kbcs*(ts-ti)/delta
            j10 = ds*abs_*dp/delta
            j11 = ds*ab*float(p.b(cs))/delta+hc
            x00 = -kb/delta
            x01 = kbci*(ts-ti)/delta
            x10 = ds*abi*dp/delta
            x11 = -ds*ab*float(p.b(ci))/delta
            determinant = j00*j11-j01*j10
            if not math.isfinite(determinant) or determinant <= 0:
                raise ArithmeticError("Surface implicit Jacobian lost the verified positive branch")
            dx = -np.array([[j11*x00-j01*x10, j11*x01-j01*x11],
                            [-j10*x00+j00*x10, -j10*x01+j00*x11]])/determinant
            root_derivative = j11-j10*j01/j00
            hr = kb*(ts-ti)/delta+ht*(ts-ta)
            return float(ts), dx, float(hr), float(g), float(root_derivative)

        if hc == 0 or ci == ce:
            cs, iterations = ci, 0
        else:
            lo, hi = sorted((ci, ce))
            # 将速度量纲残差缩放为含水率量级，不改变方程根。
            scale = hc+float(p.D(ci, ti))/delta
            cs, info = brentq(lambda c: calc(c)/scale, lo, hi,
                             xtol=self.root_xtol, rtol=8.881784197001252e-16,
                             full_output=True, maxiter=80)
            if not info.converged:
                raise ArithmeticError("Bracketed surface solve failed")
            iterations = int(info.iterations)
        ts, dx, hr, cr, gp = calc(cs, True)
        equiv = abs(cr/gp)
        if equiv > 10*self.root_xtol+3e-14*abs(cs):
            raise ArithmeticError("Surface residual exceeds concentration-scaled tolerance")
        self.boundary_calls += 1
        self.max_root_equiv = max(self.max_root_equiv, equiv)
        self.max_heat_residual = max(self.max_heat_residual, abs(hr))
        self.max_moisture_residual = max(self.max_moisture_residual, abs(cr))
        self.max_root_iterations = max(self.max_root_iterations, iterations)
        return Surface(ts, cs, ht*(ts-ta), hc*(cs-ce), dx, hr, cr, equiv,
                       tuple(sorted((ci, ce))), iterations)

    def interior(self, t, c, *, jacobian=False):
        p, tau = self.props, self.grid.tau
        w, k, wp, kp, wpp = p.thermal(c)
        a, ap = p.a(t)
        b = p.b(c)
        kh = harmonic(k[:-1], k[1:])
        ft = tau*kh*(t[:-1]-t[1:])
        if self.flux == "potential":
            ah = harmonic(a[:-1], a[1:])
            diff = p.psi_difference(c[:-1], c[1:])
            fc = tau*p.dstar*ah*diff
        else:
            d = p.dstar*a*b
            dh = harmonic(d[:-1], d[1:])
            diff = c[:-1]-c[1:]
            fc = tau*dh*diff
        if not jacobian:
            return ft, fc
        n = len(c)
        deriv = np.zeros((n-1, 2, 2, 2))
        hx, hy = harmonic_partials(k[:-1], k[1:])
        deriv[:, 0, 0, 0] = tau*kh
        deriv[:, 0, 1, 0] = -tau*kh
        deriv[:, 0, 0, 1] = tau*hx*kp[:-1]*(t[:-1]-t[1:])
        deriv[:, 0, 1, 1] = tau*hy*kp[1:]*(t[:-1]-t[1:])
        if self.flux == "potential":
            hx, hy = harmonic_partials(a[:-1], a[1:])
            deriv[:, 1, 0, 0] = tau*p.dstar*hx*ap[:-1]*diff
            deriv[:, 1, 1, 0] = tau*p.dstar*hy*ap[1:]*diff
            deriv[:, 1, 0, 1] = tau*p.dstar*ah*b[:-1]
            deriv[:, 1, 1, 1] = -tau*p.dstar*ah*b[1:]
        else:
            hx, hy = harmonic_partials(d[:-1], d[1:])
            dc = d*p.gamma/c**2
            dt = p.dstar*ap*b
            deriv[:, 1, 0, 0] = tau*hx*dt[:-1]*diff
            deriv[:, 1, 1, 0] = tau*hy*dt[1:]*diff
            deriv[:, 1, 0, 1] = tau*(hx*dc[:-1]*diff+dh)
            deriv[:, 1, 1, 1] = tau*(hy*dc[1:]*diff-dh)
        return ft, fc, deriv

    def evaluate(self, state, ta, ce, *, jacobian=False, source=None):
        n, p = self.grid.cells, self.props
        yy = np.asarray(state[:2*n]).reshape(n, 2)
        t, c = yy[:, 0], yy[:, 1]
        if not np.all(np.isfinite(yy)) or np.any(c <= 0):
            raise ValueError("Nonphysical trial state; no clipping or reuse of old result")
        w, _, wp, _, wpp = p.thermal(c)
        v, area = self.grid.volumes, self.grid.areas[-1]
        sf = self.surface(float(t[-1]), float(c[-1]), ta, ce)
        inner = self.interior(t, c, jacobian=jacobian)
        ft, fc = np.r_[0., inner[0], area*sf.flux_T], np.r_[0., inner[1], area*sf.flux_C]
        rhs = np.column_stack((-np.diff(ft)/(v*w), -np.diff(fc)/v))
        # 构造解检验时才引入体积平均热源与水分源项。
        if source is not None:
            rhs[:, 0] += source[:, 0]/w
            rhs[:, 1] += source[:, 1]
        if not jacobian:
            return rhs.ravel(), sf
        rows, cols, vals = [], [], []
        faces = np.arange(n-1)
        d = inner[2]
        for field in range(2):
            storage = v*w if field == 0 else v
            for receiver, sign in ((0, -1), (1, 1)):
                for neighbor in range(2):
                    for variable in range(2):
                        rows.append(2*(faces+receiver)+field)
                        cols.append(2*(faces+neighbor)+variable)
                        vals.append(sign*d[:, field, neighbor, variable]/storage[faces+receiver])
        for field, h in ((0, self.hT), (1, self.hC)):
            storage = v[-1]*(w[-1] if field == 0 else 1)
            for variable in range(2):
                rows.append(np.array([2*(n-1)+field]))
                cols.append(np.array([2*(n-1)+variable]))
                vals.append(np.array([-area*h*sf.derivatives[field, variable]/storage]))
        rows.append(2*np.arange(n)); cols.append(2*np.arange(n)+1)
        vals.append(-rhs[:, 0]*wp/w)
        mat = coo_matrix((np.concatenate(vals), (np.concatenate(rows), np.concatenate(cols))),
                         shape=(2*n, 2*n)).tocsc()
        return rhs.ravel(), sf, mat

    def augmented(self, state, ta, ce, *, jacobian=False, source=None):


        n, p, v = self.grid.cells, self.props, self.grid.volumes
        total = float(v.sum())
        w0 = float(p.thermal(2.55)[0])
        scale = total*w0
        result = self.evaluate(state, ta, ce, jacobian=jacobian, source=source)
        dy, sf = result[:2]
        t, c, dc = state[:2*n:2], state[1:2*n:2], dy[1::2]
        _, _, wp, _, wpp = p.thermal(c)
        coeff = v*wp*(t-28)/scale
        extras = np.array([self.grid.areas[-1]*sf.flux_C/total,
                           -self.grid.areas[-1]*sf.flux_T/scale, coeff@dc])
        if not jacobian:
            return np.r_[dy, extras]
        mat = result[2]
        bottom = np.zeros((3, 2*n))
        bottom[0, -2:] = self.grid.areas[-1]*self.hC*sf.derivatives[1]/total
        bottom[1, -2:] = -self.grid.areas[-1]*self.hT*sf.derivatives[0]/scale
        bottom[2] = np.asarray(coeff@mat[1::2]).ravel()
        bottom[2, 0::2] += v*wp*dc/scale
        bottom[2, 1::2] += v*wpp*(t-28)*dc/scale
        return bmat([[mat, None], [csr_matrix(bottom), csr_matrix((3, 3))]], format="csc")


def read_environment(root, config):
    import openpyxl
    spec = config["inputs"]["environment"]
    path = Path(root)/spec["path"]
    if sha256(path) != spec["sha256"]:
        raise ValueError("Environmental source fingerprint changed")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = np.array([r[:3] for r in wb[spec["sheet"]].iter_rows(min_row=2, values_only=True)], float)
    wb.close()
    if rows.shape != (241, 3) or not np.all(np.isfinite(rows)):
        raise ValueError("Missing/non-numeric input; stop rather than fill")
    if not np.array_equal(rows[:, 0], np.arange(241)*60):
        raise ValueError("Original sampling differs from the audited 60 s nodes")
    return LinearDrive(rows[:, 0], rows[:, 1]), LinearDrive(rows[:, 0], rows[:, 2]), rows


def run_short(model, drives, times, radii, *, method="Radau", rtol=1e-9,
              atol_T=1e-10, atol_C=1e-12, max_step=1., first_step_factor=.01,
              checkpoint_dir=None, timeout_per_segment=600., verbose=True):
    times, radii = np.asarray(times, float), np.asarray(radii, float)
    if times[0] != 0 or times[-1] > 121 or np.any(np.diff(times) <= 0):
        raise ValueError("This entry is authorized only for 0--121 s diagnostics")
    td, cd = drives
    n, grid, p = model.grid.cells, model.grid, model.props
    y = np.r_[np.tile([28., 2.55], n), np.zeros(3)]
    weights = grid.point_weights(radii)
    T = np.full((len(times), len(radii)), np.nan); T[0] = 28.
    C = np.full_like(T, np.nan); C[0] = 2.55
    exchange = np.full((len(times), 3), np.nan); exchange[0] = 0
    balance = np.full((len(times), 2), np.nan); balance[0] = 0
    atols = np.r_[np.tile([atol_T, atol_C], n), atol_C, atol_T, atol_T]
    w0 = float(p.thermal(2.55)[0]); volume = float(grid.volumes.sum())
    alpha = float(p.thermal(2.55)[1])/w0
    first = min(max_step, first_step_factor*np.min(np.diff(grid.faces))**2/alpha)
    now, segments = 0., []
    accepted_range = [28., 28., 2.55, 2.55]
    clock = time.perf_counter()
    for end in np.r_[td.times[(td.times > 0)&(td.times < times[-1])], times[-1]]:
        started = time.perf_counter()
        def fun(t, z):
            if time.perf_counter()-started > timeout_per_segment:
                raise TimeoutError(f"Q2 segment deadline at t={t:.17g}")
            return model.augmented(z, td(t), cd(t))
        def jac(t, z):
            return model.augmented(z, td(t), cd(t), jacobian=True)
        sol = solve_ivp(fun, (now, float(end)), y, method=method, jac=jac,
                        rtol=rtol, atol=atols, max_step=max_step,
                        first_step=min(first, float(end)-now), dense_output=True)
        if not sol.success:
            raise ArithmeticError(sol.message)
        select = (times > now)&(times <= end)
        state = sol.sol(times[select]).T
        if not np.all(np.isfinite(state)) or np.any(state[:, 1:2*n:2] <= 0):
            raise ArithmeticError("Nonpositive or nonfinite accepted/dense-output state")
        surfaces = [model.surface(z[-5], z[-4], td(t), cd(t))
                    for t, z in zip(times[select], state)]
        T[select] = state[:, :2*n:2]@weights.T
        C[select] = state[:, 1:2*n:2]@weights.T
        surf_mask = radii == grid.radius
        T[np.ix_(select, surf_mask)] = np.array([s.T for s in surfaces])[:, None]
        C[np.ix_(select, surf_mask)] = np.array([s.C for s in surfaces])[:, None]
        exchange[select] = state[:, -3:]
        balance[select, 0] = (state[:, 1:2*n:2]-2.55)@grid.volumes+volume*state[:, -3]
        w = p.thermal(state[:, 1:2*n:2])[0]
        balance[select, 1] = (w*(state[:, :2*n:2]-28))@grid.volumes-volume*w0*(state[:, -2]+state[:, -1])
        accepted_range[0] = min(accepted_range[0], float(sol.y[:2*n:2].min()))
        accepted_range[1] = max(accepted_range[1], float(sol.y[:2*n:2].max()))
        accepted_range[2] = min(accepted_range[2], float(sol.y[1:2*n:2].min()))
        accepted_range[3] = max(accepted_range[3], float(sol.y[1:2*n:2].max()))
        y, now = sol.y[:, -1], float(end)
        info = dict(end_s=now, accepted_steps=len(sol.t)-1, nfev=sol.nfev,
                    njev=sol.njev, nlu=sol.nlu, wall_s=time.perf_counter()-started,
                    rejected_steps="not exposed by solve_ivp; not reported as zero")
        segments.append(info)
        if checkpoint_dir:
            out = Path(checkpoint_dir); out.mkdir(parents=True, exist_ok=True)
            np.savez_compressed(out/f"checkpoint-{now:g}.npz", time_s=now,
                                state=y, faces_m=grid.faces, config_method=method,
                                code_sha256=sha256(__file__))
        if verbose:
            print(json.dumps({"Q2_progress": info, "cells": n, "flux": model.flux,
                              "method": method}, ensure_ascii=False), flush=True)
    if np.any(~np.isfinite(T)) or np.any(~np.isfinite(C)):
        raise ArithmeticError("Missing diagnostic output")
    stats = dict(cells=n, flux=model.flux, method=method, rtol=rtol, atol_T=atol_T,
                 atol_C=atol_C, max_step_s=max_step, first_step_s=first,
                 first_step_factor=first_step_factor, segments=segments,
                 wall_s=time.perf_counter()-clock, full_final_state_saved=bool(checkpoint_dir),
                 root_calls=model.boundary_calls, root_max_equivalent_C=model.max_root_equiv,
                 root_max_heat_residual_W_m2=model.max_heat_residual,
                 root_max_moisture_residual_m_s=model.max_moisture_residual,
                 root_max_iterations=model.max_root_iterations,
                 maximum_C_balance_m3_C=float(np.max(abs(balance[:, 0]))),
                 maximum_C_balance_normalized=float(np.max(abs(balance[:, 0]))/(volume*2.55)),
                 maximum_heat_chain_balance_J=float(np.max(abs(balance[:, 1]))),
                 maximum_heat_chain_balance_normalized=float(np.max(abs(balance[:, 1]))/(volume*w0)),
                 accepted_ring_range_T_C=accepted_range[:2],
                 accepted_ring_range_C_kg_kg=accepted_range[2:],
                 diagnostic_only=True, original_t0_surface_preserved=True)
    return dict(time_s=times, radius_m=radii, T_C=T, C_kg_kg=C,
                exchanges=exchange, balances=balance, final_state=y, stats=stats)


# 问题三：随温度和含水率变化的物性扩展。


import math
import numpy as np
from scipy.special import expn

class ScalarProperties(Properties):
    def thermal(self,c):
        if np.ndim(c):return super().thermal(c)
        c=float(c)
        if self.test_q1:return 820.*2600,.36,0.,0.,0.
        rho=650+128*c;cp=1450+2736*c/(1+c);cpp=2736/(1+c)**2
        return rho*cp,.21+.38*c/(1+c),128*cp+rho*cpp,.38/(1+c)**2,256*cpp-2*rho*2736/(1+c)**3
    def a(self,t):
        if np.ndim(t):return super().a(t)
        tk=float(t)+273.15
        if tk<=0 or not math.isfinite(tk):raise ValueError('Nonphysical temperature')
        if self.test_q1:return 1.,0.
        a=math.exp(-3850/tk);return a,a*3850/tk**2
    def b(self,c):
        if np.ndim(c):return super().b(c)
        c=float(c)
        if c<0 or not math.isfinite(c):raise ValueError('Nonphysical concentration')
        return math.exp(-self.gamma/c) if c>0 else 0.
    def psi_difference(self,left,right):
        if np.ndim(left) or np.ndim(right):return super().psi_difference(left,right)
        left=float(left);right=float(right)
        if min(left,right)<0 or not math.isfinite(left+right):raise ValueError('Nonphysical potential')
        delta=left-right
        if delta==0:return 0.
        if (abs(delta)<=.2*min(left,right) and min(left,right)>self.gamma/700
                and abs(self.gamma/left-self.gamma/right)<=1):
            mid=(left+right)/2;half=delta/2
            return half*math.fsum(float(w)*math.exp(-self.gamma/(mid+half*float(x))) for x,w in zip(GLX,GLW))
        def psi(c):return c*float(expn(2,self.gamma/c)) if c>0 else 0.
        return psi(left)-psi(right)


# 共用数值监测工具。


import ctypes,gc,time

def current_memory():
    class Counters(ctypes.Structure):
        _fields_=[('cb',ctypes.c_ulong),('PageFaultCount',ctypes.c_ulong)]+[(k,ctypes.c_size_t) for k in
          ['PeakWorkingSetSize','WorkingSetSize','QuotaPeakPagedPoolUsage','QuotaPagedPoolUsage',
           'QuotaPeakNonPagedPoolUsage','QuotaNonPagedPoolUsage','PagefileUsage','PeakPagefileUsage','PrivateUsage']]
    p=Counters();p.cb=ctypes.sizeof(p)
    ctypes.windll.kernel32.GetCurrentProcess.restype=ctypes.c_void_p
    h=ctypes.windll.kernel32.GetCurrentProcess()
    if not ctypes.windll.psapi.GetProcessMemoryInfo(ctypes.c_void_p(h),ctypes.byref(p),p.cb):
        raise OSError('GetProcessMemoryInfo failed')
    return {k:float(getattr(p,k)/2**20) for k in ['WorkingSetSize','PeakWorkingSetSize','PrivateUsage','PeakPagefileUsage']}

def collect_completed_cycles():
    before=current_memory();start=time.perf_counter();objects=gc.collect()
    return dict(before_MB=before,after_MB=current_memory(),collected_objects=objects,wall_s=time.perf_counter()-start)

def probe():
    # 仅诊断已完成分段的不可达循环，不改变求解过程。
    import weakref,json
    from pathlib import Path
    import numpy as np
    from scipy.integrate import BDF
    from threadpoolctl import threadpool_limits
    from q2_model import CoupledFV
    from q3_scalar_properties import ScalarProperties
    from q3_linear import install
    from q3_full_run import environment,ROOT,BASE
    from q2_export import digest,save_json
    td,cd,_=environment();n=4096;m=CoupledFV(n,flux='fick',properties=ScalarProperties())
    initial=np.r_[np.tile([28.,2.55],n),np.zeros(3)];references=[];outputs=[];times=[]
    gc.collect();start=time.perf_counter();before=current_memory();enabled=gc.isenabled();gc.disable()
    try:
        with threadpool_limits(1):
            for _ in range(16):
                s=BDF(lambda t,y:m.augmented(y,td(t),cd(t)),0.,initial.copy(),1.,
                    jac=lambda t,y:m.augmented(y,td(t),cd(t),jacobian=True),
                    rtol=1e-11,atol=np.r_[np.tile([1e-12,1e-14],n),[1e-14,1e-12,1e-12]],first_step=1e-8)
                install(s,2*n,{})
                s.step()
                if s.status=='failed':raise ValueError('Diagnostic first step failed')
                references.append(weakref.ref(s));outputs.append(s.y.copy());times.append(float(s.t));del s
        held=current_memory();live=sum(w() is not None for w in references)
        collected=collect_completed_cycles();remaining=sum(w() is not None for w in references)
    finally:
        if enabled:gc.enable()
    result={'scope':'16 independent one-step probes, not a full PDE trajectory. GC was disabled only inside this isolated diagnostic process.',
            'cells':n,'before_MB':before,'before_collection_MB':held,'live_unreachable_solvers_before':live,
            'remaining_solvers_after':remaining,'collection':collected,
            'equal_initial_step_output_max_gap':float(np.max(abs(np.array(outputs)-outputs[0]))),
            'equal_initial_step_time_max_gap':float(np.ptp(times)),
            'wall_s':time.perf_counter()-start,'source_sha256':digest(Path(__file__)),
            'limitation':'Shows collectable cycles and their memory in this probe; does not prove the entire cause or exact peak of the earlier failed process.'}
    save_json(BASE/'memory-cycle-probe.json',result);print(json.dumps(result))


from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
import numpy as np


@dataclass(frozen=True)
class PolynomialProfile:
    edges_x: np.ndarray
    origin_x: np.ndarray
    scale_x: np.ndarray
    coefficients: np.ndarray
    centre_C: float
    surface_C: float

    @property
    def radius_m(self):
        return float(np.sqrt(self.edges_x[-1]))


class RingMaximumPlan:


    def __init__(self, faces_m):
        faces = np.asarray(faces_m, dtype=float)
        if (faces.ndim != 1 or len(faces) < 4 or faces[0] != 0
                or not np.all(np.isfinite(faces)) or np.any(np.diff(faces) <= 0)):
            raise ValueError("Invalid radial faces")
        width = np.diff(faces)
        if not np.allclose(width, faces[-1]/len(width), rtol=2e-13,
                           atol=16*np.finfo(float).eps*faces[-1]):
            raise ValueError("Frozen two-ring centre rule requires a uniform grid")
        self.faces = faces.copy()
        self.n = len(width)
        self.edges = faces**2
        cell = np.arange(self.n)
        start = np.clip(cell-1, 0, self.n-3)
        self.ids = start[:, None]+np.arange(3)
        self.origin = (self.edges[:-1]+self.edges[1:])/2
        self.scale = self.edges[start+3]-self.edges[start]
        lo = (self.edges[self.ids]-self.origin[:, None])/self.scale[:, None]
        hi = (self.edges[self.ids+1]-self.origin[:, None])/self.scale[:, None]
        moments = np.stack((np.ones_like(lo), (lo+hi)/2,
                            (lo*lo+lo*hi+hi*hi)/3), axis=2)
        self.inverse = np.linalg.inv(moments)
        self.lo_s = (self.edges[:-1]-self.origin)/self.scale
        self.hi_s = (self.edges[1:]-self.origin)/self.scale
        # 给出环带平均值误差向重构场传播的上界，不代表有限体积离散误差。
        weight_polys = np.swapaxes(self.inverse, 1, 2).reshape(-1, 3)
        lo_s, hi_s = np.repeat(self.lo_s, 3), np.repeat(self.hi_s, 3)
        self.weight_abs_sup = _quadratic_abs_max(weight_polys, lo_s, hi_s).reshape(self.n, 3)

    def profile(self, averages_C, surface_C):
        averages = np.asarray(averages_C, dtype=float)
        if averages.shape != (self.n,) or not np.all(np.isfinite(averages)):
            raise ValueError("Need all finite ring means on this grid")
        if not np.isfinite(surface_C):
            raise ValueError("Need the independently solved finite surface point")
        coefficients = np.einsum("nij,nj->ni", self.inverse, averages[self.ids])
        return PolynomialProfile(
            self.edges.copy(), self.origin, self.scale, coefficients,
            float(1.25*averages[0]-.25*averages[1]), float(surface_C))

    def propagate_supplied_errors(self, ring_abs_errors_C, surface_abs_error_C):
        error = np.asarray(ring_abs_errors_C, dtype=float)
        if (error.shape != (self.n,) or not np.all(np.isfinite(error))
                or np.any(error < 0) or not np.isfinite(surface_abs_error_C)
                or surface_abs_error_C < 0):
            raise ValueError("Need finite nonnegative error bounds/estimates")
        cell = np.sum(self.weight_abs_sup*error[self.ids], axis=1)
        centre = 1.25*error[0]+.25*error[1]
        return dict(
            maximum_C=float(max(np.max(cell), centre, surface_abs_error_C)),
            cell_C=cell, centre_C=float(centre), surface_C=float(surface_abs_error_C),
            scope="Propagation through reconstruction only; inherits the status of supplied errors")


def _values(coef, s):
    return (coef[:, 2]*s+coef[:, 1])*s+coef[:, 0]


def _stationary(coef, lo, hi):
    roots = np.zeros(len(coef))
    np.divide(-coef[:, 1], 2*coef[:, 2], out=roots, where=coef[:, 2] != 0)
    valid = (coef[:, 2] != 0) & np.isfinite(roots) & (roots > lo) & (roots < hi)
    return roots, valid


def _quadratic_abs_max(coef, lo, hi):
    result = np.maximum(np.abs(_values(coef, lo)), np.abs(_values(coef, hi)))
    roots, valid = _stationary(coef, lo, hi)
    if np.any(valid):
        result[valid] = np.maximum(result[valid], np.abs(_values(coef[valid], roots[valid])))
    return result


def evaluate_profile(profile, radii_m):
    radius = np.asarray(radii_m, dtype=float)
    if not np.all(np.isfinite(radius)) or np.any(radius < 0) or np.any(radius > profile.radius_m):
        raise ValueError("Radius outside domain")
    x = radius**2
    ids = np.clip(np.searchsorted(profile.edges_x, x, side="right")-1,
                  0, len(profile.coefficients)-1)
    s = (x-profile.origin_x[ids])/profile.scale_x[ids]
    values = _values(profile.coefficients[ids], s)
    values = np.where(radius == 0, profile.centre_C, values)
    return np.where(radius == profile.radius_m, profile.surface_C, values)


def scan_profile(profile):


    coef = profile.coefficients
    lo = (profile.edges_x[:-1]-profile.origin_x)/profile.scale_x
    hi = (profile.edges_x[1:]-profile.origin_x)/profile.scale_x
    vl, vr = _values(coef, lo), _values(coef, hi)
    roots, valid = _stationary(coef, lo, hi)
    root_values = _values(coef[valid], roots[valid])
    root_x = profile.origin_x[valid]+profile.scale_x[valid]*roots[valid]
    values = np.r_[profile.centre_C, profile.surface_C, vl, vr, root_values]
    positions = np.sqrt(np.r_[0., profile.edges_x[-1],
                             profile.edges_x[:-1], profile.edges_x[1:], root_x])
    kinds = np.r_[np.array(["centre_point", "surface_point"]),
                  np.repeat("left_piece_limit", len(coef)),
                  np.repeat("right_piece_limit", len(coef)),
                  np.repeat("interior_stationary", np.count_nonzero(valid))]
    imax, imin = int(np.argmax(values)), int(np.argmin(values))
    seam = vr[:-1]-vl[1:]
    tie_tolerance = 32*np.finfo(float).eps*max(1., float(np.max(np.abs(values))))
    ties = np.flatnonzero(values >= values[imax]-tie_tolerance)
    return dict(
        maximum_C=float(values[imax]), argmax_radius_m=float(positions[imax]),
        argmax_kind=str(kinds[imax]),
        minimum_C=float(values[imin]), argmin_radius_m=float(positions[imin]),
        argmin_kind=str(kinds[imin]),
        centre_C=float(profile.centre_C), surface_C=float(profile.surface_C),
        centre_right_limit_C=float(vl[0]), surface_left_limit_C=float(vr[-1]),
        centre_seam_C=float(vl[0]-profile.centre_C),
        surface_seam_C=float(vr[-1]-profile.surface_C),
        max_internal_seam_abs_C=float(np.max(np.abs(seam))) if len(seam) else 0.,
        stationary_candidates=int(np.count_nonzero(valid)),
        evaluated_candidates=int(len(values)),
        near_tied_candidates=[dict(value_C=float(values[i]), radius_m=float(positions[i]),
                                   kind=str(kinds[i])) for i in ties[:12]],
        scope="Full reconstructed-radius supremum/minimum, including one-sided limits; no PDE error bound")


def combine_profiles(terms):


    if not terms:
        raise ValueError("Empty combination")
    if any(not np.isfinite(weight) for weight, _ in terms):
        raise ValueError("Nonfinite combination weight")
    end = terms[0][1].edges_x[-1]
    if any(profile.edges_x[0] != 0 or profile.edges_x[-1] != end for _, profile in terms):
        raise ValueError("The physical domains must agree exactly")
    edges = np.unique(np.concatenate([profile.edges_x for _, profile in terms]))
    origin = (edges[:-1]+edges[1:])/2
    scale = np.diff(edges)
    result = np.zeros((len(origin), 3))
    for weight, profile in terms:
        ids = np.clip(np.searchsorted(profile.edges_x, origin, side="right")-1,
                      0, len(profile.coefficients)-1)
        coef = profile.coefficients[ids]
        a = (origin-profile.origin_x[ids])/profile.scale_x[ids]
        b = scale/profile.scale_x[ids]
        result[:, 0] += weight*_values(coef, a)
        result[:, 1] += weight*b*(coef[:, 1]+2*coef[:, 2]*a)
        result[:, 2] += weight*b*b*coef[:, 2]
    return PolynomialProfile(
        edges, origin, scale, result,
        float(sum(weight*profile.centre_C for weight, profile in terms)),
        float(sum(weight*profile.surface_C for weight, profile in terms)))


def profile_difference_sup(first, second):
    result = scan_profile(combine_profiles([(1., first), (-1., second)]))
    if abs(result["minimum_C"]) > abs(result["maximum_C"]):
        error, radius, kind = abs(result["minimum_C"]), result["argmin_radius_m"], result["argmin_kind"]
    else:
        error, radius, kind = abs(result["maximum_C"]), result["argmax_radius_m"], result["argmax_kind"]
    return dict(max_abs_difference_C=float(error), radius_m=float(radius), kind=kind,
                scope="Supremum of two reconstructed fields on merged partitions, not true PDE error")


def threshold_side(scan, supplied_error_C, *, threshold_C=.15, error_status):


    if (not np.isfinite(supplied_error_C) or supplied_error_C < 0
            or not np.isfinite(threshold_C) or not error_status):
        raise ValueError("Supply a nonnegative budget and state whether it is empirical or rigorous")
    maximum = float(scan["maximum_C"])
    if not np.isfinite(maximum):
        raise ValueError("Nonfinite reconstructed maximum")
    lower = float(np.nextafter(maximum-supplied_error_C, -np.inf))
    upper = float(np.nextafter(maximum+supplied_error_C, np.inf))
    if upper < threshold_C:
        side = "strictly_below_within_supplied_budget"
    elif lower >= threshold_C:
        side = "not_strictly_below_within_supplied_budget"
    else:
        side = "unresolved"
    return dict(maximum_C=maximum, lower_C=lower, upper_C=upper,
                threshold_C=float(threshold_C), strict_margin_C=float(threshold_C-upper),
                side=side, error_status=str(error_status),
                supplied_error_C=float(supplied_error_C))


def event_bracket_report(left_time_s, left_side, right_time_s, right_side,
                         *, single_transition_evidence, hour_digits=4):

    if not (np.isfinite(left_time_s) and np.isfinite(right_time_s)
            and 0 <= left_time_s < right_time_s):
        raise ValueError("Need increasing nonnegative finite times")
    if left_side["side"] != "not_strictly_below_within_supplied_budget":
        raise ValueError("Left endpoint is not established on the not-dry side")
    if right_side["side"] != "strictly_below_within_supplied_budget":
        raise ValueError("Right endpoint lacks a strictly positive dry-side margin")
    quantum = Decimal(1).scaleb(-hour_digits)
    def rounded_h(t):
        return str((Decimal.from_float(float(t))/Decimal(3600)).quantize(
            quantum, rounding=ROUND_HALF_UP))
    lo, hi = rounded_h(left_time_s), rounded_h(right_time_s)
    return dict(
        crossing_interval_s=[float(left_time_s), float(right_time_s)],
        width_s=float(right_time_s-left_time_s),
        verified_strict_side_time_s=float(right_time_s),
        strict_side_margin_C=float(right_side["strict_margin_C"]),
        hour_rounding_at_left=lo, hour_rounding_at_right=hi,
        crossing_hour_rounding_stable=bool(lo == hi),
        single_transition_evidence=str(single_transition_evidence),
        scope=("First crossing bracket only conditional on supplied error budgets and the stated "
               "single-transition/no-earlier-crossing evidence; right endpoint is not the exact infimum"))
