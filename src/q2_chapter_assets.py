"""Generate Q2 chapter tables and traceable numeric macros; never solve a PDE."""
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import csv
import hashlib
import json
import sys

import numpy as np
import openpyxl
from q2_export import verify_delivery

ROOT = Path(__file__).resolve().parents[1]
REC = ROOT / "records/q2-chapter-20260911"
OUT = ROOT / "paper/generated"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def fixed(value, digits=4):
    return str(Decimal.from_float(float(value)).quantize(
        Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP))


def scientific(value, digits=4):
    a, b = f"{float(value):.{digits}e}".split("e")
    return a + r"\times10^{" + str(int(b)) + "}"


def save(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False),
                    encoding="utf-8", newline="\n")


def main():
    REC.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(exist_ok=True)
    authority = ROOT / "results/q2/q2-authority.npz"
    workbook = ROOT / "results/q2/result2.xlsx"
    acceptance = ROOT / "records/q2-full-20260911/numerical-acceptance.json"
    config = ROOT / "config/q2-accepted.json"
    # Reuse the existing read-only delivery verifier, not the numerical solver.
    verified = verify_delivery(authority, acceptance, config, ROOT / "results/q2")
    save(REC / "current-numerical-readback.json", verified)
    inputs = [
        authority, workbook, acceptance, config, ROOT / "config/q2.json",
        ROOT / "src/q2_model.py", ROOT / "src/q2_full_run.py",
        ROOT / "src/q2_grid_extrapolation.py", ROOT / "src/q1_model.py",
        ROOT / "A题/A题.pdf", ROOT / "A题/附件/附件1.xlsx",
        ROOT / "A题/附件/附件3/result2.xlsx",
        ROOT / "records/q2-full-20260911/analysis/final-r16-v2/analysis.json",
        ROOT / "records/q2-full-20260911/run-audits/primary-8192.json",
        ROOT / "records/q2-full-20260911/analysis/final-r16-v2/table60-evidence.csv",
        ROOT / "results/q2/table3.csv", ROOT / "results/q2/table4.csv",
        ROOT / "results/q2/table3.tex", ROOT / "results/q2/table4.tex",
    ]
    # Preserve the entire existing Q1 paper and all accepted Q2 bound evidence.
    inputs.extend(p for p in (ROOT / "paper").rglob("*")
                  if p.is_file() and "build" not in p.parts
                  and ("q1" in p.name.lower() or p.name in ("references.bib", "main.tex")))
    before = {str(p.relative_to(ROOT)): sha(p) for p in inputs}
    with np.load(authority, allow_pickle=False) as z:
        times, radii = z["time_s"].copy(), z["radius_m"].copy()
        fields = [z["T_C"].copy(), z["C_kg_kg"].copy()]
    assert np.array_equal(times, np.arange(10801))
    assert np.array_equal(radii, np.arange(21) * .001)
    assert np.all(fields[0][0] == 28) and np.all(fields[1][0] == 2.55)
    selected_times = [1800, 3600, 5400, 7200, 9000, 10800]
    selected_columns = [0, 5, 10, 15, 20]
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    checked = []
    for k, (sheet, filename, caption, label) in enumerate([
        ("温度", "q2-table-temperature.tex", "3小时内药材的温度（单位：℃）", "tab:q2-temperature"),
        ("水分浓度", "q2-table-moisture.tex", "3小时内药材的水分浓度（干基，单位：kg/kg）", "tab:q2-moisture"),
    ]):
        with (ROOT / f"results/q2/table{k+3}.csv").open(encoding="utf-8-sig", newline="") as f:
            old = list(csv.reader(f))
        # Stream once; retain only the six source workbook rows.
        excel = {i: row for i, row in enumerate(wb[sheet].iter_rows(values_only=True))
                 if i in selected_times}
        lines = [
            "% Generated from frozen q2-authority.npz; no manual numeric edits.",
            "% SHA-256: " + sha(authority),
            r"\begin{table}[!htbp]", r"\centering", r"\small",
            r"\caption{" + caption + "}", r"\label{" + label + "}",
            r"\setlength{\tabcolsep}{9pt}", r"\renewcommand{\arraystretch}{1.12}",
            r"\begin{tabular*}{0.86\textwidth}{@{\extracolsep{\fill}}rccccc}",
            r"\toprule", r"时间/h & \multicolumn{5}{c}{到药材中心的距离/cm} \\",
            r"\cmidrule(lr){2-6}", r" & 0 & 0.5 & 1 & 1.5 & 2 \\", r"\midrule",
        ]
        for row_index, t in enumerate(selected_times, 1):
            row = [fixed(fields[k][t, j]) for j in selected_columns]
            assert Decimal(old[row_index][0]) == Decimal(str(t / 3600))
            for col_index, (j, value) in enumerate(zip(selected_columns, row), 1):
                assert Decimal(str(excel[t][j+1])) == Decimal(value)
                assert Decimal(old[row_index][col_index]) == Decimal(value)
                checked.append(dict(table=k+3, field=sheet, time_s=t,
                                    radius_cm=float(radii[j]*100), full_precision=float(fields[k][t,j]),
                                    printed=value))
            lines.append(f"{t/3600:.1f}" + " & " + " & ".join(row) + r" \\")
        lines.extend([r"\bottomrule", r"\end{tabular*}", r"\end{table}"])
        (OUT / filename).write_text("\n".join(lines)+"\n", encoding="utf-8", newline="\n")
    wb.close()

    ew = openpyxl.load_workbook(ROOT / "A题/附件/附件1.xlsx", read_only=True, data_only=True)
    environment = np.array([r[:3] for r in ew["Sheet1"].iter_rows(min_row=2, values_only=True)], float)
    ew.close()
    assert environment.shape == (241,3) and np.array_equal(environment[:,0], np.arange(241)*60)
    assert np.all(np.isfinite(environment))
    T, C = fields
    analysis = json.loads(inputs[12].read_text(encoding="utf-8"))
    audit = json.loads(inputs[13].read_text(encoding="utf-8"))
    accepted = json.loads(config.read_text(encoding="utf-8"))
    for key in ["T_C", "C_kg_kg"]:
        assert accepted["numerical_acceptance"]["error_evidence"][key] == analysis["error_evidence"][key]
    nums, derivations = {}, {}

    def add(name, value, source, formula, style="fixed", digits=4):
        rendered = scientific(value, digits) if style == "scientific" else fixed(value, digits)
        nums[name] = rendered
        derivations[name] = dict(value=float(value), rendered=rendered,
                                 source=source, formula=formula)

    for t, suffix in [(1800,"Half"), (10800,"End")]:
        for field, arr in [("T",T),("C",C)]:
            for col, place in [(0,"Centre"),(20,"Surface")]:
                add(f"QTwo{field}{place}{suffix}", arr[t,col], "results/q2/q2-authority.npz",
                    f"{field}[time_s={t}, radius_m={radii[col]}]")
        add("QTwoTGap"+suffix, T[t,-1]-T[t,0], "results/q2/q2-authority.npz",
            f"T_C[{t},20]-T_C[{t},0]")
    add("QTwoCCentreHalfFine", C[1800,0], "results/q2/q2-authority.npz",
        "C_kg_kg[1800,0]", digits=10)
    add("QTwoCCentreHalfDrop", 2.55-C[1800,0], "results/q2/q2-authority.npz",
        "2.55-C_kg_kg[1800,0]", style="scientific")
    for name, t1, t2 in [("QTwoCentreGainEarly",1800,3600), ("QTwoCentreGainLate",9000,10800)]:
        add(name, T[t2,0]-T[t1,0], "results/q2/q2-authority.npz", f"T_C[{t2},0]-T_C[{t1},0]")
    add("QTwoTSurfaceReverse", T[10500,-1], "results/q2/q2-authority.npz", "T_C[10500,20]")
    add("QTwoTaReverse", environment[175,1], "A题/附件/附件1.xlsx", "Sheet1 time_s=10500, temperature")
    add("QTwoTaEnd", environment[180,1], "A题/附件/附件1.xlsx", "Sheet1 time_s=10800, temperature")
    add("QTwoSurfaceEnvironmentGap", T[10500,-1]-environment[175,1],
        "authority T_C[10500,20] and attachment1 time=10500", "T_s-T_a")
    # Evaluate the given constitutive law on the accepted physical point outputs.
    D0 = .0024*np.exp(-.45/2.55)*np.exp(-3850/(28+273.15))
    add("QTwoDInitial", D0, "题目附录3及题定初态", "0.0024*exp(-0.45/2.55)*exp(-3850/301.15)", "scientific")
    for col, place in [(0,"Centre"),(20,"Surface")]:
        Ic = .45*(1/2.55-1/C[-1,col])
        It = 3850*(1/301.15-1/(T[-1,col]+273.15))
        direct = .0024*np.exp(-.45/C[-1,col])*np.exp(-3850/(T[-1,col]+273.15))
        assert np.isclose(D0*np.exp(Ic+It), direct, rtol=2e-14, atol=0)
        for tag, val, formula in [
            ("LogC", Ic, "0.45*(1/C0-1/C)"),
            ("LogT", It, "3850*(1/TK0-1/TK)"),
            ("FactorC", np.exp(Ic), "exp(I_C)"),
            ("FactorT", np.exp(It), "exp(I_T)"),
            ("RatioD", direct/D0, "D(C,TK)/D(C0,TK0)"),
        ]:
            add("QTwo"+place+tag, val, f"authority at time_s=10800, radius_m={radii[col]}, appendix3",
                formula)
        add("QTwoD"+place+"End", direct, "authority at time_s=10800 and appendix3",
            f"0.0024*exp(-0.45/C[-1,{col}])*exp(-3850/(T[-1,{col}]+273.15))", "scientific")
    for field, letter in [("T_C","T"),("C_kg_kg","C")]:
        data = analysis["error_evidence"][field]
        for tag, key in [("Independent","independent_difference_max"), ("Envelope","engineering_envelope_max"),
                         ("IntegerEnvelope","integer_envelope_max")]:
            add("QTwo"+letter+tag, data[key]["value"],
                "records/q2-full-20260911/analysis/final-r16-v2/analysis.json",
                f"error_evidence.{field}.{key}.value", "scientific")
        c = analysis["raw_main_convergence"][field]
        for tag, key in [("CoarseDifference","coarse_to_middle"),("FineDifference","middle_to_fine")]:
            add("QTwo"+letter+tag, c[key]["value"], "final-r16-v2/analysis.json",
                f"raw_main_convergence.{field}.{key}.value", "scientific")
        add("QTwo"+letter+"Order", c["norm_order"], "final-r16-v2/analysis.json",
            f"raw_main_convergence.{field}.norm_order", digits=5)
    for name, value, formula in [
        ("QTwoMoistureLedger", audit["independent_checkpoint_audit"]["C_integral_m3_C"], "independent_checkpoint_audit.C_integral_m3_C"),
        ("QTwoHeatLedger", audit["independent_checkpoint_audit"]["heat_chain_J"], "independent_checkpoint_audit.heat_chain_J"),
        ("QTwoMoistureGauss", audit["actual_run_stats"]["Gauss_C_gap_m3_C"], "actual_run_stats.Gauss_C_gap_m3_C"),
        ("QTwoHeatGauss", audit["actual_run_stats"]["Gauss_heat_chain_gap_J"], "actual_run_stats.Gauss_heat_chain_gap_J"),
    ]:
        add(name, value, "records/q2-full-20260911/run-audits/primary-8192.json", formula, "scientific")
    (OUT/"q2-numbers.tex").write_text(
        "% Derived from frozen Q2 data. See chapter assets.json for each numeric source/formula.\n" +
        "\n".join("\\newcommand{\\"+key+"}{"+value+"}" for key,value in nums.items())+"\n",
        encoding="utf-8", newline="\n")
    for rel, expected in before.items():
        assert sha(ROOT/rel) == expected, f"Frozen input changed: {rel}"
    save(REC/"assets.json", dict(
        status="Q2_CHAPTER_ASSETS_PASS", created_utc=datetime.now(timezone.utc).isoformat(),
        command=[sys.executable,*sys.argv], source_sha256=sha(__file__), PDE_solved=False,
        inputs_sha256=before, preserved_evidence_sha256=verified["gate"]["current_evidence_sha256"],
        numeric_macros=derivations, checked_paper_values=checked,
        source_excel_values_checked=verified["workbook"]["checked_numeric_results"],
        output_files={str(p.relative_to(ROOT)):sha(p) for p in OUT.glob("q2-*.tex")},
        input_and_accepted_results_unchanged=True, human_review_performed=False))
    print(json.dumps(dict(status="Q2_CHAPTER_ASSETS_PASS", paper_cells=len(checked),
                         source_excel_cells=453600, numeric_macros=len(nums), PDE_solved=False),
                     ensure_ascii=False))


if __name__ == "__main__":
    main()
