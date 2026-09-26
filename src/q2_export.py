"""Q2 authority -> untouched-template copy, tables 3/4, exhaustive readback.

This module does not call a PDE solver or claim scientific accuracy on its own.
artifact-tool authors the workbook. The verified Q1-style adapter is incorporated
here: restore original template styles without changing any authored value.
Fixture mode is confined to records and can never declare a formal result.
"""
from __future__ import annotations

import argparse
import csv
from copy import copy
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = Path(__file__).resolve()
PREP = ROOT / "records/q2-full-20260911/export-preparation"
TEMPLATE = ROOT / "A题/附件/附件3/result2.xlsx"
TEMPLATE_SHA = "23b261b295c1b787d000eebbca6521c37075107b6fcf78724f8d395ce1798ff4"
NAMES = ("温度", "水分浓度")
A1 = "时间\\到药材中心的距离"
END_S = 10800
TABLE_TIMES = (1800, 3600, 5400, 7200, 9000, 10800)
TABLE_RADIAL_INDEX = (0, 5, 10, 15, 20)
STYLE_KEYS = ("font", "fill", "border", "alignment", "protection")
NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
RAW_INPUTS = (ROOT / "A题/A题.pdf", ROOT / "A题/附件/附件1.xlsx")
RAW_TEMPLATES = tuple(ROOT / f"A题/附件/附件3/result{i}.xlsx" for i in range(1, 5))


def digest(path: Path) -> str:
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def save_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def project_path(value: str | Path) -> Path:
    path = Path(value)
    path = path.resolve() if path.is_absolute() else (ROOT / path).resolve()
    if not path.is_relative_to(ROOT):
        raise ValueError(f"Path leaves this project: {path}")
    return path


def check_python() -> None:
    config = json.loads((ROOT / "project_config.json").read_text(encoding="utf-8"))
    expected = Path(os.path.expandvars(config["python"])).resolve()
    if Path(sys.executable).resolve() != expected:
        raise ValueError(f"Use project Python: {expected}")


def round_half_up(values: np.ndarray) -> np.ndarray:
    """Identical Q1 rule: round the exact binary64 value once, at export only."""
    array = np.asarray(values, dtype=np.float64)
    if not np.all(np.isfinite(array)) or np.any(array < 0):
        raise ValueError("This Q2 export requires finite nonnegative physical values")
    quantum = Decimal("0.0001")
    return np.fromiter((float(Decimal.from_float(float(v)).quantize(quantum, rounding=ROUND_HALF_UP))
                        for v in array.flat), dtype=np.float64, count=array.size).reshape(array.shape)


def load_authority(path: Path, *, fixture: bool = False) -> tuple[np.ndarray, np.ndarray, list[np.ndarray]]:
    with np.load(path, allow_pickle=False) as data:
        required = {"time_s", "radius_m", "T_C", "C_kg_kg"}
        if not required.issubset(data.files):
            raise ValueError(f"Authority missing keys: {required - set(data.files)}")
        if "fixture_only" in data.files and (data["fixture_only"].shape != () or data["fixture_only"].dtype != np.bool_):
            raise ValueError("NPZ fixture_only, if present, must be a scalar boolean")
        tagged_fixture = "fixture_only" in data.files and bool(data["fixture_only"].item())
        if tagged_fixture != fixture:
            raise ValueError("NPZ fixture tag and explicit --fixture mode must agree")
        t = np.array(data["time_s"], dtype=np.float64)
        r = np.array(data["radius_m"], dtype=np.float64)
        fields = [np.array(data[k], dtype=np.float64) for k in ("T_C", "C_kg_kg")]
    if not np.array_equal(t, np.arange(END_S + 1, dtype=float)):
        raise ValueError("Authority requires exactly t=0..10800 s, including the uniform initial state")
    if r.shape != (21,) or not np.allclose(r, np.arange(21) * .001, rtol=0, atol=1e-15):
        raise ValueError("Authority radii must be 0..0.020 m with spacing 0.001 m")
    for field, initial in zip(fields, (28., 2.55)):
        if field.shape != (END_S + 1, 21) or not np.all(np.isfinite(field)) or np.any(field < 0):
            raise ValueError("Authority fields must be finite nonnegative (10801,21) arrays")
        if not np.all(field[0] == initial):
            raise ValueError("The exact original uniform t=0 state must be preserved")
    return t, r, fields


def check_numerical_gate(acceptance: dict, authority: Path, config: Path,
                         *, candidate: bool = False, fixture: bool = False) -> dict:
    """Verify a current numerical verdict and every supplied evidence binding.

    Mandatory coverage ties the verdict to this exporter, Q2 model, selected run
    config, project runtime config and raw inputs. Additional numerical scripts,
    arrays and reports must be listed by the numerical reviewer in evidence.
    Candidate/fixture modes never bypass stale-evidence rejection.
    """
    if type(candidate) is not bool or type(fixture) is not bool:
        raise ValueError("candidate and fixture must be booleans")
    authority, config = project_path(authority), project_path(config)
    if acceptance.get("question") != "Q2":
        raise ValueError("Acceptance question must be Q2")
    if acceptance.get("authority_sha256") != digest(authority):
        raise ValueError("Authority SHA differs from numerical acceptance")
    fixture_tag = acceptance.get("fixture_only", False)
    if type(fixture_tag) is not bool or fixture_tag != fixture:
        raise ValueError("Fixture acceptance can never authorize a real/formal export")
    count = acceptance.get("unresolved_rounding_count")
    if type(count) is not int or count < 0:
        raise ValueError("unresolved_rounding_count must be an explicit nonnegative integer")
    if type(acceptance.get("numerical_pass")) is not bool:
        raise ValueError("numerical_pass must be an explicit boolean")
    passed = acceptance["numerical_pass"]
    if not (candidate or fixture) and (not passed or count != 0):
        raise ValueError("Formal output requires numerical_pass=true and unresolved_rounding_count=0")
    if fixture and not authority.is_relative_to(PREP):
        raise ValueError("Fixture authority must stay inside export-preparation records")
    evidence = acceptance.get("evidence_sha256")
    if not isinstance(evidence, dict) or not evidence:
        raise ValueError("A nonempty evidence_sha256 absolute-project-path mapping is required")
    checked = {}

    def verify_file(value, expected, label):
        if not isinstance(value, str) or not Path(value).is_absolute():
            raise ValueError(f"{label} requires an absolute project path")
        path = project_path(value)
        if not isinstance(expected, str) or not re.fullmatch(r"[0-9a-fA-F]{64}", expected):
            raise ValueError(f"Invalid {label} SHA-256")
        if not path.is_file() or digest(path) != expected.lower():
            raise ValueError(f"{label} is missing or changed: {path}")
        return path

    for value, sha in evidence.items():
        path = verify_file(value, sha, "Evidence")
        if str(path) in checked:
            raise ValueError(f"Duplicate normalized evidence path: {path}")
        checked[str(path)] = sha.lower()
    required = (SCRIPT, ROOT / "src/q2_model.py", config, ROOT / "project_config.json", TEMPLATE, *RAW_INPUTS)
    missing = [str(p.resolve()) for p in required if str(p.resolve()) not in checked]
    if missing:
        raise ValueError(f"Evidence omits required current code/config/input files: {missing}")
    analysis = verify_file(acceptance.get("analysis_path"), acceptance.get("analysis_sha256"), "Analysis")
    if str(analysis) not in checked or checked[str(analysis)] != acceptance["analysis_sha256"].lower():
        raise ValueError("Analysis must also appear in evidence_sha256")
    if digest(TEMPLATE) != TEMPLATE_SHA:
        raise ValueError("The original result2.xlsx template version changed")
    return {"question": "Q2", "authority_sha256": digest(authority),
            "numerical_pass": passed, "unresolved_rounding_count": count,
            "candidate_requested": candidate, "fixture_only": fixture,
            "current_evidence_sha256": checked, "analysis_path": str(analysis),
            "analysis_sha256": digest(analysis), "checked_evidence_files": len(checked)}


def template_metadata() -> dict:
    if digest(TEMPLATE) != TEMPLATE_SHA:
        raise ValueError("Original template SHA changed")
    book = openpyxl.load_workbook(TEMPLATE)
    try:
        if book.sheetnames != list(NAMES):
            raise ValueError("Template sheet names/order changed")
        info = {}
        for sheet in book:
            if sheet["A1"].value != A1 or [sheet.cell(i, 1).value for i in (2, 3, 4)] != [1, 2, 3]:
                raise ValueError("Original schematic axes changed")
            if sheet.calculate_dimension() != "A1:F5" or list(sheet.merged_cells.ranges):
                raise ValueError("Unexpected template native layout")
            info[sheet.title] = {"dimension": sheet.calculate_dimension(), "A1": A1,
                "a_width": sheet.column_dimensions["A"].width,
                "row1_height": sheet.row_dimensions[1].height,
                "freeze_panes": sheet.freeze_panes, "merges": [],
                "rows": [list(row) for row in sheet.iter_rows(values_only=True)]}
        return {"path": str(TEMPLATE), "sha256": TEMPLATE_SHA, "sheets": info,
                "endpoint_source": "User-approved 10800 s; original template is schematic"}
    finally:
        book.close()


def _xml_sheet_paths(archive: ZipFile) -> list[tuple[str, str]]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    links = {r.attrib["Id"]: r.attrib["Target"] for r in relations}
    output = []
    for sheet in workbook.findall("x:sheets/x:sheet", NS):
        rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
        target = links[rid]
        member = target.lstrip("/") if target.startswith("/") else "xl/" + target
        if ".." in member.split("/") or member not in archive.namelist():
            raise ValueError("Unexpected workbook relationship path")
        output.append((sheet.attrib["name"], member))
    return output


def verify_value_xml(path: Path, rounded: list[np.ndarray]) -> dict:
    """Second parser: direct ZIP/XML confirms every numeric value and unique cell.

    Independent of openpyxl's cell construction and artifact-tool's authoring.
    Repeated physical values are valid; duplicate row/cell coordinates are not.
    """
    count = 0
    with ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise AssertionError("XLSX ZIP checksum failure")
        sheets = _xml_sheet_paths(archive)
        if [n for n, _ in sheets] != list(NAMES):
            raise AssertionError("XML sheet names/order mismatch")
        for (name, member), field in zip(sheets, rounded):
            root = ET.fromstring(archive.read(member))
            rows = root.findall("x:sheetData/x:row", NS)
            if len(rows) != END_S + 1 or [int(row.attrib["r"]) for row in rows] != list(range(1, END_S + 2)):
                raise AssertionError(f"Missing, duplicate or unsorted XML rows: {name}")
            for row_index, row in enumerate(rows, 1):
                cells = row.findall("x:c", NS)
                if len(cells) != 22:
                    raise AssertionError(f"Wrong XML cell count: {name}, row {row_index}")
                for column, cell in enumerate(cells, 1):
                    expected_ref = f"{openpyxl.utils.get_column_letter(column)}{row_index}"
                    if cell.attrib.get("r") != expected_ref:
                        raise AssertionError(f"Duplicate, missing or unsorted XML cell: {name}!{expected_ref}")
                    if column == row_index == 1:
                        continue  # A1 text and style checked by independent workbook readback.
                    value = cell.find("x:v", NS)
                    if cell.find("x:f", NS) is not None or cell.attrib.get("t", "n") != "n" or value is None:
                        raise AssertionError(f"Non-numeric/static XML value: {name}!{expected_ref}")
                    actual = float(value.text)
                    expected = (column - 2) / 10 if row_index == 1 else row_index - 1 if column == 1 else float(field[row_index - 2, column - 2])
                    if not math.isfinite(actual) or actual != expected:
                        raise AssertionError(f"XML numeric mismatch: {name}!{expected_ref}, {actual} != {expected}")
                    if row_index > 1 and column > 1:
                        count += 1
    if count != 453600:
        raise AssertionError("Incomplete XML verification")
    return {"xml_readback_pass": True, "checked_numeric_results": count,
            "coordinate_uniqueness_checked": True, "all_values_static_numeric": True}


def restore_template_styles(path: Path, rounded: list[np.ndarray], audit: Path) -> dict:
    """Restore the original template's actual style properties, never data."""
    before = verify_value_xml(path, rounded)
    preserved = audit / "artifact-before-style-restoration.xlsx"
    if preserved.exists():
        raise FileExistsError(preserved)
    preserved.write_bytes(path.read_bytes())
    ref, book = openpyxl.load_workbook(TEMPLATE), openpyxl.load_workbook(path)
    try:
        book.loaded_theme = ref.loaded_theme
        for name in NAMES:
            original, sheet = ref[name], book[name]
            for anchor in ("A1", "B1", "A2", "B2"):
                for attr in (*STYLE_KEYS, "number_format"):
                    setattr(sheet[anchor], attr, copy(getattr(original[anchor], attr)))
            sheet["B1"].number_format = "0.0"
            sheet["A2"].number_format = "0"
            sheet["B2"].number_format = "0.0000"
            # These IDs now belong to this output workbook after actual-property copying.
            styles = {anchor: copy(sheet[anchor]._style) for anchor in ("B1", "A2", "B2")}
            sheet.column_dimensions["A"].width = original.column_dimensions["A"].width
            sheet.row_dimensions[1].height = original.row_dimensions[1].height
            for row in sheet.iter_rows(min_row=1, max_row=END_S + 1, max_col=22):
                for cell in row:
                    if cell.coordinate == "A1":
                        continue
                    source = "B1" if cell.row == 1 else "A2" if cell.column == 1 else "B2"
                    cell._style = copy(styles[source])
        book.save(path)
    finally:
        book.close()
        ref.close()
    after = verify_value_xml(path, rounded)
    return {"status": "ORIGINAL_STYLES_RESTORED_ALL_VALUES_UNCHANGED",
            "before_sha256": digest(preserved), "after_sha256": digest(path),
            "before_xml_readback": before, "after_xml_readback": after}


def verify_xlsx(path: Path, rounded: list[np.ndarray]) -> dict:
    """Read all 453600 results; verify layout, numeric types, format and 60 links."""
    book = openpyxl.load_workbook(path, read_only=False, data_only=False)
    ref = openpyxl.load_workbook(TEMPLATE)
    count = links = 0
    try:
        if book.sheetnames != list(NAMES):
            raise AssertionError("Wrong worksheet names/order")
        checks = []
        for name, field in zip(NAMES, rounded):
            sheet, original = book[name], ref[name]
            if (sheet.max_row, sheet.max_column) != (END_S + 1, 22):
                raise AssertionError(f"Wrong sheet dimensions: {name}")
            if sheet["A1"].value != original["A1"].value or sheet["A1"].value != A1:
                raise AssertionError("A1 content changed")
            if list(map(str, sheet.merged_cells.ranges)) != list(map(str, original.merged_cells.ranges)):
                raise AssertionError("Merged cells changed")
            if sheet.freeze_panes != original.freeze_panes:
                raise AssertionError("Freeze pane layout changed")
            for attr in (*STYLE_KEYS, "number_format"):
                if copy(getattr(sheet["A1"], attr)) != copy(getattr(original["A1"], attr)):
                    raise AssertionError(f"A1 {attr} changed: {name}")
            if sheet.column_dimensions["A"].width != original.column_dimensions["A"].width or sheet.row_dimensions[1].height != original.row_dimensions[1].height:
                raise AssertionError("A1 row/column size changed")
            for row in sheet.iter_rows(min_row=1, max_row=END_S + 1, max_col=22):
                for cell in row:
                    if cell.coordinate == "A1":
                        continue
                    if type(cell.value) not in (int, float) or cell.data_type != "n" or not math.isfinite(cell.value):
                        raise AssertionError(f"Non-numeric value: {name}!{cell.coordinate}")
                    expected = (cell.column - 2) / 10 if cell.row == 1 else cell.row - 1 if cell.column == 1 else float(field[cell.row - 2, cell.column - 2])
                    fmt = "0.0" if cell.row == 1 else "0" if cell.column == 1 else "0.0000"
                    if cell.value != expected or cell.number_format != fmt:
                        raise AssertionError(f"Value/format mismatch: {name}!{cell.coordinate}")
                    if cell.row > 1 and cell.column > 1:
                        count += 1
            for t in TABLE_TIMES:
                for j in TABLE_RADIAL_INDEX:
                    if sheet.cell(t + 1, j + 2).value != float(field[t - 1, j]):
                        raise AssertionError("Paper-table index mismatch")
                    links += 1
            checks.append({"sheet": name, "data_shape": [END_S, 21], "time_s": [1, END_S],
                           "radius_cm": [0, 2], "A1_and_styles_preserved": True})
    finally:
        book.close()
        ref.close()
    if count != 453600 or links != 60:
        raise AssertionError("Incomplete workbook readback coverage")
    return {"readback_pass": True, "checked_numeric_results": count, "checked_paper_table_links": links, "sheets": checks}


def paper_tables(output_dir: Path, rounded: list[np.ndarray], authority_sha: str, status: str) -> list[Path]:
    files = []
    for index, (name, field, unit) in enumerate(zip(NAMES, rounded, ("℃", "kg/kg")), 3):
        selected = field[np.array(TABLE_TIMES) - 1][:, TABLE_RADIAL_INDEX]
        header = ["时间/h", "r=0 cm", "r=0.5 cm", "r=1 cm", "r=1.5 cm", "r=2 cm"]
        rows = [[f"{t / 3600:g}", *(f"{v:.4f}" for v in values)] for t, values in zip(TABLE_TIMES, selected)]
        csv_path = output_dir / f"table{index}.csv"
        with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
            csv.writer(handle).writerows([header, *rows])
        md_path = output_dir / f"table{index}.md"
        md = [f"表 {index}　{name}（{unit}）", "", f"状态：{status}", f"权威数据 SHA-256：`{authority_sha}`", "",
              "| " + " | ".join(header) + " |", "| " + " | ".join(["---:"] * 6) + " |"]
        md.extend("| " + " | ".join(row) + " |" for row in rows)
        md_path.write_text("\n".join(md) + "\n", encoding="utf-8")
        tex_path = output_dir / f"table{index}.tex"
        tex = [f"% {status}; authority SHA256: {authority_sha}", r"\begin{table}[htbp]", r"\centering",
               f"\\caption{{{name}（{unit}）}}", f"\\label{{tab:q2-{index}}}", r"\begin{tabular}{rrrrrr}",
               r"\hline", r"时间/h & 0 cm & 0.5 cm & 1 cm & 1.5 cm & 2 cm \\", r"\hline"]
        tex.extend(" & ".join(row) + r" \\" for row in rows)
        tex.extend([r"\hline", r"\end{tabular}", r"\end{table}"])
        tex_path.write_text("\n".join(tex) + "\n", encoding="utf-8")
        files.extend([csv_path, md_path, tex_path])
    return files


def verify_paper_tables(output_dir: Path, rounded: list[np.ndarray]) -> dict:
    counts = {"csv": 0, "md": 0, "tex": 0}
    for index, field in enumerate(rounded, 3):
        expected = field[np.array(TABLE_TIMES) - 1][:, TABLE_RADIAL_INDEX]
        with (output_dir / f"table{index}.csv").open(encoding="utf-8-sig", newline="") as handle:
            csv_rows = list(csv.reader(handle))[1:]
        md_rows = [line.strip().strip("|").split("|") for line in (output_dir / f"table{index}.md").read_text(encoding="utf-8").splitlines()
                   if re.match(r"^\|\s*[0-9]", line)]
        tex_rows = [line.removesuffix(r" \\").split("&") for line in (output_dir / f"table{index}.tex").read_text(encoding="utf-8").splitlines()
                    if re.match(r"^[0-9]", line)]
        for kind, rows in (("csv", csv_rows), ("md", md_rows), ("tex", tex_rows)):
            if len(rows) != 6 or any(len(row) != 6 for row in rows):
                raise AssertionError(f"Wrong table {index} {kind} dimensions")
            arr = np.array([[float(item.strip()) for item in row] for row in rows])
            if not np.array_equal(arr[:, 0], np.array(TABLE_TIMES) / 3600) or not np.array_equal(arr[:, 1:], expected):
                raise AssertionError(f"Table {index} {kind} differs from authority/Excel")
            for row in rows:
                if any(not re.fullmatch(r"\d+\.\d{4}", token.strip()) for token in row[1:]):
                    raise AssertionError(f"Table {index} {kind} lacks four-place display")
            counts[kind] += 30
    return {"paper_tables_pass": True, "unique_physical_values": 60, "checked_values_per_format": counts,
            "time_unit": "h", "radius_unit": "cm", "latex_compilation_claimed": False}


BUILDER = r'''// Q2 static values only. Supported public artifact-tool API.
import fs from "node:fs/promises";
import path from "node:path";
import { createRequire } from "node:module";
import { pathToFileURL } from "node:url";
const args={};
for(let i=2;i<process.argv.length;i+=2) args[process.argv[i].replace(/^--/,"")]=process.argv[i+1];
const req=createRequire(path.join(path.resolve(args.runtime),"runtime-entry.js"));
const {FileBlob,SpreadsheetFile}=await import(pathToFileURL(req.resolve("@oai/artifact-tool")).href);
const wb=await SpreadsheetFile.importXlsx(await FileBlob.load(args.template));
const names=["温度","水分浓度"];
await fs.mkdir(args.preview,{recursive:true});
const info=await wb.inspect({kind:"workbook,sheet,table",maxChars:4500,tableMaxRows:5,tableMaxCols:6});
await fs.writeFile(path.join(args.preview,"inspection.ndjson"),info.ndjson);
if(args.payload){
  const p=JSON.parse(await fs.readFile(args.payload,"utf8"));
  if(p.times.length!==10800 || p.radius_cm.length!==21) throw new Error("Wrong axes");
  for(let i=0;i<2;i++){
    const s=wb.worksheets.getItem(names[i]);
    if(s.getRange("A1").values[0][0]!=="时间\\到药材中心的距离") throw new Error("A1 changed");
    if(p.values[i].length!==10800 || p.values[i].some(r=>r.length!==21 || r.some(v=>!Number.isFinite(v)))) throw new Error("Bad field");
    s.getRange("A2:A10801").copyFrom(s.getRange("A2"),"all");
    s.getRange("B1:V1").copyFrom(s.getRange("B1"),"all");
    s.getRange("B2:V10801").copyFrom(s.getRange("B2"),"all");
    s.getRange("A2:A10801").values=p.times.map(t=>[t]);
    s.getRange("B1:V1").values=[p.radius_cm];
    s.getRange("B2:V10801").values=p.values[i];
    s.getRange("A2:A10801").setNumberFormat("0");
    s.getRange("B1:V1").setNumberFormat("0.0");
    s.getRange("B2:V10801").setNumberFormat("0.0000");
    s.getRange("B1:V10801").format.columnWidth=9.625;
    s.getRange("A2:V10801").format.rowHeight=14.1;
  }
  wb.recalculate();
  const errors=await wb.inspect({kind:"match",searchTerm:"#REF!|#DIV/0!|#VALUE!|#NAME\\?|#N/A|#NUM!|#NULL!|#SPILL!|#CALC!",options:{useRegex:true,maxResults:20},summary:"Q2 static value scan"});
  await fs.writeFile(path.join(args.preview,"formula-error-scan.ndjson"),errors.ndjson);
  const out=await SpreadsheetFile.exportXlsx(wb);
  await out.save(args.output);
}
for(let i=0;i<2;i++){
  const ranges=args.mode==="template" ? ["A1:F5"] : ["A1:V8","A10795:V10801"];
  for(let j=0;j<ranges.length;j++){
    const blob=await wb.render({sheetName:names[i],range:ranges[j],scale:1.5,format:"png"});
    await fs.writeFile(path.join(args.preview,`sheet-${i+1}-range-${j+1}.png`),new Uint8Array(await blob.arrayBuffer()));
  }
}
console.log(JSON.stringify({exported:!!args.payload,output:args.output||null,preview:args.preview}));
'''


def artifact_runtime() -> tuple[Path, Path, Path]:
    """Only the loader-provided bundle; never install or inspect its internals."""
    runtime_root = Path(os.environ["USERPROFILE"]) / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node"
    node, packages = runtime_root / "bin/node.exe", runtime_root / "node_modules"
    marker = Path(os.environ["USERPROFILE"]) / ".codex/plugins/cache/openai-primary-runtime/spreadsheets/26.909.12148/skills/spreadsheets/container_tools/mark_artifact_operation_started.mjs"
    if not all(path.exists() for path in (node, packages, marker)):
        raise FileNotFoundError("Previously verified bundled spreadsheet runtime unavailable; do not install automatically")
    runtime = PREP / "runtime"
    runtime.mkdir(parents=True, exist_ok=True)
    link = runtime / "node_modules"
    if not link.exists():
        def quote(value):
            return "'" + str(value).replace("'", "''") + "'"
        subprocess.run(["powershell.exe", "-NoProfile", "-Command", f"New-Item -ItemType Junction -Path {quote(link)} -Target {quote(packages)} | Out-Null"], check=True, cwd=ROOT)
    if link.resolve() != packages.resolve():
        raise ValueError("Runtime junction points at an unexpected dependency directory")
    return node, runtime, marker


def run_node(node: Path, args: list[str | Path], audit: Path, stem: str) -> dict:
    command = [str(node), *map(str, args)]
    started = time.perf_counter()
    try:
        done = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=300)
    except subprocess.TimeoutExpired as error:
        save_json(audit / f"{stem}.json", {"command": command, "status": "TIMEOUT", "timeout_s": 300})
        raise RuntimeError(f"Artifact operation timed out; see {audit}") from error
    (audit / f"{stem}.stdout.log").write_text(done.stdout, encoding="utf-8")
    (audit / f"{stem}.stderr.log").write_text(done.stderr, encoding="utf-8")
    result = {"command": command, "returncode": done.returncode, "elapsed_s": time.perf_counter() - started}
    save_json(audit / f"{stem}.json", result)
    if done.returncode:
        raise RuntimeError(f"Artifact operation failed; see {audit / (stem + '.stderr.log')}")
    return result


def verify_delivery(authority: Path, acceptance_path: Path, config: Path, output: Path,
                    *, candidate: bool = False, fixture: bool = False) -> dict:
    """Read-only entry point that rejects stale evidence or changed exported files."""
    acceptance = json.loads(acceptance_path.read_text(encoding="utf-8"))
    gate = check_numerical_gate(acceptance, authority, config, candidate=candidate, fixture=fixture)
    manifest_path = output / "q2-export-manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    allowed_status = "FIXTURE_ONLY_EXPORT_READBACK_PASS" if fixture else "CANDIDATE_EXPORT_READBACK_PASS" if candidate else "Q2_NUMERICAL_EXPORT_READBACK_PASS"
    if manifest.get("status") != allowed_status or manifest.get("acceptance_sha256") != digest(acceptance_path):
        raise ValueError("Manifest status or acceptance binding no longer current")
    if manifest.get("formal_result") is not (not (candidate or fixture)) or manifest.get("fixture_only") is not fixture:
        raise ValueError("Manifest cannot change a fixture/candidate into a formal result")
    if (project_path(manifest["authority_path"]) != authority or project_path(manifest["acceptance_path"]) != acceptance_path
            or project_path(manifest["config_path"]) != config or manifest.get("authority_sha256") != digest(authority)):
        raise ValueError("Readback arguments must refer to the manifest's exact authority, acceptance, and configuration")
    for group in ("input_sha256", "source_sha256", "output_sha256"):
        for value, expected in manifest[group].items():
            path = project_path(value)
            if not path.is_file() or digest(path) != expected:
                raise ValueError(f"Current delivery file changed: {path}")
    _, _, fields = load_authority(authority, fixture=fixture)
    rounded = [round_half_up(field[1:]) for field in fields]
    xml = verify_value_xml(output / "result2.xlsx", rounded)
    workbook = verify_xlsx(output / "result2.xlsx", rounded)
    tables = verify_paper_tables(output, rounded)
    return {"status": allowed_status, "mode": "CURRENT_READONLY_VERIFY", "gate": gate,
            "xml": xml, "workbook": workbook, "tables": tables,
            "human_review_claimed": False, "contest_submission_release_claimed": False}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--authority")
    parser.add_argument("--acceptance")
    parser.add_argument("--config", default="config/q2-full.json")
    parser.add_argument("--output-dir", default="results/q2")
    parser.add_argument("--audit-dir", default="records/q2-full-20260911/export")
    parser.add_argument("--candidate", action="store_true")
    parser.add_argument("--fixture", action="store_true")
    parser.add_argument("--preflight", action="store_true")
    parser.add_argument("--verify", action="store_true")
    parser.add_argument("--inspect-template", action="store_true")
    args = parser.parse_args()
    check_python()
    if args.inspect_template:
        audit = PREP / "template-intake"
        if audit.exists():
            raise FileExistsError("Template intake already exists; preserve the current evidence")
        audit.mkdir(parents=True)
        save_json(audit / "template-metadata.json", template_metadata())
        builder = audit / "q2-workbook.mjs"
        builder.write_text(BUILDER, encoding="utf-8")
        node, runtime, _ = artifact_runtime()
        run_node(node, [builder, "--runtime", runtime, "--template", TEMPLATE, "--preview", audit / "preview", "--mode", "template"], audit, "template-render")
        print(json.dumps({"status": "READONLY_TEMPLATE_INTAKE_COMPLETE", "audit": str(audit)}, ensure_ascii=False))
        return 0
    if not args.authority or not args.acceptance:
        parser.error("--authority and --acceptance are required")
    authority, acceptance_path, config, output, audit = map(project_path, (args.authority, args.acceptance, args.config, args.output_dir, args.audit_dir))
    if args.fixture:
        if not all(path.is_relative_to(PREP) for path in (authority, acceptance_path, output, audit)):
            raise ValueError("Fixture I/O is confined to export-preparation records")
    elif not output.is_relative_to(ROOT / "results") or not audit.is_relative_to(ROOT / "records"):
        raise ValueError("Real exports belong under results; audit logs under records")
    if args.verify:
        print(json.dumps(verify_delivery(authority, acceptance_path, config, output, candidate=args.candidate, fixture=args.fixture), ensure_ascii=False))
        return 0
    acceptance = json.loads(acceptance_path.read_text(encoding="utf-8"))
    gate = check_numerical_gate(acceptance, authority, config, candidate=args.candidate, fixture=args.fixture)
    _, _, fields = load_authority(authority, fixture=args.fixture)
    template = template_metadata()
    if args.preflight:
        print(json.dumps({"status": "PREFLIGHT_PASS_NO_OUTPUT_CREATED", "gate": gate, "template": template}, ensure_ascii=False))
        return 0
    targets = [output / "result2.xlsx", output / "q2-export-manifest.json", *[output / f"table{i}.{ext}" for i in (3, 4) for ext in ("csv", "md", "tex")]]
    if any(p.exists() for p in targets) or audit.exists():
        raise FileExistsError("Refuse overwrite; use a fresh output/audit directory and preserve older evidence")
    output.mkdir(parents=True, exist_ok=True)
    audit.mkdir(parents=True)
    started = time.perf_counter()
    input_hashes = {str(p): digest(p) for p in (authority, acceptance_path, config, ROOT / "project_config.json", *RAW_INPUTS, *RAW_TEMPLATES)}
    source_hashes = {str(SCRIPT): digest(SCRIPT)}
    status = "FIXTURE_ONLY_EXPORT_READBACK_PASS" if args.fixture else "CANDIDATE_EXPORT_READBACK_PASS" if args.candidate else "Q2_NUMERICAL_EXPORT_READBACK_PASS"
    manifest = {"status": "EXPORT_RUNNING_NOT_FORMAL", "formal_result": False,
        "fixture_only": args.fixture, "question": "Q2", "authority_path": str(authority),
        "authority_sha256": digest(authority), "acceptance_path": str(acceptance_path),
        "acceptance_sha256": digest(acceptance_path), "config_path": str(config), "gate": gate,
        "input_sha256": input_hashes, "source_sha256": source_hashes,
        "command": [sys.executable, *sys.argv], "started_utc": datetime.now(timezone.utc).isoformat(),
        "python": sys.executable, "versions": {p: importlib.metadata.version(p) for p in ("numpy", "openpyxl")},
        "artifact_bundle": "26.909.12148", "seed": None, "seed_reason": "Deterministic export, no random process",
        "workers": 1,
        "export_rule": "Decimal.from_float(binary64).quantize(0.0001, ROUND_HALF_UP); typed numeric cells; 0.0000 display",
        "units": {"time_axis": "s", "radius_axis": "cm", "temperature": "degree C", "moisture": "dry-basis kg/kg", "paper_time": "h"},
        "human_review_claimed": False, "contest_submission_release_claimed": False}
    save_json(audit / "export-progress.json", manifest)
    try:
        rounded = [round_half_up(field[1:]) for field in fields]
        payload = audit / "rounded-payload.json"
        save_json(payload, {"times": list(range(1, END_S + 1)), "radius_cm": [j / 10 for j in range(21)], "values": [field.tolist() for field in rounded]})
        builder = audit / "q2-workbook.mjs"
        builder.write_text(BUILDER, encoding="utf-8")
        node, runtime, marker = artifact_runtime()
        # Marker is executed exactly once, immediately before this authoring command.
        run_node(node, [marker, "--operation-kind", "edit", "--expected-output-count", "1", "--output-format", "xlsx"], audit, "operation-marker")
        manifest["artifact_author"] = run_node(node, [builder, "--runtime", runtime, "--template", TEMPLATE, "--payload", payload,
            "--output", output / "result2.xlsx", "--preview", audit / "artifact-preview", "--mode", "result"], audit, "artifact-author")
        # The library can emit a large inspection sidecar beside the workbook.
        # Keep tool diagnostics in records, not in the concise results directory.
        sidecar = output / "result2.xlsx.inspect.ndjson"
        if sidecar.exists():
            relocated = audit / "artifact-generated-inspection.ndjson"
            if relocated.exists():
                raise FileExistsError(relocated)
            sidecar.rename(relocated)
            manifest["artifact_generated_inspection"] = {"path": str(relocated), "sha256": digest(relocated), "bytes": relocated.stat().st_size}
        manifest["style_restoration"] = restore_template_styles(output / "result2.xlsx", rounded, audit)
        manifest["workbook_readback"] = verify_xlsx(output / "result2.xlsx", rounded)
        paper_files = paper_tables(output, rounded, digest(authority), status)
        manifest["paper_tables_readback"] = verify_paper_tables(output, rounded)
        # Final saved workbook preview, after the scoped original-style restoration.
        manifest["final_render"] = run_node(node, [builder, "--runtime", runtime, "--template", output / "result2.xlsx",
            "--preview", audit / "final-preview", "--mode", "result"], audit, "final-render")
        for value, expected in {**input_hashes, **source_hashes}.items():
            if digest(Path(value)) != expected:
                raise AssertionError(f"Input/code changed during export: {value}")
        manifest["gate_after_export"] = check_numerical_gate(acceptance, authority, config, candidate=args.candidate, fixture=args.fixture)
        manifest["output_sha256"] = {str(p): digest(p) for p in (output / "result2.xlsx", *paper_files)}
        manifest["generated_builder_sha256"] = digest(builder)
        manifest["status"] = status
        manifest["formal_result"] = not (args.candidate or args.fixture)
        manifest["visual_review"] = "Previews generated; actual image inspection must be recorded by the reviewing agent. No human review is claimed."
        manifest["elapsed_s"] = time.perf_counter() - started
        save_json(output / "q2-export-manifest.json", manifest)
        save_json(audit / "export-progress.json", manifest)
        print(json.dumps({"status": status, "formal_result": manifest["formal_result"], "fixture_only": args.fixture,
                          "numeric_values_read_back": 453600, "paper_values_linked": 60,
                          "elapsed_s": manifest["elapsed_s"], "output": str(output), "audit": str(audit)}, ensure_ascii=False))
    except Exception as error:
        manifest.update(status="EXPORT_FAILED_NOT_FORMAL", formal_result=False,
                        error_type=type(error).__name__, error=str(error), elapsed_s=time.perf_counter() - started)
        save_json(audit / "export-progress.json", manifest)
        raise
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
