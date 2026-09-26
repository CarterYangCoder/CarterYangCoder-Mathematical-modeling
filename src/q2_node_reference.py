"""Independent node-control-volume reference for the approved Q2-R PDE.

N intervals give N+1 point unknowns, including a dynamic surface point.  Point
mass lumping over truncated nodal control volumes is a spatial approximation.
Robin supplies the exterior flux; no algebraic surface elimination or Phi is
used.  A finite-grid surface difference is NOT asserted to satisfy pointwise
Robin exactly.  This module imports no Q1/Q2 production solver.

Only the approved 0--121 s diagnostic is exposed by run_reference.  The
standalone --verify command solves explicit test configurations, never a real
three-hour run.  The shared SciPy BDF library is not an independent time code.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
from importlib.metadata import version
import json
from pathlib import Path
import sys
import time
from typing import Callable

import numpy as np
from numpy.polynomial.legendre import leggauss
import openpyxl
from scipy.integrate import BDF
from scipy.optimize import brentq
from scipy.sparse import coo_matrix, csr_matrix
from scipy.special import j0, j1, jn_zeros
from threadpoolctl import threadpool_limits

ROOT = Path(__file__).resolve().parents[1]
RECORDS = ROOT / 'records/q2-stage1-20260911/node-reference'
RADIUS = 0.02
LENGTH = 0.25
T_INITIAL = 28.0
C_INITIAL = 2.55
ENV_SHA256 = '7ef32870abeef420b89560b2530ff60dfe4255917805151d89988d0311af9dd7'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save_json(path, obj):
    Path(path).write_text(json.dumps(obj, ensure_ascii=False, indent=2,
                                    allow_nan=False) + '\n', encoding='utf-8')


def properties(T, C, constant=None):
    """W, k, D and exact partials; complex extension only aids Jacobian tests."""
    if np.any(np.real(C) <= 0) or np.any(np.real(T) + 273.15 <= 0):
        raise ValueError('Reference encountered C<=0 or absolute T<=0; no clipping.')
    if not (np.all(np.isfinite(T)) and np.all(np.isfinite(C))):
        raise FloatingPointError('Nonfinite reference state.')
    if constant is not None:
        shape = np.broadcast(T, C).shape
        W, k, D = [np.full(shape, val) for val in constant]
        zero = np.zeros(shape)
        return W, k, D, zero, zero, zero, zero
    rho = 650.0 + 128.0*C
    cp = 1450.0 + 2736.0*C/(1.0+C)
    W = rho*cp
    Wc = 128.0*cp + rho*2736.0/(1.0+C)**2
    k = 0.21 + 0.38*C/(1.0+C)
    kc = 0.38/(1.0+C)**2
    TK = T + 273.15
    D = 0.0024*np.exp(-0.45/C - 3850.0/TK)
    Dc = D*0.45/C**2
    DT = D*3850.0/TK**2
    return W, k, D, Wc, kc, Dc, DT


@dataclass
class OriginalDrive:
    times: np.ndarray
    temperature: np.ndarray
    concentration: np.ndarray
    source_sha256: str

    @classmethod
    def read(cls):
        path = ROOT / 'A题/附件/附件1.xlsx'
        digest = sha(path)
        if digest != ENV_SHA256:
            raise ValueError('Environment version differs from the read-only Q2 input review.')
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            if wb.sheetnames != ['Sheet1']:
                raise ValueError('Unexpected environment sheets.')
            ws = wb['Sheet1']
            rows = list(ws.iter_rows(values_only=True))
            if tuple(rows[0]) != ('时间', '温度', '水分浓度'):
                raise ValueError('Unexpected environment headers.')
            a = np.asarray(rows[1:], dtype=float)
        finally:
            wb.close()
        if a.shape != (241, 3) or not np.all(np.isfinite(a)):
            raise ValueError('Unexpected or nonfinite environment cells.')
        if not np.array_equal(a[:, 0], np.arange(241, dtype=float)*60.0):
            raise ValueError('Missing/reordered original environment knots.')
        return cls(a[:, 0], a[:, 1], a[:, 2], digest)

    def __call__(self, t):
        if not self.times[0] <= t <= self.times[-1]:
            raise ValueError('Environment extrapolation is forbidden.')
        return (float(np.interp(t, self.times, self.temperature)),
                float(np.interp(t, self.times, self.concentration)))


class NodeModel:
    """Shared fluxes, local midpoint constitutive values, nodal mass lumping."""
    def __init__(self, n, drive, *, hT=25.0, hC=8e-7, constant=None,
                 source: Callable | None = None):
        if int(n) != n or n < 4:
            raise ValueError('At least four radial intervals required.')
        self.n = int(n)
        self.r = np.linspace(0.0, RADIUS, self.n+1)
        self.h = RADIUS/self.n
        self.faces = np.r_[0.0, (self.r[:-1]+self.r[1:])/2.0, RADIUS]
        # Omit one common factor 2*pi*L from ALL volumes and areas.
        self.v = np.diff(self.faces**2)/2.0
        self.face_g = self.faces[1:-1]/np.diff(self.r)
        self.physical_volume = self.v*(2.0*np.pi*LENGTH)
        self.drive = drive
        self.hT = float(hT)
        self.hC = float(hC)
        self.constant = constant
        self.source = source
        assert np.all(self.v > 0)
        assert abs(np.sum(self.v)-RADIUS**2/2) < 1e-18

    def values(self, t, y):
        T, C = y[0::2], y[1::2]
        W, k, D, Wc, kc, Dc, DT = properties(T, C, self.constant)
        Tm = (T[:-1]+T[1:])/2
        Cm = (C[:-1]+C[1:])/2
        _, km, Dm, _, kcm, Dcm, DTm = properties(Tm, Cm, self.constant)
        Ta, Ce = self.drive(t)
        FT = self.face_g*km*(T[:-1]-T[1:])
        FC = self.face_g*Dm*(C[:-1]-C[1:])
        FTs = RADIUS*self.hT*(T[-1]-Ta)
        FCs = RADIUS*self.hC*(C[-1]-Ce)
        netT = np.zeros_like(T)
        netC = np.zeros_like(C)
        netT[:-1] -= FT
        netT[1:] += FT
        netC[:-1] -= FC
        netC[1:] += FC
        netT[-1] -= FTs
        netC[-1] -= FCs
        qT = np.zeros_like(T)
        qC = np.zeros_like(C)
        if self.source is not None:
            # Explicit manufactured source in the ORIGINAL PDE units:
            # qT [W/m^3], qC [(kg/kg)/s]. It is not fitted to a computed state.
            qT, qC = self.source(self.r, t)
        dT = (netT/self.v + qT)/W
        dC = netC/self.v + qC
        return T, C, W, Wc, km, kcm, Dm, Dcm, DTm, dT, dC, FTs, FCs, qT, qC

    def rhs(self, t, y):
        v = self.values(t, y)
        out = np.empty_like(y)
        out[0::2] = v[9]
        out[1::2] = v[10]
        return out

    def jac(self, t, y):
        T, C, W, Wc, km, kcm, Dm, Dcm, DTm, dT, *_ = self.values(t, y)
        nodes = np.arange(self.n+1)
        left, right = nodes[:-1], nodes[1:]
        # Four derivatives of each face flux, ordered TL, CL, TR, CR.
        dFT = [self.face_g*km,
               self.face_g*kcm*(T[:-1]-T[1:])/2,
               -self.face_g*km,
               self.face_g*kcm*(T[:-1]-T[1:])/2]
        dFC = [self.face_g*DTm*(C[:-1]-C[1:])/2,
               self.face_g*(Dm+Dcm*(C[:-1]-C[1:])/2),
               self.face_g*DTm*(C[:-1]-C[1:])/2,
               self.face_g*(-Dm+Dcm*(C[:-1]-C[1:])/2)]
        columns = [2*left, 2*left+1, 2*right, 2*right+1]
        rr, cc, vv = [], [], []
        for component, derivs in [(0, dFT), (1, dFC)]:
            for row_nodes, sign in [(left, -1.0), (right, 1.0)]:
                denom = self.v[row_nodes]*(W[row_nodes] if component == 0 else 1.0)
                for col, derivative in zip(columns, derivs):
                    rr.append(2*row_nodes+component)
                    cc.append(col)
                    vv.append(sign*derivative/denom)
        # Derivative of division by W(C), including explicit heat source.
        rr += [2*nodes, np.array([2*self.n, 2*self.n+1])]
        cc += [2*nodes+1, np.array([2*self.n, 2*self.n+1])]
        vv += [-dT*Wc/W,
               np.array([-RADIUS*self.hT/(self.v[-1]*W[-1]),
                         -RADIUS*self.hC/self.v[-1]])]
        size = 2*(self.n+1)
        return coo_matrix((np.concatenate(vv),
                           (np.concatenate(rr), np.concatenate(cc))),
                          shape=(size, size)).tocsc()


def point_sampling_matrix(nodes, radii):
    """Independent local quadratic point interpolation in even coordinate r^2.

    Actual r=0 and r=R nodes are returned directly. This is NOT a conversion
    from ring averages; interpolation is checked by manufactured solutions.
    """
    q = np.asarray(radii, dtype=float)
    if q.ndim != 1 or np.any(q < 0) or np.any(q > RADIUS):
        raise ValueError('Requested radii must be in [0,R].')
    if np.any(np.diff(q) <= 0):
        raise ValueError('Requested radii must strictly increase.')
    rows, cols, vals = [], [], []
    for j, r in enumerate(q):
        nearest = int(np.argmin(np.abs(nodes-r)))
        if abs(nodes[nearest]-r) <= 8*np.finfo(float).eps*RADIUS:
            inds, weights = [nearest], [1.0]
        else:
            center = min(max(nearest, 1), len(nodes)-2)
            inds = np.arange(center-1, center+2)
            xx = nodes[inds]**2
            target = r*r
            weights = []
            for i in range(3):
                other = [k for k in range(3) if k != i]
                weights.append(np.prod((target-xx[other])/(xx[i]-xx[other])))
        for i, w in zip(inds, weights):
            rows.append(j); cols.append(i); vals.append(w)
    return csr_matrix((vals, (rows, cols)), shape=(len(q), len(nodes)))


def _integrate(model, times, radii, initial, *, rtol, atol_T, atol_C,
               max_step, stop_s, segment_knots, audit_integrals=True):
    times = np.asarray(times, dtype=float)
    radii = np.asarray(radii, dtype=float)
    if (times.ndim != 1 or len(times) == 0 or np.any(np.diff(times) <= 0)
            or times[0] < 0 or times[-1] > stop_s):
        raise ValueError('Strictly increasing output times in [0,stop_s] required.')
    if any(x <= 0 for x in [rtol, atol_T, atol_C, max_step, stop_s]):
        raise ValueError('Positive tolerances, max_step and stop required.')
    bounds = np.unique(np.r_[0.0, np.asarray(segment_knots), stop_s])
    bounds = bounds[(bounds >= 0) & (bounds <= stop_s)]
    sample = point_sampling_matrix(model.r, radii)
    y = np.asarray(initial, dtype=float).copy()
    if y.shape != (2*(model.n+1),):
        raise ValueError('Initial full point state has wrong shape.')
    y0 = y.copy()
    fields = np.empty((len(times), 2, len(radii)))
    next_output = 0
    if times[0] == 0:
        fields[0, 0] = sample@y[0::2]
        fields[0, 1] = sample@y[1::2]
        next_output = 1
    atol = np.tile([atol_T, atol_C], model.n+1)
    start = time.perf_counter()
    stats = {'n_intervals': model.n, 'n_nodes': model.n+1,
             'state_definition': 'physical point values, nodal-control-volume mass lumping',
             'surface': 'dynamic point node; physical Robin exterior flux, finite half-ring storage',
             'interior_flux': 'direct midpoint-local Fick/Fourier; no Phi, no algebraic surface root',
             'sampling': 'local quadratic interpolation in r^2 of point values; endpoints exact nodes',
             'method': 'SciPy BDF; shared library, independent spatial/surface discretization',
             'rtol': rtol, 'atol_T': atol_T, 'atol_C': atol_C,
             'max_step_s': max_step, 'segments': [], 'accepted_steps': 0,
             'internal_rejected_steps': None,
             'rejection_note': 'Internal BDF Newton/step rejections are not exposed or counted here.',
             'accepted_state_T_min': float(y[0::2].min()),
             'accepted_state_C_min': float(y[1::2].min()),
             'accepted_state_T_max': float(y[0::2].max()),
             'accepted_state_C_max': float(y[1::2].max()),
             'min_step_s': None, 'max_step_actual_s': 0.0,
             'BDF_dense_output': 'SciPy piecewise polynomial at requested times; temporal convergence required',
             'nonphysical_state_clipping': False, 'failed': False}
    qx, qw = leggauss(3)
    mass_in, heat_in, heat_chain = 0.0, 0.0, 0.0
    max_alg_C = 0.0
    max_alg_T = 0.0
    for lo, hi in zip(bounds[:-1], bounds[1:]):
        if hi <= lo:
            continue
        clock0 = time.perf_counter()
        # BDF's RMS-based automatic first-step probe can overshoot a thin
        # surface startup layer because most interior derivatives are zero.
        # Supply a bounded first step instead of clipping that trial state.
        p0 = properties(y[0::2], y[1::2], model.constant)
        fastest = max(float(np.max(p0[1]/p0[0])), float(np.max(p0[2])))
        first_step = min(max_step, hi-lo, 0.01*model.h**2/fastest)
        f0 = model.rhs(lo, y)
        for val, der in [(y[1::2], f0[1::2]), (y[0::2]+273.15, f0[0::2])]:
            declining = der < 0
            if np.any(declining):
                first_step = min(first_step, 0.1*float(np.min(val[declining]/(-der[declining]))))
        solver = BDF(model.rhs, lo, y, hi, rtol=rtol, atol=atol,
                     first_step=first_step, max_step=max_step,
                     jac=model.jac, vectorized=False)
        count = 0
        while solver.status == 'running':
            old_t = solver.t
            message = solver.step()
            if solver.status == 'failed':
                raise RuntimeError(f'BDF failed at t={solver.t}: {message}')
            dt = solver.t-old_t
            stats['min_step_s'] = dt if stats['min_step_s'] is None else min(stats['min_step_s'], dt)
            stats['max_step_actual_s'] = max(stats['max_step_actual_s'], dt)
            count += 1
            dense = solver.dense_output()
            while next_output < len(times) and times[next_output] <= solver.t:
                out = solver.y if times[next_output] == solver.t else dense(times[next_output])
                fields[next_output, 0] = sample@out[0::2]
                fields[next_output, 1] = sample@out[1::2]
                next_output += 1
            T, C = solver.y[0::2], solver.y[1::2]
            if not np.all(np.isfinite(solver.y)) or np.min(C) <= 0 or np.min(T) <= -273.15:
                raise RuntimeError('BDF accepted a nonphysical state; no clipping performed.')
            for key, val in [('T_min', T.min()), ('C_min', C.min())]:
                stats['accepted_state_'+key] = min(stats['accepted_state_'+key], float(val))
            for key, val in [('T_max', T.max()), ('C_max', C.max())]:
                stats['accepted_state_'+key] = max(stats['accepted_state_'+key], float(val))
            if audit_integrals:
                # Three Gauss points integrate the exterior flux exactly for
                # each BDF degree<=5 dense polynomial and affine environment.
                # The heat-chain product is nonlinear, so its quadrature is
                # an independent empirical check, not an exact identity.
                for gx, gw in zip(qx, qw):
                    tq = (old_t+solver.t)/2 + dt*gx/2
                    Y = dense(tq)
                    vv = model.values(tq, Y)
                    TT, CC, WW, WWc, _, _, _, _, _, dTT, dCC, FT, FC, srcT, srcC = vv
                    mass_rhs = -FC + np.dot(model.v, srcC)
                    heat_rhs = -FT + np.dot(model.v, srcT)
                    mass_in += dt*gw*mass_rhs/2
                    heat_in += dt*gw*heat_rhs/2
                    heat_chain += dt*gw*np.dot(model.v, WWc*(TT-T_INITIAL)*dCC)/2
                    max_alg_C = max(max_alg_C, abs(float(np.dot(model.v, dCC)-mass_rhs)))
                    max_alg_T = max(max_alg_T, abs(float(np.dot(model.v*WW, dTT)-heat_rhs)))
        y = solver.y.copy()
        stats['accepted_steps'] += count
        stats['segments'].append({'start_s': float(lo), 'end_s': float(hi),
                                  'requested_first_step_s': first_step,
                                  'first_step_rule': 'min(max_step,segment_length,0.01*dr^2/max(k/W,D),0.1*Euler_positivity_time)',
                                  'accepted_steps': count, 'nfev': solver.nfev,
                                  'njev': solver.njev, 'nlu': solver.nlu,
                                  'wall_seconds': time.perf_counter()-clock0,
                                  'endpoint_reached': bool(solver.t == hi)})
    if next_output != len(times):
        raise RuntimeError('An output time was not produced.')
    stats['wall_seconds'] = time.perf_counter()-start
    factor = 2*np.pi*LENGTH
    W0 = properties(y0[0::2], y0[1::2], model.constant)[0]
    W1 = properties(y[0::2], y[1::2], model.constant)[0]
    mass_change = float(np.dot(model.v, y[1::2]-y0[1::2]))
    energy_change = float(np.dot(model.v, W1*(y[0::2]-T_INITIAL)-W0*(y0[0::2]-T_INITIAL)))
    stats['integral_audit'] = {'enabled': audit_integrals,
        'geometry_common_factor': factor, 'length_m': LENGTH,
        'total_control_volume_m3': float(np.sum(model.physical_volume)),
        'C_integral_change_m3_kg_kg': factor*mass_change,
        'C_integrated_input_m3_kg_kg': factor*mass_in,
        'C_balance_defect_m3_kg_kg': factor*(mass_change-mass_in),
        'C_balance_defect_normalized_by_initial_integral': (mass_change-mass_in)/float(np.dot(model.v,y0[1::2])),
        'Eref_change_J': factor*energy_change,
        'Qin_J': factor*heat_in, 'chain_correction_J': factor*heat_chain,
        'Eref_chain_defect_J': factor*(energy_change-heat_in-heat_chain),
        'algebraic_C_rhs_residual_reduced_units': max_alg_C,
        'algebraic_T_rhs_residual_reduced_units': max_alg_T,
        'qualification': 'C is not converted to kilograms of water; heat is effective PDE chain balance, not full physical enthalpy. Algebraic cancellation is not a spatial accuracy certificate.'}
    stats['source_sha256'] = sha(__file__)
    return {'times': times, 'radii': radii, 'T_C': fields[:, 0],
            'C_kg_kg': fields[:, 1], 'final_point_state': y,
            'node_radii_m': model.r, 'stats': stats}


def run_reference(n, times, radii, *, rtol, atol_T, atol_C, max_step,
                  stop_s=121):
    """Read original environment and solve ONLY the approved short interval."""
    if stop_s > 121:
        raise ValueError('This reference is authorized only through 121 s.')
    drive = OriginalDrive.read()
    model = NodeModel(n, drive)
    initial = np.tile([T_INITIAL, C_INITIAL], n+1)
    result = _integrate(model, times, radii, initial, rtol=rtol,
                        atol_T=atol_T, atol_C=atol_C, max_step=max_step,
                        stop_s=stop_s, segment_knots=drive.times)
    result['stats']['environment_sha256'] = drive.source_sha256
    result['stats']['test_configuration'] = False
    return result


def manufactured(r, t):
    """Analytical even nonpolynomial fields and derivatives, never solution-derived."""
    x = r/RADIUS
    aT, aC = 1.1, 1.5
    eT, eC = 2.0*np.exp(-t/80.0), 0.2*np.exp(-t/100.0)
    T, C = 30.0 + eT*np.cos(aT*x), 1.5 + eC*np.cos(aC*x)
    Tt, Ct = -eT*np.cos(aT*x)/80, -eC*np.cos(aC*x)/100
    Tr, Cr = -eT*aT/RADIUS*np.sin(aT*x), -eC*aC/RADIUS*np.sin(aC*x)
    Trr = -eT*(aT/RADIUS)**2*np.cos(aT*x)
    Crr = -eC*(aC/RADIUS)**2*np.cos(aC*x)
    Tr_r = np.empty_like(r); Cr_r = np.empty_like(r)
    np.divide(Tr, r, out=Tr_r, where=r != 0)
    np.divide(Cr, r, out=Cr_r, where=r != 0)
    Tr_r[r == 0] = Trr[r == 0]
    Cr_r[r == 0] = Crr[r == 0]
    return T, C, Tt, Ct, Tr, Cr, Trr+Tr_r, Crr+Cr_r


def manufactured_source(r, t):
    T, C, Tt, Ct, Tr, Cr, lapT, lapC = manufactured(r, t)
    W, k, D, Wc, kc, Dc, DT = properties(T, C)
    divT = k*lapT + kc*Cr*Tr
    divC = D*lapC + Dc*Cr*Cr + DT*Tr*Cr
    return W*Tt-divT, Ct-divC


def manufactured_drive(t):
    T, C, _, _, Tr, Cr, *_ = manufactured(np.array([RADIUS]), t)
    _, k, D, *_ = properties(T, C)
    return float(T[0]+k[0]*Tr[0]/25), float(C[0]+D[0]*Cr[0]/8e-7)


def record_solution(directory, result):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=False)
    np.savez_compressed(directory/'solution.npz', **{k:v for k,v in result.items() if k != 'stats'})
    save_json(directory/'stats.json', result['stats'])


def _max_record(error, times, radii):
    ii = np.unravel_index(np.argmax(np.abs(error)), error.shape)
    return {'max_abs': float(abs(error[ii])), 'signed': float(error[ii]),
            'time_s': float(times[ii[0]]), 'radius_m': float(radii[ii[1]])}


def verify_components(out):
    """Meaningful test-only runs; no original-environment PDE run here."""
    out = Path(out).resolve()
    out.mkdir(parents=True, exist_ok=False)
    started = time.perf_counter()
    test = {'created_utc': datetime.now(timezone.utc).isoformat(),
            'scope': 'Independent reference geometry/Jacobian/eigenmode/manufactured tests only',
            'script_sha256': sha(__file__), 'command': [sys.executable,*sys.argv],
            'versions': {x: version(x) for x in ['numpy','scipy','openpyxl','threadpoolctl']},
            'skills_read': ['cumcm-quality/SKILL.md','cumcm-ab-validation/SKILL.md'],
            'not_original_environment_run': True, 'checks': {}}
    # Verify exact shared-face signs and center behavior with u=a+b*r^2.
    geo = []
    for n in [8, 32, 128]:
        bT, bC = 200.0, 10.0
        const = (3.0e6, 0.4, 6e-9)
        m = NodeModel(n, lambda t: (28+bT*RADIUS**2+2*const[1]*bT*RADIUS/25,
                                    2.55+bC*RADIUS**2+2*const[2]*bC*RADIUS/8e-7),
                      constant=const)
        y = np.empty(2*(n+1)); y[0::2] = 28+bT*m.r**2; y[1::2] = 2.55+bC*m.r**2
        rhs = m.rhs(0,y)
        geo.append({'n':n,'volume_error_m3':float(np.sum(m.physical_volume)-np.pi*RADIUS**2*LENGTH),
                    'T_laplacian_error':float(np.max(np.abs(rhs[0::2]*const[0]/const[1]-4*bT))),
                    'C_laplacian_error':float(np.max(np.abs(rhs[1::2]/const[2]-4*bC)))})
    test['checks']['quadratic_operator'] = geo
    m = NodeModel(8, lambda t:(31.0,1.0), source=manufactured_source)
    y = np.empty(18); y[0::2] = 29+2*np.cos(1.2*m.r/RADIUS); y[1::2] = 1.5+0.3*np.cos(1.5*m.r/RADIUS)
    J = m.jac(0.4,y).toarray()
    Jcs = np.zeros_like(J)
    for j in range(len(y)):
        yy = y.astype(complex); yy[j] += 1e-28j
        Jcs[:,j] = np.imag(m.rhs(0.4,yy))/1e-28
    test['checks']['jacobian_complex_step'] = {
        'max_abs':float(np.max(np.abs(J-Jcs))),
        'max_relative_to_max1_entry':float(np.max(np.abs(J-Jcs)/np.maximum(1,np.abs(Jcs)))),
        'includes_W_derivative_and_manufactured_source':True}
    yzero = y.copy(); yzero[1::2] = 2.0
    plain = NodeModel(8, lambda t:(31.,2.0))
    cf = plain.rhs(0,yzero)[1::2]
    yzero = y.copy(); yzero[0::2] = 31.0
    tf = plain.rhs(0,yzero)[0::2]
    test['checks']['no_cross_gradient_flux'] = {'C_uniform_T_nonuniform_Ct_max':float(np.max(np.abs(cf))),
                                               'T_uniform_C_nonuniform_Tt_max':float(np.max(np.abs(tf)))}
    # True uniform no-exchange and exact-equilibrium tests of the nonlinear code.
    invariants = []
    for label, mm in [('zero_exchange',NodeModel(64,lambda t:(40.,.02),hT=0,hC=0)),
                      ('equilibrium',NodeModel(64,lambda t:(28.,2.55)))]:
        z = _integrate(mm,np.array([0.,1.,10.]),np.linspace(0,RADIUS,21),
                       np.tile([28.,2.55],65),rtol=1e-11,atol_T=1e-11,atol_C=1e-13,
                       max_step=0.5,stop_s=10,segment_knots=[0.,10.])
        invariants.append({'test':label,'T_max_delta':float(np.max(np.abs(z['T_C']-28))),
                           'C_max_delta':float(np.max(np.abs(z['C_kg_kg']-2.55))),
                           'wall_seconds':z['stats']['wall_seconds']})
    test['checks']['uniform_invariants'] = invariants
    times = np.array([0.,1.,5.,15.,30.,60.,90.,120.,121.])
    radii = np.linspace(0,RADIUS,41)
    base = properties(np.array([28.]),np.array([2.55]))
    const = tuple(float(base[j][0]) for j in range(3))
    BiT, BiC = 25*RADIUS/const[1], 8e-7*RADIUS/const[2]
    jt = float(jn_zeros(0,1)[0])
    lt = brentq(lambda x:x*j1(x)-BiT*j0(x),1e-10,jt,xtol=1e-14)
    lc = brentq(lambda x:x*j1(x)-BiC*j0(x),1e-10,jt,xtol=1e-14)
    test['checks']['constant_robin_roots'] = {'lambda_T':lt,'lambda_C':lc,
        'residual_T':float(lt*j1(lt)-BiT*j0(lt)),'residual_C':float(lc*j1(lc)-BiC*j0(lc))}
    for mode in ['constant_robin_eigenmode','nonpolynomial_manufactured']:
        records = []
        for n in [64,128,256]:
            if mode == 'constant_robin_eigenmode':
                mm = NodeModel(n,lambda t:(28.,1.),constant=const)
                initial = np.empty(2*(n+1))
                initial[0::2] = 28+0.5*j0(lt*mm.r/RADIUS)
                initial[1::2] = 1+0.1*j0(lc*mm.r/RADIUS)
                Tex = 28+0.5*np.exp(-const[1]/const[0]*(lt/RADIUS)**2*times[:,None])*j0(lt*radii[None,:]/RADIUS)
                Cex = 1+0.1*np.exp(-const[2]*(lc/RADIUS)**2*times[:,None])*j0(lc*radii[None,:]/RADIUS)
            else:
                mm = NodeModel(n,manufactured_drive,source=manufactured_source)
                T0,C0,*_ = manufactured(mm.r,0.)
                initial = np.empty(2*(n+1)); initial[0::2]=T0; initial[1::2]=C0
                Tex = np.array([manufactured(radii,t)[0] for t in times])
                Cex = np.array([manufactured(radii,t)[1] for t in times])
            result = _integrate(mm,times,radii,initial,rtol=1e-11,
                                atol_T=1e-11,atol_C=1e-13,max_step=0.25,
                                stop_s=121,segment_knots=[0.,60.,120.,121.])
            path = out/f'{mode}-N{n}'
            record_solution(path,result)
            records.append({'n':n,'T_error':_max_record(result['T_C']-Tex,times,radii),
                            'C_error':_max_record(result['C_kg_kg']-Cex,times,radii),
                            'wall_seconds':result['stats']['wall_seconds'],
                            'accepted_steps':result['stats']['accepted_steps'],
                            'result_path':str(path.relative_to(ROOT))})
        for prev,cur in zip(records[:-1],records[1:]):
            cur['observed_order_T']=float(np.log2(prev['T_error']['max_abs']/cur['T_error']['max_abs']))
            cur['observed_order_C']=float(np.log2(prev['C_error']['max_abs']/cur['C_error']['max_abs']))
        test['checks'][mode]=records
    test['wall_seconds'] = time.perf_counter()-started
    test['limitations'] = [
        'Known/manufactured tests do not establish accuracy of the incompatible real startup boundary layer.',
        'Nodal lumping uses point constitutive coefficients times finite control volume; truncation is measured rather than declared zero.',
        'Surface node carries finite half-ring storage; pointwise Robin derivative equality only emerges with refinement.',
        'The BDF library is shared with other references, while spatial fluxes/surface treatment are independent.',
        'No physical-model validation or official-answer claim follows from these tests.']
    save_json(out/'verification.json',test)
    return test


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--verify', action='store_true')
    parser.add_argument('--out', type=Path)
    args=parser.parse_args()
    if not args.verify:
        parser.error('Only --verify is provided as a standalone action; actual short runs require explicit orchestration.')
    out=args.out or (RECORDS/('verification-'+datetime.now().strftime('%Y%m%d-%H%M%S')))
    try:
        with threadpool_limits(limits=1):
            result=verify_components(out)
        print(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False))
    except Exception as exc:
        out.mkdir(parents=True,exist_ok=True)
        save_json(out/'failure.json',{'error':str(exc),'type':type(exc).__name__,
                                    'source_sha256':sha(__file__),'command':[sys.executable,*sys.argv]})
        raise


if __name__ == '__main__':
    main()
