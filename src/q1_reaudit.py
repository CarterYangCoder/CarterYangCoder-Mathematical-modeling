"""Read-only re-audit of Q1 inputs, published numbers and mechanistic scales.

Writes only a new audit directory. Does not solve a changed model, edit source
workbooks, draw figures, or overwrite the accepted numerical authority.
"""
from __future__ import annotations

import csv
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import sys

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "records/q1-reaudit-20260911"
FULL = ROOT / "records/q1-full-20260911"


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_new(path, value):
    with Path(path).open("x", encoding="utf-8") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)


def rounded(value):
    return Decimal.from_float(float(value)).quantize(Decimal(".0001"), rounding=ROUND_HALF_UP)


def rounding_margin(value):
    x = Decimal.from_float(float(value))
    y = rounded(value)
    half = Decimal(".00005")
    return float(min(abs(x - (y - half)), abs(x - (y + half))))


def maximum(matrix, times, radii):
    i, j = np.unravel_index(np.argmax(np.abs(matrix)), matrix.shape)
    return {"absolute_value": float(abs(matrix[i, j])), "time_s": float(times[i]),
            "radius_cm": float(radii[j] * 100), "signed_value": float(matrix[i, j])}


def main():
    OUT.mkdir(exist_ok=True)
    config = read_json(ROOT / "project_config.json")
    expected_python = Path(config["python"].replace("${USERPROFILE}", str(Path.home())))
    assert Path(sys.executable).resolve() == expected_python.resolve()
    accepted = read_json(ROOT / "config/q1-accepted.json")
    p = accepted["physical_parameters"]
    inputs = read_json(ROOT / "records/stage0-a-20260910/input-inventory.json")["files"]
    original_checks = []
    for item in inputs:
        path = ROOT / item["path"]
        found = digest(path)
        original_checks.append({"path": item["path"], "current_sha256": found,
                                "equals_initial_inventory": found == item["sha256_before"]})
    assert len(original_checks) == 7 and all(x["equals_initial_inventory"] for x in original_checks)

    wb = openpyxl.load_workbook(ROOT / accepted["inputs"]["environment"], read_only=True, data_only=False)
    ws = wb[accepted["inputs"]["environment_sheet"]]
    rows = list(ws.iter_rows(values_only=True))
    headers, values = rows[0], np.array(rows[1:], dtype=float)
    assert values.shape == (241, 3) and np.isfinite(values).all()
    assert np.array_equal(values[:, 0], np.arange(0, 14401, 60))
    wb.close()
    q1 = values[values[:, 0] <= 1800]
    assert q1.shape == (31, 3)

    authority_path = ROOT / "results/q1/q1-authority.npz"
    authority_before = digest(authority_path)
    with np.load(authority_path, allow_pickle=False) as z:
        time = z["time_s"].copy()
        radius = z["radius_m"].copy()
        T, C = z["T_C"].copy(), z["C_kg_kg"].copy()
        budgets = [z["T_empirical_budget_C"].copy(), z["C_empirical_budget_kg_kg"].copy()]
        saved_margins = [z["T_rounding_margin_C"].copy(), z["C_rounding_margin_kg_kg"].copy()]
    assert np.array_equal(time, np.arange(1801, dtype=float))
    assert np.array_equal(radius, np.arange(21) * .001)
    assert T.shape == C.shape == (1801, 21)
    assert np.isfinite(T).all() and np.isfinite(C).all()
    assert np.all(T[0] == p["initial_temperature_C"])
    assert np.all(C[0] == p["initial_moisture_kg_kg"])
    Ta, Ce = np.interp(time, q1[:, 0], q1[:, 1]), np.interp(time, q1[:, 0], q1[:, 2])

    analysis = read_json(FULL / "final-candidate-analysis.json")
    with np.load(FULL / "heat/heat-solution.npz", allow_pickle=False) as z:
        assert np.array_equal(z["times_s"], time) and np.array_equal(z["radii_m"], radius)
        Tref = z["reference_2048_temperature_C"].copy()
        Tref_budget = z["reference_empirical_budget_C"].copy()
    Cref_path = ROOT / analysis["independent_comparison"]["source"]
    assert digest(Cref_path) == analysis["independent_comparison"]["sha256"]
    with np.load(Cref_path, allow_pickle=False) as z:
        ix = np.searchsorted(z["time_s"], time)
        assert np.array_equal(z["time_s"][ix], time)
        assert np.allclose(z["radius_m"], radius, atol=1e-17, rtol=0)
        Cref = z["C_kg_kg"][ix].copy()

    results_path = ROOT / "results/q1/result1.xlsx"
    result_before = digest(results_path)
    wb = openpyxl.load_workbook(results_path, read_only=True, data_only=False)
    template = openpyxl.load_workbook(ROOT / accepted["inputs"]["result_template"], read_only=True, data_only=False)
    assert wb.sheetnames == template.sheetnames == ["温度", "水分浓度"]
    report, certificates, margins_all = [], [], []
    paper_times = [100, 300, 600, 900, 1200, 1500, 1800]
    paper_columns = [0, 5, 10, 15, 20]
    for k, (name, field, budget, reference) in enumerate(zip(wb.sheetnames, [T, C], budgets, [Tref, Cref])):
        sheet = wb[name]
        grid = list(sheet.iter_rows(values_only=True))
        assert sheet.max_row == 1801 and sheet.max_column == 22
        assert grid[0][0] == template[name]["A1"].value == "时间\\到药材中心的距离"
        assert np.array_equal(np.array([row[0] for row in grid[1:]], float), time[1:])
        assert np.allclose(np.array(grid[0][1:], float), radius * 100, atol=2e-15, rtol=0)
        expected = np.array([[float(rounded(v)) for v in row] for row in field[1:]])
        exported = np.array([row[1:] for row in grid[1:]], dtype=float)
        assert exported.shape == (1800, 21) and np.isfinite(exported).all()
        assert np.array_equal(exported, expected)
        assert all(cell.number_format == "0.0000"
                   for row in sheet.iter_rows(min_row=2, max_row=1801, min_col=2, max_col=22)
                   for cell in row)
        margin = np.array([[rounding_margin(v) for v in row] for row in field])
        margins_all.append(margin)
        assert np.max(np.abs(margin - saved_margins[k])) < 1e-13
        unresolved = margin[1:] <= budget[1:]
        with (ROOT / f"results/q1/table{k+1}.csv").open(encoding="utf-8-sig", newline="") as stream:
            paper = list(csv.reader(stream))
        assert len(paper) == 8
        checked_paper = 0
        for row_number, t in enumerate(paper_times, 1):
            assert float(paper[row_number][0]) == t
            for col_number, j in enumerate(paper_columns, 1):
                assert Decimal(paper[row_number][col_number]) == rounded(field[t, j])
                assert float(paper[row_number][col_number]) == exported[t-1, j]
                certificates.append({"quantity": name, "unit": "degC" if k == 0 else "kg/kg dry basis",
                    "time_s": t, "radius_cm": float(radius[j] * 100),
                    "authority_full_precision": format(field[t, j], ".17g"),
                    "published_4dp": str(rounded(field[t, j])),
                    "empirical_numerical_budget": float(budget[t, j]),
                    "distance_to_rounding_boundary": float(margin[t, j]),
                    "margin_over_budget": float(margin[t, j] / budget[t, j]),
                    "independent_method_value": format(reference[t, j], ".17g"),
                    "absolute_independent_difference": float(abs(field[t, j] - reference[t, j])),
                    "environment_Ta_degC": float(Ta[t]), "effective_Ce_kg_kg": float(Ce[t]),
                    "reason_code": ("HEAT_BESSEL_ROBIN" if k == 0 else "LOCAL_D_CONSERVATIVE_ROBIN"),
                    "claim_limit": "Numerical evidence in R1; neither physical error nor strict global error bound"})
                checked_paper += 1
        paper_budget = budget[np.ix_(paper_times, paper_columns)]
        report.append({"field": name, "checked_excel_values": exported.size,
            "checked_paper_values": checked_paper, "unresolved_4dp_under_declared_budget": int(unresolved.sum()),
            "full_budget_maximum": maximum(budget[1:], time[1:], radius),
            "paper_budget_maximum": maximum(paper_budget, np.array(paper_times), radius[paper_columns]),
            "independent_method_maximum_difference": maximum((field-reference)[1:], time[1:], radius),
            "maximum_export_rounding_change": maximum(expected-field[1:], time[1:], radius),
            "minimum_margin_over_budget": float(np.min(margin[1:] / budget[1:])),
            "exact_decimal_vs_saved_margin_maxdiff": float(np.max(abs(margin-saved_margins[k])))})
    wb.close()
    template.close()

    R, L = p["radius_m"], p["length_m"]
    rho, cp, conductivity = p["density_kg_m3"], p["heat_capacity_J_kgK"], p["conductivity_W_mK"]
    hT, hC = p["heat_transfer_W_m2K"], p["moisture_transfer_m_s"]
    prefactor, exponent = p["diffusivity_prefactor_m2_s"], p["diffusivity_exponent_kg_kg"]
    alpha = conductivity/(rho*cp)
    Dinitial = prefactor * np.exp(-exponent/p["initial_moisture_kg_kg"])
    Dsurface = prefactor * np.exp(-exponent/C[:, -1])
    volumes = np.pi * L * np.diff(np.linspace(0., R, 32769)**2)
    final_checkpoint = FULL / "Radau-N32768-r1e-12-t1800-candidate/checkpoint-1800.npz"
    with np.load(final_checkpoint, allow_pickle=False) as z:
        assert float(z["t_s"]) == 1800
        ring_mean = float(volumes @ z["ring_C"] / volumes.sum())
    mechanics = {
        "alpha_m2_s": float(alpha), "Bi_T_radius": hT*R/conductivity,
        "D_initial_m2_s": float(Dinitial), "Bi_C_initial_radius": float(hC*R/Dinitial),
        "radial_thermal_diffusion_time_s": R**2/alpha,
        "radial_initial_moisture_diffusion_time_s": float(R**2/Dinitial),
        "Fo_T_at1800": alpha*1800/R**2, "Fo_C_initial_at1800": float(Dinitial*1800/R**2),
        "sqrt_alpha_t_at1800_m": float(np.sqrt(alpha*1800)),
        "sqrt_Dinitial_t_at1800_m": float(np.sqrt(Dinitial*1800)),
        "half_length_m": L/2, "length_over_diameter": L/(2*R),
        "end_to_side_area_ratio": R/L,
        "volume_m3": np.pi*R**2*L,
        "final_T_center_degC": float(T[-1, 0]), "final_T_surface_degC": float(T[-1, -1]),
        "final_Ta_degC": float(Ta[-1]), "final_C_center_kg_kg": float(C[-1, 0]),
        "final_C_surface_kg_kg": float(C[-1, -1]), "final_Ce_effective_kg_kg": float(Ce[-1]),
        "final_surface_D_m2_s": float(Dsurface[-1]),
        "final_T_r_surface_degC_m_from_Robin": float(hT*(Ta[-1]-T[-1, -1])/conductivity),
        "final_C_r_surface_per_m_from_Robin": float(-hC*(C[-1, -1]-Ce[-1])/Dsurface[-1]),
        "final_inward_heat_flux_W_m2": float(hT*(Ta[-1]-T[-1, -1])),
        "final_outward_C_flux_m_s": float(hC*(C[-1, -1]-Ce[-1])),
        "final_ring_volume_mean_C_kg_kg": ring_mean,
        "initial_corner_moisture_flux_m_s": hC*(p["initial_moisture_kg_kg"]-Ce[0]),
        "limits": ["Diffusion lengths are scale estimates, not penetration fronts or error bounds.",
                   "Robin-derived gradients check sign and magnitude implications; they are not independent boundary residuals.",
                   "Volume-mean C and its integral are not measured water mass; no dry-density conversion is introduced."]}
    assert digest(authority_path) == authority_before and digest(results_path) == result_before
    result = {"status": "READBACK_AND_POINT_CERTIFICATE_PASS", "created_utc": datetime.now(timezone.utc).isoformat(),
        "command": [sys.executable, *sys.argv], "source_sha256": digest(Path(__file__)),
        "versions": {"python": sys.version, "numpy": np.__version__, "openpyxl": openpyxl.__version__},
        "original_input_checks": original_checks,
        "environment": {"headers": list(headers), "all_rows": 241, "Q1_nodes": q1.tolist(), "time_step_s": 60,
                        "no_missing_duplicate_or_irregular_time": True},
        "authority_sha256": authority_before, "excel_sha256": result_before,
        "fields": report, "mechanistic_scales": mechanics,
        "physical_accuracy_certified": False, "official_standard_answer_available": False,
        "existing_numerical_outputs_unchanged": True, "figures_drawn": 0,
        "human_review_performed": False, "full_PDE_replay": "Separate q1_reproduce.py --run; consult its actual run.json"}
    with (OUT / "paper-70-evidence.csv").open("x", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(certificates[0]))
        writer.writeheader()
        writer.writerows(certificates)
    write_new(OUT / "data-and-point-audit.json", result)
    print(json.dumps({k: result[k] for k in ["status", "fields", "mechanistic_scales"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
