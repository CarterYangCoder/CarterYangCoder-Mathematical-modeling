"""Q1 authority -> original workbook, paper tables and independent full readback.

No PDE solver is called. Project Python handles NPZ/gating/rounding/readback.
Bundled artifact-tool imports the template, writes typed values and renders.
"""
from __future__ import annotations

import argparse
import csv
from copy import copy
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import subprocess
import sys
import time

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
NAMES = ("温度", "水分浓度")
TEMPLATE = ROOT / "A题/附件/附件3/result1.xlsx"
TEMPLATE_SHA = "23b261b295c1b787d000eebbca6521c37075107b6fcf78724f8d395ce1798ff4"
A1 = "时间\\到药材中心的距离"
TABLE_TIMES = (100, 300, 600, 900, 1200, 1500, 1800)
TABLE_RADIAL_INDEX = (0, 5, 10, 15, 20)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def project_path(value: str | Path) -> Path:
    path = Path(value)
    path = path.resolve() if path.is_absolute() else (ROOT / path).resolve()
    if not path.is_relative_to(ROOT):
        raise ValueError(f"Q1 project path leaves workspace: {path}")
    return path


def round_half_up(values: np.ndarray) -> np.ndarray:
    """Round exact stored binary64 values; do not multiply by a float scale."""
    array = np.asarray(values, dtype=np.float64)
    if not np.all(np.isfinite(array)) or np.any(array < 0):
        raise ValueError("Export requires finite nonnegative physical values")
    quantum = Decimal("0.0001")
    return np.array([float(Decimal.from_float(float(x)).quantize(quantum, rounding=ROUND_HALF_UP))
                     for x in array.flat], dtype=np.float64).reshape(array.shape)


def load_authority(path: Path) -> tuple[np.ndarray, np.ndarray, list[np.ndarray]]:
    with np.load(path, allow_pickle=False) as result:
        required = {"time_s", "radius_m", "T_C", "C_kg_kg"}
        if not required.issubset(result.files):
            raise ValueError(f"Authority missing keys: {sorted(required - set(result.files))}")
        times = np.array(result["time_s"], dtype=np.float64)
        radius = np.array(result["radius_m"], dtype=np.float64)
        fields = [np.array(result[key], dtype=np.float64) for key in ("T_C", "C_kg_kg")]
    if not np.array_equal(times, np.arange(1801, dtype=np.float64)):
        raise ValueError("Authority must contain exactly t=0,...,1800 seconds")
    if radius.shape != (21,) or not np.allclose(radius, np.arange(21)*.001, rtol=0, atol=1e-15):
        raise ValueError("Authority must contain r=0,...,0.020 m at 0.001 m spacing")
    for field, initial in zip(fields, (28., 2.55)):
        if field.shape != (1801, 21) or not np.all(np.isfinite(field)):
            raise ValueError("Authority field must be finite, shape (1801,21)")
        if not np.all(field[0] == initial):
            raise ValueError("Authority t=0 must preserve the uniform initial condition exactly")
    return times, radius, fields


def check_evidence_bindings(acceptance: dict) -> dict:
    """Verify current evidence bytes, not just the saved acceptance report."""
    checked = {}

    def verify_file(value, expected, label, require_absolute=False):
        if not isinstance(value, str) or not value:
            raise ValueError(f"{label} path must be a nonempty string")
        if require_absolute and not Path(value).is_absolute():
            raise ValueError(f"{label} requires absolute project paths")
        path = project_path(value)  # resolve also rejects junctions outside ROOT
        if not path.is_file():
            raise ValueError(f"{label} file missing: {path}")
        if (not isinstance(expected, str) or len(expected) != 64
                or any(char not in "0123456789abcdefABCDEF" for char in expected)):
            raise ValueError(f"{label} SHA-256 must contain 64 hexadecimal characters")
        actual = digest(path)
        if actual != expected.lower():
            raise ValueError(f"{label} changed since numerical acceptance: {path}")
        return str(path), actual

    if "evidence_sha256" in acceptance:
        evidence = acceptance["evidence_sha256"]
        if not isinstance(evidence, dict) or not evidence:
            raise ValueError("evidence_sha256 must be a nonempty absolute-path-to-SHA mapping")
        for value, expected in evidence.items():
            path, actual = verify_file(value, expected, "Evidence", require_absolute=True)
            checked[path] = actual
    analysis = None
    if "analysis_path" in acceptance or "analysis_sha256" in acceptance:
        if not {"analysis_path", "analysis_sha256"}.issubset(acceptance):
            raise ValueError("analysis_path and analysis_sha256 must be provided together")
        path, actual = verify_file(acceptance["analysis_path"], acceptance["analysis_sha256"], "Analysis")
        analysis = {"path": path, "sha256": actual}
    return {"evidence_sha256": checked, "checked_evidence_files": len(checked), "analysis": analysis}


def check_numerical_gate(acceptance: dict, authority: Path, candidate: bool) -> dict:
    if type(candidate) is not bool:
        raise ValueError("candidate must be a JSON/Python boolean, never a string")
    authority_sha = digest(authority)
    if acceptance.get("authority_sha256") != authority_sha:
        raise ValueError("Numerical acceptance is not bound to this authority SHA-256")
    count = acceptance.get("unresolved_rounding_count")
    if type(count) is not int or count < 0:
        raise ValueError("Acceptance must give explicit nonnegative unresolved_rounding_count")
    passed = acceptance.get("numerical_pass") is True
    if not candidate and (not passed or count != 0):
        raise ValueError("Formal export requires numerical_pass=true and unresolved_rounding_count=0")
    evidence = check_evidence_bindings(acceptance)
    return {"numerical_pass": passed, "unresolved_rounding_count": count,
            "authority_sha256": authority_sha, "candidate_requested": candidate,
            "current_evidence": evidence}


def template_styles() -> dict:
    workbook = openpyxl.load_workbook(TEMPLATE, read_only=False, data_only=False)
    if workbook.sheetnames != list(NAMES):
        raise ValueError("Original template sheet names/order changed")
    info = {}
    for name in NAMES:
        sheet = workbook[name]
        if sheet["A1"].value != A1 or [sheet.cell(row, 1).value for row in (2, 3, 4)] != [1, 2, 3]:
            raise ValueError("Original template axes changed")
        info[name] = {key: copy(getattr(sheet["A1"], key)) for key in
                      ("font", "fill", "border", "alignment", "protection", "number_format")}
        info[name]["merged"] = list(map(str, sheet.merged_cells.ranges))
        info[name]["a_width"] = sheet.column_dimensions["A"].width
    workbook.close()
    return info


def verify_xlsx(path: Path, rounded: list[np.ndarray], styles: dict) -> dict:
    """Read every one of 75,600 numeric results and all 70 paper-table links."""
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    if workbook.sheetnames != list(NAMES):
        raise AssertionError("Exported sheet names/order incorrect")
    checks = []
    count = paper_count = 0
    for name, field in zip(NAMES, rounded):
        sheet = workbook[name]
        if sheet.max_row != 1801 or sheet.max_column != 22:
            raise AssertionError(f"Wrong worksheet dimensions: {name}, {sheet.max_row}x{sheet.max_column}")
        if sheet["A1"].value != A1:
            raise AssertionError("A1 changed")
        if list(map(str, sheet.merged_cells.ranges)) != styles[name]["merged"]:
            raise AssertionError("Merged cells changed")
        style_checks = {key: copy(getattr(sheet["A1"], key)) == styles[name][key] for key in
                        ("font", "fill", "border", "alignment", "protection", "number_format")}
        if not all(style_checks.values()):
            raise AssertionError(f"A1 template style not preserved: {name}: {style_checks}")
        if abs(sheet.column_dimensions["A"].width - styles[name]["a_width"]) > 1e-12:
            raise AssertionError("A1 column width changed")
        for row in range(2, 1802):
            cell = sheet.cell(row, 1)
            if cell.value != row-1 or cell.data_type != "n":
                raise AssertionError(f"Wrong/non-numeric time at {name}!A{row}")
        for col in range(2, 23):
            cell = sheet.cell(1, col)
            if cell.value != (col-2)/10 or cell.data_type != "n":
                raise AssertionError(f"Wrong/non-numeric radius at {name}!{cell.coordinate}")
        for row in range(2, 1802):
            for col in range(2, 23):
                cell = sheet.cell(row, col)
                expected = float(field[row-2, col-2])
                if cell.data_type != "n" or type(cell.value) not in (int, float):
                    raise AssertionError(f"Non-numeric result at {name}!{cell.coordinate}")
                if not np.isfinite(cell.value) or cell.value != expected:
                    raise AssertionError(f"Value mismatch at {name}!{cell.coordinate}: {cell.value} != {expected}")
                if cell.number_format != "0.0000":
                    raise AssertionError(f"Wrong format at {name}!{cell.coordinate}")
                count += 1
        for t in TABLE_TIMES:
            for j in TABLE_RADIAL_INDEX:
                if sheet.cell(t+1, j+2).value != float(field[t-1, j]):
                    raise AssertionError("Paper table vs workbook mismatch")
                paper_count += 1
        checks.append({"sheet": name, "result_shape": [1800, 21], "time_s": [1, 1800], "radius_cm": [0, 2],
                       "a1_value_preserved": True, "a1_style_preserved": style_checks,
                       "formulas": sum(cell.data_type == "f" for row in sheet for cell in row)})
    workbook.close()
    if count != 75600 or paper_count != 70:
        raise AssertionError("Readback coverage incomplete")
    return {"readback_pass": True, "checked_numeric_results": count, "checked_paper_table_links": paper_count, "sheets": checks}


def paper_tables(output_dir: Path, rounded: list[np.ndarray], authority_sha: str) -> list[Path]:
    files = []
    for index, (name, field, unit) in enumerate(zip(NAMES, rounded, ("℃", "kg/kg")), 1):
        table = field[np.array(TABLE_TIMES)-1][:, TABLE_RADIAL_INDEX]
        labels = ["时间/s", "r=0 cm", "r=0.5 cm", "r=1 cm", "r=1.5 cm", "r=2 cm"]
        stem = output_dir / f"table{index}"
        csv_path = stem.with_suffix(".csv")
        with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(labels)
            for t, row in zip(TABLE_TIMES, table):
                writer.writerow([t, *[f"{x:.4f}" for x in row]])
        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        if rows[0] != labels or [int(row[0]) for row in rows[1:]] != list(TABLE_TIMES):
            raise AssertionError("Paper CSV axes changed")
        if not np.array_equal(np.array([[float(x) for x in row[1:]] for row in rows[1:]]), table):
            raise AssertionError("Paper CSV numeric readback mismatch")
        md = [f"表{index}　{name}（{unit}）", "", "| "+" | ".join(labels)+" |",
              "| "+" | ".join(["---:"]*6)+" |"]
        md += ["| "+" | ".join([str(t), *[f"{x:.4f}" for x in row]])+" |" for t, row in zip(TABLE_TIMES, table)]
        md += ["", f"数值来源：同一全精度权威结果 SHA-256 {authority_sha}；统一按四位小数 half-up 导出。", ""]
        md_path = stem.with_suffix(".md")
        md_path.write_text("\n".join(md), encoding="utf-8")
        tex = [f"% authority SHA-256: {authority_sha}", r"\begin{table}[htbp]", r"\centering",
               f"\\caption{{{name}（{unit}）}}", f"\\label{{tab:q1-{index}}}", r"\begin{tabular}{rrrrrr}",
               r"\hline", r"时间/s & 0 cm & 0.5 cm & 1 cm & 1.5 cm & 2 cm \\", r"\hline"]
        tex += [" & ".join([str(t), *[f"{x:.4f}" for x in row]])+r" \\" for t, row in zip(TABLE_TIMES, table)]
        tex += [r"\hline", r"\end{tabular}", r"\end{table}", ""]
        tex_path = stem.with_suffix(".tex")
        tex_path.write_text("\n".join(tex), encoding="utf-8")
        files += [csv_path, md_path, tex_path]
    return files


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--authority")
    parser.add_argument("--acceptance")
    parser.add_argument("--output-dir", default="results/q1")
    parser.add_argument("--audit-dir", default="records/q1-full-20260911/export")
    parser.add_argument("--candidate", action="store_true")
    parser.add_argument("--preflight", action="store_true", help="Read-only validation; no outputs or workbook authoring")
    parser.add_argument("--job", type=Path)
    args = parser.parse_args()
    if args.job:
        parameters = json.loads(args.job.read_text(encoding="utf-8"))["parameters"]
        for key in ("authority", "acceptance", "output_dir", "audit_dir", "candidate"):
            if key in parameters:
                setattr(args, key, parameters[key])
    if not args.authority or not args.acceptance:
        parser.error("--authority and --acceptance required (directly or in --job parameters)")
    if type(args.candidate) is not bool:
        raise ValueError("Job candidate must be an explicit JSON boolean")
    config = json.loads((ROOT/"project_config.json").read_text(encoding="utf-8"))
    project_python = Path(config["python"].replace("$"+"{USERPROFILE}", os.environ["USERPROFILE"])).resolve()
    if Path(sys.executable).resolve() != project_python:
        raise ValueError("Use project_config.json's isolated Python")
    authority = project_path(args.authority)
    acceptance_path = project_path(args.acceptance)
    output_dir = project_path(args.output_dir)
    audit_dir = project_path(args.audit_dir)
    if not output_dir.is_relative_to(ROOT/"results") or not audit_dir.is_relative_to(ROOT/"records"):
        raise ValueError("Outputs belong in results; export audit belongs in records")
    outfile = output_dir/"result1.xlsx"
    planned_files = [outfile, output_dir/"q1-export-manifest.json"]
    planned_files += [output_dir/f"table{i}.{suffix}" for i in (1, 2) for suffix in ("csv", "md", "tex")]
    if any(path.exists() for path in planned_files):
        raise FileExistsError("Existing export retained; choose a new explicit version directory")
    if (audit_dir/"export-audit.json").exists():
        raise FileExistsError("Existing export audit retained; choose a new explicit audit directory")
    if digest(TEMPLATE) != TEMPLATE_SHA:
        raise ValueError("Original result1.xlsx template SHA changed")
    acceptance = json.loads(acceptance_path.read_text(encoding="utf-8"))
    gate = check_numerical_gate(acceptance, authority, args.candidate)
    times, radius_m, fields = load_authority(authority)
    rounded = [round_half_up(field[1:]) for field in fields]
    styles = template_styles()
    originals = {str(p.relative_to(ROOT)): digest(p) for p in sorted(TEMPLATE.parent.glob("*.xlsx"))}
    if args.preflight:
        print(json.dumps({"status": "READ_ONLY_EXPORT_PREFLIGHT_PASS", "numerical_gate": gate,
                          "authority_field_shape": [1801, 21], "template_sha256": TEMPLATE_SHA,
                          "expected_export_values": 75600, "outputs_written": False}, ensure_ascii=False))
        return 0
    output_dir.mkdir(parents=True, exist_ok=True)
    audit_dir.mkdir(parents=True, exist_ok=True)
    manifest = {"status": "EXPORTING_CANDIDATE" if args.candidate else "EXPORTING_AWAITING_READBACK",
                "created_utc": datetime.now(timezone.utc).isoformat(), "command": [sys.executable, *sys.argv],
                "numerical_gate": gate, "acceptance_sha256": digest(acceptance_path),
                "template_sha256": TEMPLATE_SHA, "original_templates": originals,
                "export_parameters": {"authority": str(authority.relative_to(ROOT)),
                                      "acceptance": str(acceptance_path.relative_to(ROOT)),
                                      "output_dir": str(output_dir.relative_to(ROOT)),
                                      "audit_dir": str(audit_dir.relative_to(ROOT)),
                                      "candidate": args.candidate, "workers": 1, "random_seed": None,
                                      "time_unit": "s", "radial_header_unit": "cm",
                                      "temperature_unit": "degC", "moisture_unit": "kg/kg"},
                "project_config_sha256": digest(ROOT/"project_config.json"),
                "job_sha256": digest(args.job) if args.job else None,
                "rounding": "Stored binary64 -> Decimal ROUND_HALF_UP to 0.0001; numeric values, 0.0000 display",
                "versions": {key: importlib.metadata.version(key) for key in ("numpy", "openpyxl")},
                "python": sys.version, "source_sha256": {str(p.relative_to(ROOT)): digest(p) for p in
                         (Path(__file__), ROOT/"src/q1_export_workbook.mjs")},
                "human_review_performed": False, "submission_release_ready": False}
    manifest_path = audit_dir/"export-audit.json"
    save_json(manifest_path, manifest)
    start = time.perf_counter()
    try:
        payload_path = audit_dir/"export-payload.json"
        save_json(payload_path, {"times": times[1:].astype(int).tolist(), "radius_cm": (np.arange(21)/10).tolist(),
                                 "values": [x.tolist() for x in rounded], "authority_sha256": gate["authority_sha256"]})
        user = Path(os.environ["USERPROFILE"])
        deps = user/".cache/codex-runtimes/codex-primary-runtime/dependencies"
        node = deps/"node/bin/node.exe"
        runtime = ROOT/"records/q1-full-20260911/artifact-preview"
        if not node.is_file() or not (runtime/"node_modules").is_dir():
            raise RuntimeError("Discovered bundled Node/artifact junction unavailable; no automatic installation")
        marker = user/".codex/plugins/cache/openai-primary-runtime/spreadsheets/26.909.12148/skills/spreadsheets/container_tools/mark_artifact_operation_started.mjs"
        marker_state = runtime/"authoring-operation-started.json"
        if not marker_state.exists():
            mark_command = [str(node), str(marker), "--operation-kind", "edit", "--expected-output-count", "1", "--output-format", "xlsx"]
            marked = subprocess.run(mark_command, cwd=marker.parent.parent, capture_output=True, text=True, encoding="utf-8", timeout=90)
            (audit_dir/"artifact-operation-marker.log").write_text(marked.stdout+marked.stderr, encoding="utf-8")
            if marked.returncode:
                raise RuntimeError("Artifact operation marker failed before workbook authoring")
            save_json(marker_state, {"command": mark_command, "time_utc": datetime.now(timezone.utc).isoformat()})
        command = [str(node), str(ROOT/"src/q1_export_workbook.mjs"), "--runtime", str(runtime),
                   "--template", str(TEMPLATE), "--preview", str(audit_dir/"preview"),
                   "--payload", str(payload_path), "--output", str(outfile)]
        executed = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=180)
        (audit_dir/"artifact-export.log").write_text(executed.stdout+executed.stderr, encoding="utf-8")
        if executed.returncode:
            raise RuntimeError("Artifact workbook export failed; see export log")
        manifest["readback"] = verify_xlsx(outfile, rounded, styles)
        table_files = paper_tables(output_dir, rounded, gate["authority_sha256"])
        if {relative: digest(ROOT/relative) for relative in originals} != originals:
            raise AssertionError("An original template changed")
        if digest(authority) != gate["authority_sha256"] or digest(acceptance_path) != manifest["acceptance_sha256"]:
            raise AssertionError("Authority or acceptance changed during export")
        if any(digest(ROOT/relative) != expected for relative, expected in manifest["source_sha256"].items()):
            raise AssertionError("Export source changed during execution")
        if digest(ROOT/"project_config.json") != manifest["project_config_sha256"]:
            raise AssertionError("Project runtime configuration changed during export")
        if args.job and digest(args.job) != manifest["job_sha256"]:
            raise AssertionError("Export job changed during execution")
        if check_evidence_bindings(acceptance) != gate["current_evidence"]:
            raise AssertionError("Numerical evidence bindings changed during export")
        manifest.update(status="CANDIDATE_EXPORTED_READBACK_PASS" if args.candidate else "Q1_NUMERICAL_EXPORT_READBACK_PASS",
                        export_readback_pass=True, formal_q1_result=not bool(args.candidate),
                        visual_review="Previews generated; actual AI visual inspection is separately recorded before final delivery",
                        original_templates_unchanged=True, elapsed_seconds=time.perf_counter()-start,
                        output_sha256={str(p.relative_to(ROOT)): digest(p) for p in [outfile, *table_files]})
        save_json(manifest_path, manifest)
        save_json(output_dir/"q1-export-manifest.json", manifest)
        print(json.dumps({"status": manifest["status"], "workbook": str(outfile),
                          "checked_numeric_results": 75600, "checked_paper_table_links": 70,
                          "human_review_performed": False}, ensure_ascii=False))
        return 0
    except Exception as exc:
        manifest.update(status="EXPORT_FAILED_NOT_FORMAL", export_readback_pass=False, formal_q1_result=False,
                        error=str(exc), elapsed_seconds=time.perf_counter()-start)
        save_json(manifest_path, manifest)
        raise


if __name__ == "__main__":
    sys.exit(main())
