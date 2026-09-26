"""Known-answer checks in temporary folders; no demo outputs in production."""
import importlib.metadata
from pathlib import Path
import tempfile
import sys


def run():
    if sys.flags.optimize: raise RuntimeError('Functional checks cannot run with Python optimization enabled')
    import numpy as np
    import pandas as pd
    import sympy as sp
    from scipy.optimize import milp,Bounds,LinearConstraint
    from shapely.geometry import Polygon
    from pymoo.indicators.hv import HV
    from docx import Document
    from ingest import extract
    from production import compare
    from precheck import identity_findings,scan_file
    sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
    from plotstyle import style,save
    import matplotlib.pyplot as plt
    checks={}
    def require(condition,name):
        if not condition: raise AssertionError(name)
        checks[name]=True
    with tempfile.TemporaryDirectory(prefix='cumcm-selftest-') as temporary:
        root=Path(temporary)
        doc=Document();doc.add_paragraph('附件读取核验')
        table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text='变量';table.cell(1,1).text='42'
        doc.save(root/'input.docx')
        parsed=extract(root/'input.docx')
        require(parsed['paragraphs']==['附件读取核验'] and parsed['tables'][0][1][1]=='42','docx_content')
        pd.DataFrame({'位置':[-1,0,1],'水深':[72,70,68]}).to_excel(root/'input.xlsx',index=False)
        parsed=extract(root/'input.xlsx')
        require(parsed['sheets']['Sheet1'][2]==[0,70],'xlsx_values')
        with style('中文坐标负数测量覆盖宽度距离米'):
            fig,ax=plt.subplots(figsize=(5,3),layout='constrained')
            ax.plot([-1,0,1],[72,70,68],marker='o')
            ax.set(xlabel='距离（米）',ylabel='测量',title='中文坐标：负数')
            save(fig,root/'figure');plt.close(fig)
        require((root/'figure.pdf').stat().st_size>1000,'chinese_plot_missing_glyph_guard')
        pdf=extract(root/'figure.pdf')
        require('距离' in ''.join(p['text'] for p in pdf['pages']),'pdf_chinese_text')
        require(bool(identity_findings('C:\\Users\\private-name\\secret')),'windows_path_rejected')
        require(bool(identity_findings('private@example.com')),'email_rejected')
        require(bool(scan_file('x.tex','待填写'.encode(),placeholders=True)),'placeholder_rejected')
        require(not scan_file('x.csv',b'x,y\n1,2\n'),'clean_csv_accepted')
    x=sp.symbols('x',real=True)
    require(sp.simplify(sp.diff(sp.exp(-x),x)+sp.exp(-x))==0,'symbolic_residual')
    solution=milp([-3,-2],integrality=[1,1],bounds=Bounds([0,0],[10,10]),constraints=LinearConstraint([[2,1],[1,2]],-np.inf,[10,10]))
    require(solution.success and abs(solution.fun+16)<1e-8,'integer_known_optimum')
    require(abs(Polygon([(0,0),(2,0),(2,2),(0,2)]).intersection(Polygon([(1,1),(3,1),(3,3),(1,3)])).area-1)<1e-12,'shapely_intersection')
    require(abs(float(HV(ref_point=np.array([3.,3.]))(np.array([[1.,2.],[2.,1.]])))-3)<1e-12,'pymoo_hypervolume')
    require(compare({'v':[1.,2.]},{'v':[1.+1e-12,2.]}) and not compare({'v':[1.,2.]},{'v':[1.,3.]}),'numeric_replay_tolerance')
    return {'passed':len(checks),'total':len(checks),'checks':checks,'scope':'Known-answer temporary fixtures, not contest model validation'}
