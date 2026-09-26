"""Read attachments without modifying original data. No automatic OCR guesses."""
from pathlib import Path
from common import digest, write_json


def extract(path):
    path = Path(path)
    suffix = path.suffix.lower()
    result = {'sha256': digest(path), 'format': suffix}
    if suffix == '.pdf':
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            result['pages'] = [{'text': p.extract_text() or '', 'tables': p.extract_tables()} for p in pdf.pages]
        result['requires_ocr_or_visual_review'] = any(not p['text'].strip() for p in result['pages'])
    elif suffix == '.docx':
        from docx import Document
        doc = Document(path)
        result['paragraphs'] = [p.text for p in doc.paragraphs]
        result['tables'] = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    elif suffix in {'.xls','.xlsx'}:
        import pandas as pd
        sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
        result['sheets'] = {k: v.where(v.notna(), None).astype(object).values.tolist() for k,v in sheets.items()}
        # String conversion retains date labels; numeric values remain numeric.
        import json
        result = json.loads(json.dumps(result, default=str, ensure_ascii=False))
    elif suffix in {'.csv','.txt','.json'}:
        for encoding in ['utf-8-sig','gb18030']:
            try:
                result['text'] = path.read_text(encoding=encoding)
                result['encoding'] = encoding
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError('Unknown text encoding; do not guess silently.')
    else:
        raise ValueError('Unsupported attachment: '+suffix)
    return result


def ingest(root):
    root = Path(root)
    paths = sorted(p for p in (root/'data/raw').rglob('*') if p.is_file())
    if not paths:
        raise ValueError('No raw attachments. Put the actual statement and data in data/raw.')
    entries = []
    for path in paths:
        result = extract(path)
        relative = path.relative_to(root/'data/raw')
        destination = root/'data/processed'/relative.parent/(relative.name+'.extracted.json')
        write_json(destination,result)
        entries.append({'path':path.relative_to(root).as_posix(),'sha256':result['sha256'],
                        'extracted':destination.relative_to(root).as_posix()})
    write_json(root/'records/input-manifest.json',entries)
    return entries
